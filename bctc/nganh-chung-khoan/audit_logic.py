# -*- coding: utf-8 -*-
r"""audit_logic.py - LIET KE LOGIC cac chi tieu/ty le trong 3 sheet trinh bay (Key ratios, FS Industry, Drivers) cua workbook goc
de audit: doc cong thuc goc, dich tham chieu o -> [Sheet][Ten dong](+-k ky), GETPIVOTDATA -> DATA[metric | dong n | ky].
Ra: AUDIT-Logic-Sheet-TrinhBay.md (bang: dong | chi tieu | logic dich | cong thuc goc).
Chay: python audit_logic.py [--wb <workbook>]
"""
import argparse, os, re, sys
import openpyxl
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import render_workbook_feed as R

XREF = re.compile(r"(?:'([^']+)'|([A-Za-z_][A-Za-z0-9_]*))!(\$?[A-Z]{1,3}\$?\d+)(?::(\$?[A-Z]{1,3}\$?\d+))?")
REF = re.compile(r"(?<![A-Za-z_!'\"@])(\$?)([A-Z]{1,3})(\$?)(\d+)(?![\d(])")
RANGE = re.compile(r"(?<![!'\"A-Za-z])(\$?[A-Z]{1,3}\$?\d+):(\$?[A-Z]{1,3}\$?\d+)")
STR = re.compile(r'"[^"]*"')
PER = re.compile(r"^Q([1-4])-(\d{4})$")

# sheet -> (cot nhan, dong header ky, cac dong phu tro ky, cot row_order)
META = {
    "Key ratios":  dict(lab=6, hdr=4, helper={4, 36, 89, 90}, ro=None, lab2=5),   # dong 149-163: nhan o cot E (metric)
    "FS Industry": dict(lab=6, hdr=4, helper={4}, ro=None),
    "Drivers":     dict(lab=6, hdr=8, helper={8, 9, 10}, ro=None),
    "IS":          dict(lab=1, hdr=5, helper={5, 6, 7}, ro=3),
    "BS":          dict(lab=1, hdr=5, helper={5, 6, 7}, ro=2),
    "NOTE":        dict(lab=1, hdr=5, helper={5, 6, 7}, ro=3),
    "Peer Data":   dict(lab=6, hdr=1, helper={1, 2, 3}, ro=4),
}


def pidx(lab):
    """Nhan ky -> chi so quy lien tuc (Qn-yyyy -> y*4+n; nam yyyy -> Q4 nam do)."""
    s = str(lab).strip()
    m = PER.match(s)
    if m:
        return int(m.group(2)) * 4 + int(m.group(1))
    if re.fullmatch(r"\d{4}(\.0)?", s):
        return int(float(s)) * 4 + 4
    return None


def load(wbp):
    sheets = {}
    wbf = openpyxl.load_workbook(wbp, read_only=True)
    wbv = openpyxl.load_workbook(wbp, read_only=True, data_only=True)
    for sn, m in META.items():
        cells, vals = {}, {}
        for row in wbf[sn].iter_rows(max_row=420):
            for c in row:
                v = c.value
                if v is None:
                    continue
                if hasattr(v, "text") and not isinstance(v, str):
                    v = "=" + str(v.text).lstrip("=")
                cells[(c.row, c.column)] = v
        for row in wbv[sn].iter_rows(min_row=1, max_row=420):
            for c in row:
                if c.value is not None:
                    vals[(c.row, c.column)] = c.value
        labels = {}                                                  # nhan = GIA TRI da tinh (o nhan co the la cong thuc)
        for (r, c), v in vals.items():
            if c == m["lab"] and str(v).strip():
                labels[r] = str(v).strip()
        if m.get("lab2"):
            for (r, c), v in vals.items():
                if c == m["lab2"] and r not in labels and str(v).strip():
                    labels[r] = str(v).strip()
        ro = m["ro"]
        hdr = {c: pidx(v) for (r, c), v in vals.items() if r == m["hdr"] and pidx(v) is not None}
        first_per = min(hdr) if hdr else 999
        sheets[sn] = dict(cells=cells, vals=vals, labels=labels, hdr=hdr, ro=ro, first=first_per, lab=m["lab"], hdr_row=m["hdr"], helper=m["helper"])
    # sheet 'Số TK mở mới': Table2 doc theo COT (header dong 7) -> tham chieu doc = ten cot
    tk = {}
    for row in wbv["Số TK mở mới"].iter_rows(min_row=7, max_row=7):
        for c in row:
            if c.value is not None:
                tk[c.column] = str(c.value).strip()
    sheets["Số TK mở mới"] = dict(cells={}, labels={}, hdr={}, ro={}, first=999, lab=0, hdr_row=7, helper=set(), tkhdr=tk)
    wbf.close(); wbv.close()
    # row_order IS/BS/NOTE la cong thuc =C8+1 -> tinh tu dong dau
    for sn in META:
        S = sheets[sn]; col = S["ro"]
        if not col:
            S["ro"] = {}; continue
        cur = None; out = {}
        for r in range(1, 421):
            v = S["cells"].get((r, col))
            if v is None:
                continue
            if isinstance(v, str) and v.startswith("="):
                mm = re.match(r"=([A-Z]+)(\d+)\+1$", v)
                cur = (out.get(int(mm.group(2))) or 0) + 1 if mm else None
            else:
                try:
                    cur = int(float(v))
                except (TypeError, ValueError):
                    cur = None
            if cur is not None:
                out[r] = cur
        S["ro"] = out
    return sheets


class Ctx:
    def __init__(self, sheets, sheet, cur_col):
        self.sh, self.s, self.cur = sheets, sheet, cur_col
        self.cur_p = sheets[sheet]["hdr"].get(cur_col)

    def name(self, tsheet, col, row, with_sheet):
        T = self.sh.get(tsheet)
        if T is None:
            return f"{tsheet}!{col}{row}"
        c = CI(col)
        pre = f"{tsheet}!" if with_sheet else ""
        if "tkhdr" in T:
            return f"SốTK[cột {T['tkhdr'].get(c, col)}]"
        if row in T["helper"] and c >= T["first"]:
            k = (T["hdr"].get(c) or 0) - (self.cur_p or 0) if (T["hdr"].get(c) and self.cur_p) else 0
            return "kỳ cột" if k == 0 else f"kỳ cột{k:+d}"
        if c < T["first"]:                                           # cot nhan/row_order/metric/o chon -> lay chu
            v = T.get("vals", {}).get((row, c))
            if v is not None and str(v).strip() != "":
                return f'"{str(v).strip()}"'
            lab = T["labels"].get(row)
            return f"{pre}[{lab}]#{row} (ô @{col}{row} trống)" if lab else f"{pre}(ô @{col}{row} trống)"
        lab = T["labels"].get(row)
        p = T["hdr"].get(c)
        k = (p - self.cur_p) if (p is not None and self.cur_p is not None) else None
        if lab is None:
            return f"{tsheet}!{col}{row}"
        ro = T["ro"].get(row) if isinstance(T["ro"], dict) else None
        s = f"{pre}[{lab}" + (f" | dòng {ro}" if ro else "") + f"]#{row}"   # #row = so dong tren sheet (phan biet nhan trung, vd 'Bonds' 80/87/94)
        if k is None:
            s += f"(@{col})"
        elif k != 0:
            s += f"({k:+d} kỳ)"
        return s

    def rng(self, tsheet, a, b, with_sheet):
        ma, mb = REF.match(a.replace("$", "")), REF.match(b.replace("$", ""))
        ca, ra, cb, rb = ma.group(2), int(ma.group(4)), mb.group(2), int(mb.group(4))
        if "tkhdr" in self.sh.get(tsheet, {}):
            return self.name(tsheet, ca, ra, with_sheet) + (f" (dòng {ra}..{rb})" if ca == cb else f"..{cb}")
        if ra == rb:
            n = CI(cb) - CI(ca) + 1
            return f"{self.name(tsheet, cb, rb, with_sheet)} {n} kỳ gần nhất"
        if ca == cb:
            T = self.sh.get(tsheet, {}); lab = T.get("labels", {})
            return f"{tsheet + '!' if with_sheet else ''}[{lab.get(ra, ca + str(ra))} .. {lab.get(rb, cb + str(rb))} ({rb - ra + 1} dòng)]"
        return f"{tsheet}!{a}:{b}"

    def expr(self, e):
        e = e.strip()
        if STR.fullmatch(e):
            return e.strip('"')
        def xr(m):
            t = m.group(1) or m.group(2); a, b = m.group(3).replace("$", ""), m.group(4)
            if b:
                return self.rng(t, a, b.replace("$", ""), True)
            mm = REF.match(a)
            return self.name(t, mm.group(2), int(mm.group(4)), True)
        e = XREF.sub(xr, e)
        e = RANGE.sub(lambda m: self.rng(self.s, m.group(1), m.group(2), False), e)
        e = REF.sub(lambda m: self.name(self.s, m.group(2), int(m.group(4)), False), e)
        return e


def gp_translate(call, ctx):
    a = R.split_args(call)
    src = a[1].strip() if len(a) > 1 else ""
    val = a[0].strip().strip('"')
    pairs = [(a[i].strip().strip('"'), ctx.expr(a[i + 1])) for i in range(2, len(a) - 1, 2)]
    d = dict(pairs)
    if "DATA" in src:
        s = "DATA["
        if "ticker" in d:
            s += f"mã={d['ticker']}; "
        s += f"{d.get('metric', '?')}"
        if "row_order" in d:
            s += f" | dòng {d['row_order']}"
        per = [d[k] for k in ("year", "quarter") if k in d]
        s += (" | " + ("kỳ cột" if any("kỳ" in p for p in per) else "/".join(per))) if per else " | mọi kỳ"
        return s + "]"
    name = "SốTK" if "TK" in src else "PV"
    return f"{name}[{val.replace('Sum of ', '')} | " + ", ".join(f"{f}={v}" for f, v in pairs) + "]"


SP_Y = re.compile(r"SUMPRODUCT\(\(RIGHT\(\$U\$4:\$BM\$4,4\)\+0=[A-Z]+\$4\)\*\$U(\d+):\$BM\1\)")
SP_Q4 = re.compile(r"SUMPRODUCT\(\(LEFT\(\$U\$4:\$BM\$4,2\)=\"(Q\d)\"\)\*\(RIGHT\(\$U\$4:\$BM\$4,4\)\+0=[A-Z]+\$4\)\*\$U(\d+):\$BM\2\)")


def translate(formula, sheets, sheet, cur_col):
    f = formula.replace("_xlfn.", "")
    ctx = Ctx(sheets, sheet, cur_col)
    if sheet == "FS Industry":                                        # khoi nam = gop cac cot quy U..BM cung nam
        labs = sheets[sheet]["labels"]
        f = SP_Q4.sub(lambda m: f"GiáTrị_{m.group(1)}_của_năm_cột([{labs.get(int(m.group(2)), m.group(2))}]#{m.group(2)})", f)
        f = SP_Y.sub(lambda m: f"Σ4quý_của_năm_cột([{labs.get(int(m.group(1)), m.group(1))}]#{m.group(1)})", f)
    for call in sorted(set(R.gp_calls(f)), key=len, reverse=True):
        f = f.replace("GETPIVOTDATA(" + call + ")", gp_translate(call, ctx))
    parts = STR.split(f); strs = STR.findall(f)
    out = ctx.expr(parts[0]) if parts[0] else ""
    for s, p in zip(strs, parts[1:]):
        out += s + (ctx.expr(p) if p else "")
    return out


def rows_table(sheets, sheet, col_pick, rows_range, title):
    S = sheets[sheet]; cells = S["cells"]; labels = S["labels"]
    lines = [f"\n## {title}\n", "| Dòng | Chỉ tiêu | Logic (đã dịch) | Công thức gốc |", "|---|---|---|---|"]
    n = 0
    for r in rows_range:
        lab = labels.get(r)
        f = None; c_used = None
        for c in col_pick:
            v = cells.get((r, c))
            if isinstance(v, str) and v.startswith("="):
                f, c_used = v, c; break
        if lab is None and f is None:
            continue
        if f is None:
            v = next((cells.get((r, c)) for c in col_pick if cells.get((r, c)) is not None), None)
            logic = "(nhập tay / giá trị)" if v is not None else "(tiêu đề nhóm)"
            raw = "" if v is None else str(v)[:60]
        else:
            logic = translate(f, sheets, sheet, c_used); raw = f
        esc = lambda s: str(s).replace("|", "\\|").replace("\n", " ")
        lines.append(f"| {r} | {esc(lab or '')} | {esc(logic)} | `{esc(raw)[:300]}` |"); n += 1
    return "\n".join(lines), n


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--wb", default=None); a = ap.parse_args()
    wbp = a.wb or os.path.join(os.path.expanduser("~"), "Downloads", "IB&Brokerage_Genea_2Q26.xlsx")
    sh = load(wbp)
    md = [f"# AUDIT logic sheet trình bày — nguồn: {os.path.basename(wbp)}\n",
          "Quy ước: `[Tên dòng]` = ô cùng kỳ của dòng đó (cùng sheet); `Sheet![Tên dòng | dòng n]` = sheet khác (n = row_order FiinProX);",
          "`(-1 kỳ)` = kỳ trước; `X 4 kỳ gần nhất` = 4 cột liền trước (TTM); `\"chuỗi\"` = nhãn lấy từ ô nhãn; `kỳ cột` = năm/quý của cột;",
          "`DATA[metric | dòng n | kỳ cột]` = GETPIVOTDATA trên pivot DATA (cộng mọi mã nếu không có `mã=`); `SốTK[...]` = pivot Số TK mở mới.",
          "Cột đại diện: quý mới nhất (Q2-2026) cho khối quý, 2025 cho khối năm; cột đầu (Q1-2016) chỉ khác ở chỗ chưa đủ 4 kỳ TTM."]
    kr = sh["Key ratios"]["hdr"]; cl = max(kr)
    t, n1 = rows_table(sh, "Key ratios", [cl, cl - 1, 9], range(5, 34), "Key ratios — khối NGÀNH (dòng 5–33)"); md.append(t)
    t, n2 = rows_table(sh, "Key ratios", [cl, cl - 1, 9], range(37, 65), "Key ratios — khối CÔNG TY chọn tại G36 (dòng 37–64)"); md.append(t)
    t, n3 = rows_table(sh, "Key ratios", [cl, cl - 1, 9], range(88, 166), "Key ratios — khối SỐ LIỆU GỐC của công ty chọn + dòng ngành 147–165 (dòng 88–165)"); md.append(t)
    fs = sh["FS Industry"]["hdr"]
    qc = [c for c, p in fs.items() if c >= CI("U")]; yc = [c for c, p in fs.items() if c < CI("U")]
    t, n4 = rows_table(sh, "FS Industry", [max(qc), max(qc) - 1], range(4, 142), "FS Industry — khối QUÝ (cột U..BM)"); md.append(t)
    t, n5 = rows_table(sh, "FS Industry", [max(yc), max(yc) - 1], range(4, 142), "FS Industry — khối NĂM (cột H..T)"); md.append(t)
    dr = sh["Drivers"]["hdr"]; cd = max(dr)
    t, n6 = rows_table(sh, "Drivers", [cd, cd - 1, cd - 2], range(8, 63), "Drivers"); md.append(t)
    out = os.path.join(HERE, "AUDIT-Logic-Sheet-TrinhBay.md")
    open(out, "w", encoding="utf-8").write("\n".join(md))
    print(f"-> {out}: Key ratios {n1}+{n2}+{n3}, FS Industry {n4}+{n5}, Drivers {n6} dòng")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    main()
