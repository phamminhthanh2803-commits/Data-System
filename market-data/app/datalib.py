# -*- coding: utf-8 -*-
r"""
datalib.py - lop DU LIEU cho Market Data App.

Doc lai TAT CA master CSV trong D:\market-data (khong ket noi mang), cache bang
parquet + st.cache_data, va tinh san cac chi bao phai suy ra (do rong, MA, chi so
nganh, nhom von hoa, GTGD).

Moi dataset khai bao 1 lan trong REGISTRY -> trang "Kho du lieu" va trang "Xuat Excel"
tu dong co day du, them pipeline moi chi can them 1 dong.
"""
from __future__ import annotations

import glob
import io
import os
import re
import sys
from datetime import datetime

import numpy as np
import pandas as pd
import streamlit as st

BASE = os.path.normpath(os.environ.get("MD_ROOT", "D:/market-data"))   # MD_ROOT (cloud); mac dinh D:\market-data
IDXD = os.path.join(BASE, "index-fetcher")
RAW = os.path.join(IDXD, "raw")
VALD = os.path.join(BASE, "market-valuation")
MACD = os.path.join(BASE, "macro-fetcher")
NSOD = os.path.join(BASE, "nso-fetcher")
TRAD = os.path.join(BASE, "transmission-fetcher")
BOND = os.path.join(BASE, "bond-pivot")
VSDC = os.path.join(BASE, "vsdc-accounts")
CACHE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "cache")
os.makedirs(CACHE, exist_ok=True)

VN_EX = ("HOSE", "HNX", "UPCOM")
VN_IDX = ["VNINDEX", "VN30", "HNXINDEX", "UPCOM"]

# ---------------------------------------------------------- KHUNG PHAN TICH (07/10/2026)
# Mot khung phan loai dung chung cho TOAN he thong: tab trang Vi mo, nhom dataset trong Kho du lieu,
# nhom chi tieu NSO, thu tu sheet Excel. Khoi -> nhom; moi dataset/chi tieu gan vao dung 1 nhom.
KHOI_THU_TU = {"Vĩ mô": 0, "Cổ phiếu": 1, "Khu vực": 2, "Trái phiếu": 3, "Dữ liệu": 4}
KHUNG_VI_MO = {       # tab trang Vi mo (thu tu hien thi) -> noi dung
    "Bảng điều khiển": "KPI + 4 biểu đồ chính",
    "Tăng trưởng & đầu tư": "GDP quý theo ngành, IIP, bán lẻ, vốn đầu tư toàn xã hội, FDI",
    "Giá cả": "CPI theo nhóm hàng, lạm phát cơ bản, vàng/USD, PPI",
    "Đối ngoại": "tỷ giá, dự trữ, XNK hàng hoá theo mặt hàng, XNK dịch vụ",
    "Tiền tệ & thanh khoản": "lãi suất liên ngân hàng/điều hành, OMO, huy động - cho vay, tín dụng, M2",
    "Hệ thống ngân hàng": "LDR, vốn ngắn hạn cho vay TDH, CAR, quy mô TCTD",
    "Lao động": "thất nghiệp, thiếu việc làm, lực lượng lao động",
    "Thế giới": "lãi suất Mỹ, DXY, vĩ mô 12 nước",
    "Dữ liệu gốc": "toàn bộ chuỗi truyền dẫn, niên giám PX-Web",
}
NHOM_NSO_TX = {       # nhom so lieu thang/quy NSO -> nhom khung phan tich
    "Tăng trưởng & đầu tư": ["GDP", "IIP", "RETAIL", "INVEST", "NSNN", "FDI"],
    "Giá cả": ["CPI", "PPI"],
    "Đối ngoại": ["XK", "NK", "SVC", "TOURIST"],
    "Lao động": ["LABOR", "UNEMP"],
}

# ---------------------------------------------------------------- REGISTRY
# key: (ten hien thi, nhom = "Khoi · Nhom" theo khung phan tich, duong dan, cot ngay, cot thuc the, mo ta)
REGISTRY = {
    "indices": ("Chỉ số thị trường (OHLCV + GTGD)", "Cổ phiếu · Giá & thanh khoản",
                os.path.join(IDXD, "indices-master.csv"), "date", "index_code",
                "25 chỉ số VN + châu Á + Mỹ; có value (triệu VND) và value_usd"),
    "tv_history": ("Giá cổ phiếu TradingView", "Cổ phiếu · Giá & thanh khoản",
                   os.path.join(IDXD, "tv-history.csv"), "date", "tv_symbol",
                   "1.245 mã VN + chỉ số, giá điều chỉnh, value_approx = close×KL"),
    "fx": ("Tỷ giá", "Vĩ mô · Đối ngoại", os.path.join(IDXD, "fx-master.csv"),
           "date", "currency", "9 đồng tiền, dùng để quy USD"),
    "flows": ("Khối ngoại & tự doanh", "Cổ phiếu · Dòng tiền",
              os.path.join(IDXD, "flows-master.csv"), "date", "index_code",
              "VN 3 sàn (ngày) + KOSPI/KOSDAQ/TAIEX/JCI (ngày), SET/Bursa (tháng)"),
    "valuation_wide": ("Định giá thị trường VN", "Cổ phiếu · Định giá",
                       os.path.join(VALD, "valuation-wide.csv"), "date", "code",
                       "P/E, P/B, P/S, cổ tức, vốn hoá, LN TTM cho VNINDEX/HNX/UPCOM/VN30"),
    "valuation_adjusted": ("Định giá loại nhóm Vin", "Cổ phiếu · Định giá",
                           os.path.join(VALD, "valuation-adjusted.csv"), "date", "index",
                           "P/E, P/B, ROE trước và sau khi loại VIC/VHM/VRE/VPL"),
    "sectors_wide": ("Định giá theo ngành ICB", "Cổ phiếu · Định giá",
                     os.path.join(VALD, "sectors-wide.csv"), "date", "code",
                     "55 mã ngành ICB × P/E, P/B, P/S, cổ tức, vốn hoá (VNDirect)"),
    "stocks_wide": ("Định giá từng mã", "Cổ phiếu · Định giá",
                    os.path.join(VALD, "stocks-wide.csv"), "date", "code",
                    "P/E, P/B, vốn hoá, BVPS, EPS, ROE theo mã"),
    "shares": ("Số CP lưu hành & vốn hoá", "Cổ phiếu · Định giá",
               os.path.join(VALD, "shares-master.csv"), "date", "code", "theo mã, theo ngày"),
    "valuation_region_wide": ("Định giá thị trường khu vực", "Khu vực · Định giá",
                              os.path.join(VALD, "valuation-region-wide.csv"), "date", "code",
                              "P/E, P/B, vốn hoá 12 thị trường (TradingView, SET, SSE, JPX, MSCI)"),
    "nso_monthly": ("Số liệu tháng & quý Cục Thống kê (GDP, CPI, XNK, IIP, bán lẻ, FDI, lao động)", "Vĩ mô · Tăng trưởng / Giá cả / Đối ngoại / Lao động",
                    os.path.join(NSOD, "nso_monthly_master.csv"), "date", "series_id",
                    "Biểu số liệu báo cáo KT-XH tháng của NSO từ 01/2023: CPI theo nhóm hàng (so cùng kỳ/tháng trước/"
                    "tháng 12/bình quân), XK-NK theo mặt hàng (trị giá + lượng), IIP theo ngành, tổng mức bán lẻ, FDI đăng ký"),
    "nso": ("Niên giám thống kê NSO (PX-Web, 333 bảng)", "Vĩ mô · Dữ liệu gốc", os.path.join(NSOD, "nso_master.csv"),
            "date", "table_id",
            "GDP, thu chi NSNN, M2/tín dụng/lãi suất, chứng khoán, CPI/PPI, XNK, bán lẻ, IIP, FDI, DN, lao động, "
            "nông nghiệp... theo năm (1986+), CPI theo tháng (1995+); bảng = table_id, chuỗi = dim1|dim2|dim3"),
    "macro_vn": ("Vĩ mô Việt Nam (IMF)", "Vĩ mô · Giá cả / Đối ngoại", os.path.join(MACD, "macro_vn_master.csv"),
                 "date", "series_id", "CPI chỉ số 2024=100 + 12 nhóm COICOP, tỷ giá, XNK, dự trữ ngoại hối, CAR FSI (IMF SDMX)"),
    "macro_region": ("Vĩ mô 12 nước", "Vĩ mô · Thế giới", os.path.join(MACD, "macro_region_master.csv"),
                     "date", "series_id", "cùng bộ chỉ tiêu IMF cho 12 nước châu Á"),
    "transmission": ("Truyền dẫn tỷ giá → lãi suất → thanh khoản", "Vĩ mô · Tiền tệ & thanh khoản",
                     os.path.join(TRAD, "transmission-master.csv"), "date", "series_id",
                     "16 node: VNIBOR, OMO, tín phiếu, tỷ giá, tín dụng, số dư ngân hàng"),
    "vsdc": ("Tài khoản nhà đầu tư (VSDC)", "Cổ phiếu · Dòng tiền",
             os.path.join(VSDC, "vsdc_tk_ndt.csv"), "date", None,
             "số tài khoản cá nhân / tổ chức / nước ngoài theo tháng"),
    "bonds": ("Trái phiếu doanh nghiệp (toàn thị trường)", "Trái phiếu · Phát hành",
              os.path.join(BOND, "data", "processed", "market_issuance_timeline.csv"),
              "ngay_phat_hanh", "parent_group",
              "từng lô TPDN (HNX CBIS): giá trị phát hành & đang lưu hành, lãi suất, "
              "ngày đáo hạn, mua lại, CTCK bảo lãnh + độ tin cậy"),
    "bond_prices": ("Giá giao dịch TPDN riêng lẻ (ngày)", "Trái phiếu · Giao dịch & lợi suất",
                    os.path.join(BOND, "data", "processed", "bond_prices.csv"), "ngay", "ma_gd",
                    "sàn TPDNRL HNX từ 19/07/2023: KL, GTGD, giá cuối, giá BQ (giá gộp, đồng/TP)"),
    "bond_yields": ("Lợi suất đáo hạn từng giao dịch TPDN", "Trái phiếu · Giao dịch & lợi suất",
                    os.path.join(BOND, "data", "processed", "bond_yields.csv"), "ngay", "ma_gd",
                    "YTM hiệu dụng năm từ giá gộp + spread vs TPCP, ngành, độ tin cậy cao/kha/thap"),
    "yield_curve": ("Đường cong lợi suất TPDN theo ngành", "Trái phiếu · Giao dịch & lợi suất",
                    os.path.join(BOND, "data", "processed", "yield_curve.csv"), "ngay_moc", "nganh",
                    "BQ gia quyền giá trị lô theo bucket kỳ hạn × ngành, 5 mốc 0/1/3/6/12M"),
    "tpcp_curve": ("Đường cong spot TPCP (HNX)", "Trái phiếu · Giao dịch & lợi suất",
                   os.path.join(BOND, "data", "processed", "tpcp_curve.csv"), "ngay", "ky_han",
                   "spot rate 11 kỳ hạn 3T–20N theo ngày, benchmark phi rủi ro"),
}


# LUU Y: tham so cua ham @st.cache_data bat dau bang "_" KHONG duoc bam vao khoa cache.
# Truoc day dat ten "_mt" -> file CSV doi ma app van tra ban cu. Tham so mtime phai ten "mt".
def _mtime(path):
    try:
        return os.path.getmtime(path)
    except OSError:
        return 0.0


# ---------------------------------------------------------------- LOADERS
@st.cache_data(show_spinner=False)
def read_csv(path: str, mt: float, **kw) -> pd.DataFrame:
    df = pd.read_csv(path, **kw)
    for c in df.columns:
        if c in ("date", "ngay_phat_hanh", "ngay_dao_han", "ngay_bao_cao", "ngay", "ngay_moc"):
            df[c] = pd.to_datetime(df[c], errors="coerce")
    return df


def load(key: str, **kw) -> pd.DataFrame:
    path = REGISTRY[key][2]
    return read_csv(path, _mtime(path), **kw)


@st.cache_data(show_spinner="Đang nạp giá cổ phiếu (lần đầu sẽ dựng cache parquet) ...")
def load_tv(mt: float) -> pd.DataFrame:
    """tv-history.csv 233 MB -> cache parquet, chi giu cot can dung."""
    pq = os.path.join(CACHE, "tv-history.parquet")
    if os.path.exists(pq) and os.path.getmtime(pq) >= mt:
        return pd.read_parquet(pq)
    df = pd.read_csv(REGISTRY["tv_history"][2],
                     usecols=["date", "tv_symbol", "exchange", "symbol", "close", "volume", "value_approx"],
                     dtype={"date": str})
    df["date"] = pd.to_datetime(df.date)
    for c in ("tv_symbol", "exchange", "symbol"):
        df[c] = df[c].astype("category")
    df.to_parquet(pq, index=False)
    return df


def tv() -> pd.DataFrame:
    return load_tv(_mtime(REGISTRY["tv_history"][2]))


@st.cache_data(show_spinner=False)
def screener_meta(mt: float) -> pd.DataFrame:
    m = pd.read_csv(os.path.join(RAW, "vn_screener_meta.csv")).drop_duplicates("name").set_index("name")
    m["nhom"] = [nhom_nganh(n, s, i) for n, s, i in zip(m.index, m.sector, m.industry)]
    return m


def meta(pn=None) -> pd.DataFrame:
    """Meta TradingView (san, so CP, von hoa) + cot nhom = nganh ICB Vietcap theo pn
    (mac dinh PN hien hanh do app dat o thanh ben). Ma chua co ICB -> giu nhom TradingView."""
    m = screener_meta(_mtime(os.path.join(RAW, "vn_screener_meta.csv")))
    icb_nhom = m.index.to_series().map(nganh_ma(pn or PN))
    return m.assign(nhom_tv=m.nhom, nhom=icb_nhom.fillna(m.nhom))


BROKERS = ("SSI", "VND", "VCI", "HCM", "VIX", "SHS", "MBS", "FTS", "BSI", "CTS",
           "VDS", "AGR", "ORS", "TCX", "VPX")


def nhom_nganh(n, s, i):
    """Phan nhom giong index-fetcher\\reports\\three_month.py va chart-pack."""
    i, s = str(i), str(s)
    if n in ("VIC", "VHM", "VRE", "VPL"):
        return "Vingroup"
    if "Bank" in i and "Investment" not in i:
        return "Ngân hàng"
    if "Investment Banks" in i or "Brokers" in i or n in BROKERS:
        return "Chứng khoán"
    if "Real Estate" in i or "Homebuilding" in i:
        return "BĐS (ngoài Vin)"
    if "Steel" in i:
        return "Thép"
    if s == "Finance":
        return "Tài chính khác"
    if s in ("Consumer Non-Durables", "Consumer Durables", "Retail Trade",
             "Consumer Services", "Distribution Services"):
        return "Tiêu dùng, bán lẻ"
    if s in ("Utilities", "Energy Minerals"):
        return "Điện, dầu khí"
    if s in ("Technology Services", "Electronic Technology", "Communications"):
        return "Công nghệ, viễn thông"
    if s in ("Process Industries", "Non-Energy Minerals", "Producer Manufacturing",
             "Industrial Services"):
        return "Công nghiệp, vật liệu"
    if s == "Transportation":
        return "Vận tải, cảng, HK"
    return "Khác"


# ------------------------------------------------------- MA TRAN GIA / GTGD
@st.cache_data(show_spinner="Đang dựng ma trận giá ...")
def price_matrix(mt: float, since: str = "2016-01-01"):
    """(px_raw, px_ffill, val) cua toan bo co phieu VN - nen tang cho do rong / MA / nganh."""
    d = tv()
    m = meta()
    d = d[(d.exchange.isin(VN_EX)) & (d.date >= since) & (d.symbol.isin(m.index))]
    px_raw = d.pivot_table(index="date", columns="symbol", values="close", aggfunc="last",
                           observed=True).sort_index()
    val = d.pivot_table(index="date", columns="symbol", values="value_approx", aggfunc="last",
                        observed=True).sort_index()
    return px_raw, px_raw.ffill(limit=20), val


def prices(since="2016-01-01"):
    return price_matrix(_mtime(REGISTRY["tv_history"][2]), since)


@st.cache_data(show_spinner="Đang tính độ rộng thị trường ...")
def breadth(mt: float, since: str = "2015-01-01") -> pd.DataFrame:
    """Do rong: tang/giam, A/D line, % tren MA, dinh-day 52 tuan, so ma duoi MA50/200/300."""
    px_raw, px, _ = prices(since="2015-01-01")
    m = meta()
    exch = m.exchange.reindex(px.columns)
    r1 = px_raw.pct_change()
    adv, dec = (r1 > 0).sum(axis=1), (r1 < 0).sum(axis=1)
    out = pd.DataFrame({
        "Số mã tăng": adv, "Số mã giảm": dec,
        "Số mã đứng giá": ((r1 == 0) & px_raw.notna()).sum(axis=1),
        "Số mã giao dịch": px_raw.notna().sum(axis=1),
    })
    out["% mã tăng"] = adv / (adv + dec).replace(0, np.nan) * 100
    out["Tăng - Giảm"] = adv - dec
    out["A/D line (luỹ kế)"] = (adv - dec).cumsum()
    for w in (20, 50, 100, 200, 300):
        ma = px.rolling(w, min_periods=w).mean()
        ok = px.notna() & ma.notna()
        below = (px < ma) & ok
        out[f"% mã trên MA{w}"] = ((px > ma) & ok).sum(axis=1) / ok.sum(axis=1).replace(0, np.nan) * 100
        if w in (50, 200, 300):
            for ex in VN_EX:
                cols = [c for c in px.columns if exch.get(c) == ex]
                out[f"Dưới MA{w} - {ex}"] = below[cols].sum(axis=1)
            out[f"Dưới MA{w} - Toàn TT"] = below.sum(axis=1)
            out[f"Số mã đủ dữ liệu MA{w}"] = ok.sum(axis=1)
            out[f"% dưới MA{w}"] = below.sum(axis=1) / ok.sum(axis=1).replace(0, np.nan) * 100
    hi = px.rolling(250, min_periods=120).max()
    lo = px.rolling(250, min_periods=120).min()
    out["Số mã đỉnh 52 tuần"] = ((px >= hi) & px.notna()).sum(axis=1)
    out["Số mã đáy 52 tuần"] = ((px <= lo) & px.notna()).sum(axis=1)
    out["Đỉnh - Đáy 52T"] = out["Số mã đỉnh 52 tuần"] - out["Số mã đáy 52 tuần"]
    out.index.name = "Ngày"
    return out.loc[since:]


def breadth_df(since="2015-01-01"):
    return breadth(_mtime(REGISTRY["tv_history"][2]), since)


@st.cache_data(show_spinner="Đang tính giá trị giao dịch ...")
def turnover(mt_idx: float, mt_tv: float) -> pd.DataFrame:
    """GTGD thi truong VN theo ngay (ty VND).

    Nguon chinh: indices-master.value (VCI accumulatedValue, trieu VND) cho 3 san.
    Bu cac phien VCI chua cap nhat bang close x KL tung ma (tv-history) - danh dau
    o cot 'Nguon' de biet dong nao la uoc tinh.
    """
    idx = load("indices")
    v = (idx[idx.index_code.isin(["VNINDEX", "HNXINDEX", "UPCOM"])]
         .pivot_table(index="date", columns="index_code", values="value", aggfunc="last") / 1e3)
    v = v.rename(columns={"VNINDEX": "HOSE", "HNXINDEX": "HNX", "UPCOM": "UPCoM"})
    v = v.reindex(columns=["HOSE", "HNX", "UPCoM"])
    _, _, val = prices(since="2016-01-01")
    m = meta()
    exch = m.exchange.reindex(val.columns)
    est = pd.DataFrame({
        ex if ex != "UPCOM" else "UPCoM":
            val[[c for c in val.columns if exch.get(c) == ex]].sum(axis=1) / 1e9
        for ex in VN_EX})
    full = v.reindex(v.index.union(est.index))
    src = pd.Series("VCI (chính thức)", index=full.index)
    miss = full["HOSE"].isna()
    src[miss] = "Ước tính close×KL"
    full = full.fillna(est.reindex(full.index))
    full["Toàn thị trường"] = full[["HOSE", "HNX", "UPCoM"]].sum(axis=1, min_count=1)
    full["MA20 toàn TT"] = full["Toàn thị trường"].rolling(20).mean()
    full["MA50 toàn TT"] = full["Toàn thị trường"].rolling(50).mean()
    vni = idx[idx.index_code == "VNINDEX"].set_index("date").close
    full["VN-Index"] = vni.reindex(full.index).ffill()
    full["Nguồn"] = src.reindex(full.index)
    now = pd.Timestamp.now()
    full["Phiên"] = np.where((full.index.normalize() == now.normalize()) & (now.hour < 15),
                             "Đang giao dịch", "Đã đóng cửa")
    # Dong DO PHIEN nam lai tu hom truoc: indices-master duoc keo GIUA gio giao dich cua chinh
    # ngay do (vd 16/09 luc 11:11 -> GTGD 5.341 ty thay vi 14.845 ty) va chua ai keo lai.
    # Nhan biet chinh xac bang gio sua file: file sua ngay D truoc 15:05 => dong ngay D la do phien.
    sua = pd.Timestamp(datetime.fromtimestamp(mt_idx)) if mt_idx else None
    if sua is not None and (sua.hour, sua.minute) < (15, 5):
        do = (full.index.normalize() == sua.normalize()) & (full["Nguồn"] == "VCI (chính thức)")
        full.loc[do & (full["Phiên"] == "Đã đóng cửa"), "Phiên"] = "Dở phiên (kéo lúc %s)" % sua.strftime("%H:%M")
    full.index.name = "Ngày"
    return full.dropna(subset=["Toàn thị trường"])


def turnover_df():
    return turnover(_mtime(REGISTRY["indices"][2]), _mtime(REGISTRY["tv_history"][2]))


@st.cache_data(show_spinner="Đang tính vốn hoá & chỉ số ngành ...")
def sector_caps(mt: float, since: str = "2018-01-01", end: str = None, pn=None):
    """(chi so nganh rebase 100 tai `since`, von hoa nhom quy mo) trong [since, end]."""
    _, px, _ = prices(since="2016-01-01")
    m = meta(pn)
    sh = m.total_shares_outstanding_fundamental
    exch = m.exchange.reindex(px.columns)
    hose = [c for c in px.columns if exch.get(c) == "HOSE" and pd.notna(sh.get(c))]
    px_h = px.loc[since:(end or px.index.max()), hose]
    mc = px_h * sh.reindex(hose)

    rank = mc.iloc[-1].rank(ascending=False)
    caps = pd.DataFrame(index=mc.index)
    caps["Top 30 (≈VN30)"] = mc.loc[:, (rank <= 30).values].sum(axis=1) / 1e12
    caps["Hạng 31-100 (≈Midcap)"] = mc.loc[:, ((rank > 30) & (rank <= 100)).values].sum(axis=1) / 1e12
    caps["Còn lại (≈Smallcap)"] = mc.loc[:, (rank > 100).values].sum(axis=1) / 1e12
    caps["Toàn HOSE"] = mc.sum(axis=1) / 1e12
    caps.index.name = "Ngày"

    nhom = m.nhom.reindex(hose)
    base = mc.columns[mc.iloc[0].notna() & mc.iloc[-1].notna()]
    sec = pd.DataFrame(index=mc.index)
    for g in sorted(nhom.dropna().unique()):
        cols = [c for c in base if nhom[c] == g]
        if cols:
            s = mc[cols].sum(axis=1)
            sec[g] = s / s.iloc[0] * 100
    t = mc[base].sum(axis=1)
    sec["Toàn HOSE"] = t / t.iloc[0] * 100
    sec.index.name = "Ngày"
    return sec, caps


def sectors(since="2018-01-01", end=None):
    return sector_caps(_mtime(REGISTRY["tv_history"][2]), since, end, PN)


def sector_summary(months: int = 12, end=None) -> pd.DataFrame:
    """Bang hieu suat nganh gia quyen von hoa, tinh den ngay `end` (mac dinh phien cuoi)."""
    _, px, _ = prices(since="2016-01-01")
    m = meta()
    sh = m.total_shares_outstanding_fundamental
    exch = m.exchange.reindex(px.columns)
    hose = [c for c in px.columns if exch.get(c) == "HOSE" and pd.notna(sh.get(c))]
    nhom = m.nhom.reindex(hose)
    if end is not None:
        px = px.loc[:pd.Timestamp(end)]
    END = px.index.max()
    p1 = px.loc[END]

    def p_at(mm=None, since=None):
        ix = px.index[px.index <= (pd.Timestamp(since) if since else END - pd.DateOffset(months=mm))]
        return px.loc[ix.max()] if len(ix) else px.iloc[0]

    def wret(cols, p0):
        s0, s1, w = p0[cols], p1[cols], sh.reindex(cols)
        ok = s0.notna() & s1.notna() & (s0 > 0) & w.notna()
        wt = s0[ok] * w[ok]
        return float(np.average((s1[ok] / s0[ok] - 1) * 100, weights=wt)) if ok.any() and wt.sum() else np.nan

    mc_now = (p1[hose] * sh.reindex(hose))
    rows = []
    for g in sorted(nhom.dropna().unique()) + ["TOÀN HOSE"]:
        cols = hose if g == "TOÀN HOSE" else [c for c in hose if nhom[c] == g]
        r = {"Nhóm ngành": g, "Số mã": len(cols),
             "Vốn hoá (nghìn tỷ)": mc_now[cols].sum() / 1e12,
             "Tỷ trọng vốn hoá (%)": mc_now[cols].sum() / mc_now.sum() * 100}
        for lab, mm in [("1T", 1), ("3T", 3), ("6T", 6), (f"{months}T", months)]:
            r[f"Hiệu suất {lab} (%)"] = wret(cols, p_at(mm=mm))
        r[f"Hiệu suất YTD {END.year} (%)"] = wret(cols, p_at(since=f"{END.year - 1}-12-31"))
        p0 = p_at(mm=months)
        rr = (p1[cols] / p0[cols] - 1).dropna()
        r[f"% mã tăng {months}T"] = (rr > 0).mean() * 100 if len(rr) else np.nan
        rows.append(r)
    return pd.DataFrame(rows).set_index("Nhóm ngành").round(2)


# ------------------------------------------------------------- KHOI NGOAI
MKT_REGION = [("VNINDEX", "Việt Nam (HOSE)", "VND"), ("KOSPI", "Hàn Quốc (KOSPI)", "KRW"),
              ("KOSDAQ", "Hàn Quốc (KOSDAQ)", "KRW"), ("TAIEX", "Đài Loan (TWSE)", "TWD"),
              ("JCI", "Indonesia (IDX)", "IDR"), ("SET", "Thái Lan (SET)", "THB"),
              ("FBMKLCI", "Malaysia (Bursa)", "MYR")]


@st.cache_data(show_spinner=False)
def flows_region(mt_fl: float, mt_fx: float, start=None, end=None) -> pd.DataFrame:
    """Khoi ngoai theo thang, trieu USD, trong khoang [start, end].

    Thang nao co dong freq=M thi dung M (SET/Bursa cong bo theo thang), con lai
    cong don cac dong ngay."""
    fl = load("flows")
    fx = load("fx")
    fxm = fx.assign(ym=fx.date.dt.to_period("M")).groupby(["currency", "ym"]).rate.mean().unstack(0)
    p_end = pd.Period(pd.Timestamp(end), "M") if end is not None else fl.date.max().to_period("M")
    p_start = pd.Period(pd.Timestamp(start), "M") if start is not None else p_end - 12
    ymi = pd.period_range(p_start, p_end, freq="M")
    out = {}
    for code, ten, cur in MKT_REGION:
        d = fl[(fl.index_code == code) & (fl.flow_type == "foreign")].dropna(subset=["net_val"]).copy()
        d["ym"] = d.date.dt.to_period("M")
        s = d[d.freq == "M"].groupby("ym").net_val.sum().reindex(ymi)
        s = s.fillna(d[d.freq != "M"].groupby("ym").net_val.sum().reindex(ymi))
        out[ten] = s / fxm[cur].reindex(ymi).ffill() / 1e6
    df = pd.DataFrame(out, index=ymi)
    df["EM Châu Á (tổng)"] = df.sum(axis=1, min_count=1)
    df.index = pd.Index([f"{p.month:02d}/{p.year}" for p in ymi], name="Tháng")
    return df.round(1)


def flows_region_df(start=None, end=None):
    return flows_region(_mtime(REGISTRY["flows"][2]), _mtime(REGISTRY["fx"][2]), start, end)


@st.cache_data(show_spinner=False)
def heatmap(mt: float, days: int, topn: int, san: tuple, end: str, pn=None) -> pd.DataFrame:
    """Du lieu ban do nhiet: top `topn` ma theo von hoa, bien dong `days` phien."""
    _, px, _ = prices()
    m = meta(pn)
    sh = m.total_shares_outstanding_fundamental
    px = px.loc[:pd.Timestamp(end)]
    cols = [c for c in px.columns if m.exchange.get(c) in san and pd.notna(sh.get(c))]
    if not cols or len(px) < 2:
        return pd.DataFrame()
    p1 = px[cols].iloc[-1]
    p0 = px[cols].iloc[max(0, len(px) - 1 - days)]
    mc = (p1 * sh.reindex(cols)) / 1e9                      # ty VND
    d = pd.DataFrame({"Mã": cols, "Nhóm ngành": m.nhom.reindex(cols).values,
                      "Vốn hoá (tỷ)": mc.values,
                      "Biến động (%)": ((p1 / p0 - 1) * 100).values,
                      "Giá": p1.values})
    return (d.dropna(subset=["Vốn hoá (tỷ)", "Biến động (%)"])
            .nlargest(topn, "Vốn hoá (tỷ)").round(2).reset_index(drop=True))


def heatmap_df(days=1, topn=300, san=("HOSE",), end=None):
    return heatmap(_mtime(REGISTRY["tv_history"][2]), days, topn, tuple(san),
                   end or str(pd.Timestamp.now().date()), PN)


# ------------------------------------------------------------------- VI MO
TM_WIDE = os.path.join(TRAD, "transmission-wide.csv")


@st.cache_data(show_spinner=False)
def tm_meta(mt: float) -> pd.DataFrame:
    """series_id -> ten tieng Viet, don vi, node, ngay cuoi. Lay ban GHI SAU CUNG vi
    SBV doi cach viet ten giua chung (vd 'qua dem' -> 'Qua đêm') nhung series_id giu nguyen."""
    t = load("transmission").sort_values("date")
    m = t.groupby("series_id").agg(node=("node_id", "last"), node_name=("node_name", "last"),
                                   ten=("series_name", "last"), unit=("unit", "last"),
                                   freq=("freq", "last"), last=("date", "max"), n=("date", "size"))
    return m


@st.cache_data(show_spinner=False)
def tm_wide(mt: float) -> pd.DataFrame:
    """transmission-wide.csv: 1 cot = 1 series_id, da gop cac ten viet khac nhau."""
    w = pd.read_csv(TM_WIDE, parse_dates=["date"]).set_index("date").sort_index()
    w.index.name = "Ngày"
    return w


# ten hien thi cho cac chuoi hay dung (nguon SBV viet luc co dau luc khong)
TM_TEN = {
    "ib_on": "LNH qua đêm", "ib_1w": "LNH 1 tuần", "ib_2w": "LNH 2 tuần", "ib_1m": "LNH 1 tháng",
    "ib_3m": "LNH 3 tháng", "ib_6m": "LNH 6 tháng", "ib_9m": "LNH 9 tháng", "ib_1y": "LNH 1 năm",
    "ib_spread_policy": "LNH qua đêm - tái cấp vốn", "ib_curve_1m_on": "Độ dốc 1 tháng - qua đêm",
    "policy_refinance": "Lãi suất tái cấp vốn", "policy_rediscount": "Lãi suất tái chiết khấu",
    "omo_win_7d_rate": "Lãi suất trúng thầu OMO 7 ngày",
    "omo_outstanding": "Repo đang lưu hành (bơm)", "bill_outstanding": "Tín phiếu (âm = hút)",
    "omo_net_outstanding": "Bơm ròng đang lưu hành", "omo_net_daily": "Bơm/hút ròng trong ngày",
    "fx_central": "Tỷ giá trung tâm", "fx_band_ceiling": "Trần biên độ +5%",
    "fx_band_floor": "Sàn biên độ -5%", "fx_vcb_sell": "VCB bán",
    "fx_vcb_transfer": "VCB mua chuyển khoản", "fx_free_sell": "Chợ đen bán",
    "fx_free_buy": "Chợ đen mua", "fx_sbv_sell_ref": "NHNN bán tham chiếu",
    "fx_vcb_sell_vs_ceiling": "VCB bán so với trần biên độ",
    "fx_reserves": "Dự trữ ngoại hối", "fx_import_cover": "Dự trữ theo tháng nhập khẩu",
    "swap_on": "Chênh lãi suất qua đêm VND-USD", "swap_1m": "Chênh 1 tháng VND-USD",
    "swap_3m": "Chênh 3 tháng VND-USD",
    "deposit_1m_avg": "Huy động 1T bình quân", "deposit_6m_avg": "Huy động 6T bình quân",
    "deposit_12m_avg": "Huy động 12T bình quân", "deposit_12m_big4": "Huy động 12T nhóm Big4",
    "deposit_12m_max": "Huy động 12T cao nhất", "deposit_rate_avg_vcb": "Huy động bình quân VCB",
    "lending_rate_avg": "Cho vay bình quân (VCB)",
    "lending_deposit_spread": "Chênh cho vay - huy động",
    "credit_growth_ytd": "Tăng trưởng tín dụng luỹ kế", "credit_outstanding": "Tổng dư nợ tín dụng",
    "m2": "M2", "m2_growth": "Tăng trưởng M2 luỹ kế", "m2_growth_yoy": "M2 so cùng kỳ",
    "deposit_growth_yoy": "Tiền gửi so cùng kỳ",
    "credit_deposit_gap": "Chênh tăng trưởng tín dụng - huy động",
    "cash_ratio_m2": "Tiền mặt trên M2", "ldr_cap": "Trần LDR",
    "sfl_cap": "Trần vốn ngắn hạn cho vay TDH", "sfl_system": "Toàn hệ thống",
    "sfl_soe": "NHTM Nhà nước", "sfl_jsc": "NHTM cổ phần",
    "us_fed_target_upper": "Trần mục tiêu Fed", "us_effr": "Fed funds hiệu lực",
    "us_sofr": "SOFR qua đêm", "us_tbill_3m": "Tín phiếu KB Mỹ 3 tháng",
    "us_ust_2y": "TPCP Mỹ 2 năm", "us_ust_10y": "TPCP Mỹ 10 năm", "us_dxy_broad": "Chỉ số USD",
    "fdi_registered": "FDI đăng ký (triệu USD, luỹ kế)",
    "fdi_realized_vnd": "FDI thực hiện (nghìn tỷ VND)",
    "current_account_bop": "Cán cân vãng lai", "financial_account_bop": "Cán cân tài chính",
    "fdi_net_bop": "FDI ròng", "fii_net_bop": "Đầu tư gián tiếp ròng",
    "bop_overall": "Cán cân tổng thể", "bop_errors": "Lỗi và sai sót",
}


def tm(ids, ten_viet=True) -> pd.DataFrame:
    """Lay cac chuoi truyen dan theo series_id, doi ten cot sang tieng Viet."""
    w = tm_wide(_mtime(TM_WIDE))
    m = tm_meta(_mtime(REGISTRY["transmission"][2]))
    cols = [c for c in ids if c in w.columns]
    d = w[cols].copy()
    if ten_viet:
        d.columns = [TM_TEN.get(c, m.ten.get(c, c)) for c in cols]
    return d.dropna(how="all")


def tm_last(sid):
    """(gia tri cuoi, ngay) cua 1 series_id."""
    w = tm_wide(_mtime(TM_WIDE))
    if sid not in w.columns:
        return np.nan, None
    s = w[sid].dropna()
    return (s.iloc[-1], s.index[-1]) if len(s) else (np.nan, None)


def mv(names=None, group=None, contains=None) -> pd.DataFrame:
    """Vi mo VN (IMF) -> wide theo series_name."""
    m = load("macro_vn")
    if group:
        m = m[m.group == group]
    if contains:
        m = m[m.series_name.str.contains(contains, case=False, na=False)]
    if names:
        m = m[m.series_name.isin(names)]
    d = m.pivot_table(index="date", columns="series_name", values="value", aggfunc="last")
    d.index.name = "Ngày"
    return d


def mv_last(name):
    m = load("macro_vn")
    s = m[m.series_name == name].dropna(subset=["value"]).sort_values("date")
    return (s.value.iloc[-1], s.date.iloc[-1]) if len(s) else (np.nan, None)


# ------------------------------------------------------------- NSO (CUC THONG KE)
# 07/10/2026: so lieu vi mo VN lay tu Cuc Thong ke (nso-fetcher) thay IMF o nhung chuoi NSO co:
#   - nso_monthly_master.csv : Bieu so lieu bao cao KT-XH thang (2023-) - CPI theo nhom (so cung ky / thang truoc /
#     thang 12 / binh quan), XK-NK theo mat hang, IIP theo nganh, ban le, FDI dang ky luy ke
#   - nso_master.csv : PX-Web nien giam (333 bang, nam) + CPI thang 2010-2025 (V11.02/05/06)
# IMF (macro_vn) chi con dung cho: chi so CPI muc 2024=100, ty gia, du tru, CAR FSI va lich su truoc 2023.
_NSO_TONE = {"oá": "óa", "oà": "òa", "oả": "ỏa", "oã": "õa", "oạ": "ọa", "uý": "úy", "uỳ": "ùy", "uỷ": "ủy",
             "uỹ": "ũy", "uỵ": "ụy", "oé": "óe", "oè": "òe", "oẻ": "ỏe", "oẽ": "õe", "oẹ": "ọe"}
_NSO_ALIAS = {"nhaovavatlieuxaydung": "nhaodiennuocchatdotvavlxd", "nhaodiennuocchatdotvavatlieuxaydung": "nhaodiennuocchatdotvavlxd",
              "maymacgiaydepvamunon": "maymacmunongiaydep", "maymacmunonvagiaydep": "maymacmunongiaydep",
              "buuchinhvienthong": "thongtinvatruyenthong", "dodungvadichvukhac": "hanghoavadichvukhac",
              "chisogiatieudung": "cpichung", "chisodolamy": "chisogiadolamy", "tongtrigia": "tongso",
              "tylethatnghiepthanhnientu1524tuoi": "tylethatnghiepcuathanhnientu1524tuoi"}
CPI_PX = {"YOY": "V11.06", "MOM": "V11.02", "VS_DEC": "V11.05"}      # bang PX-Web tuong ung metric thang


def nso_key(name: str) -> str:
    """Khoa so khop ten muc giua PX-Web va bao cao thang (hoá/hóa, dau cau, '(*)', ten doi theo nam)."""
    import re
    import unicodedata
    s = re.sub(r"^\s*trong đó\s*:?\s*", "", str(name).lower())
    for a, b in _NSO_TONE.items():
        s = s.replace(a, b)
    s = unicodedata.normalize("NFD", s)
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn").replace("đ", "d")
    s = re.sub(r"[^a-z0-9]+", "", re.sub(r"\(\*+\)", "", s))
    return _NSO_ALIAS.get(s, s)


@st.cache_data(show_spinner=False)
def nso_monthly_wide(mt: float, group: str, metric: str, section: str = "") -> pd.DataFrame:
    m = load("nso_monthly")
    m = m[(m.group == group) & (m.metric == metric)]
    if section:                                   # XK/NK: "MẶT HÀNG CHỦ YẾU"; FDI: "Phân theo ... địa phương / nước"
        m = m[m.section.fillna("").str.contains(section, case=False, regex=False)]
    d = m.pivot_table(index="date", columns="item", values="value", aggfunc="last")
    d.index.name = "Ngày"
    return d


def nm(group: str, metric: str, items=None, pct: bool = False, section: str = "") -> pd.DataFrame:
    """So lieu thang NSO -> wide (ngay x muc). pct=True: chi so 'so voi' (104,9) -> % (+4,9);
    rieng 'Lam phat co ban' NSO da ghi %, khong tru. section: chi lay dong thuoc muc (vd 'mặt hàng')."""
    d = nso_monthly_wide(_mtime(REGISTRY["nso_monthly"][2]), group, metric, section)
    if items:
        d = d[[c for c in items if c in d.columns]]
    if pct:
        d = d.copy()
        for c in d.columns:
            if nso_key(c) != "lamphatcoban":
                d[c] = d[c] - 100
    return d


def nm_last(series_id: str, pct: bool = False):
    m = load("nso_monthly")
    s = m[m.series_id == series_id].dropna(subset=["value"]).sort_values("date")
    if not len(s):
        return np.nan, None
    v = s.value.iloc[-1]
    return (v - 100 if pct and nso_key(series_id.split("|")[-1]) != "lamphatcoban" else v), s.date.iloc[-1]


@st.cache_data(show_spinner="Đang nạp niên giám NSO ...")
def nso_long(mt: float) -> pd.DataFrame:
    p = REGISTRY["nso"][2]
    pq = os.path.join(CACHE, "nso_master.parquet")
    if os.path.exists(pq) and os.path.getmtime(pq) >= mt:
        return pd.read_parquet(pq)
    df = pd.read_csv(p, dtype={"date": str, "status": str, "dim1": str, "dim2": str, "dim3": str, "unit": str},
                     keep_default_na=False, encoding="utf-8-sig",
                     usecols=["date", "freq", "table_id", "series_id", "pos", "dim1", "dim2", "dim3", "unit", "status", "value"])
    df["value"] = pd.to_numeric(df["value"], errors="coerce")
    lab = df["dim1"].copy()
    for c in ("dim2", "dim3"):
        lab = lab.where(df[c] == "", lab + " | " + df[c])
    df["label"] = lab.where(lab != "", "Chung")      # KHONG dat "Giá trị": line() melt(value_name="Giá trị") se trung ten
    df.to_parquet(pq, index=False)
    return df


def nso_catalog() -> pd.DataFrame:
    p = os.path.join(NSOD, "nso_catalog.csv")
    return read_csv(p, _mtime(p), dtype=str, keep_default_na=False, encoding="utf-8-sig")


@st.cache_data(show_spinner=False)
def nso_table_wide(mt: float, table_id: str, freq: str) -> pd.DataFrame:
    """1 bang PX-Web -> wide (ngay x chuoi). freq A: nam (01/01), M: thang, Q: quy (thang cuoi quy)."""
    d = nso_long(mt)
    d = d[(d.table_id == table_id) & (d.freq == freq)]
    if d.empty:
        return pd.DataFrame()
    if freq == "Q":
        dt = pd.PeriodIndex(d["date"].str.replace("-Q", "Q"), freq="Q").to_timestamp(how="end").normalize()
    else:
        dt = pd.to_datetime(d["date"], errors="coerce")
    d = d.assign(dt=dt).dropna(subset=["dt"])
    w = d.pivot_table(index="dt", columns="label", values="value", aggfunc="last")
    order = d.groupby("label")["pos"].min().sort_values().index
    w = w.reindex(columns=[c for c in order if c in w.columns])
    w.index.name = "Ngày"
    w.attrs["unit"] = d.groupby("label")["unit"].first().reindex(w.columns).fillna("").to_dict()
    return w


def nso_table(table_id: str, freq: str = "A") -> pd.DataFrame:
    return nso_table_wide(_mtime(REGISTRY["nso"][2]), table_id, freq)


def cpi_nso(metric: str = "YOY") -> pd.DataFrame:
    """CPI theo nhom hang, % (da tru 100): PX-Web nien giam (thang, 2010-2025, so da chot) noi voi bao cao thang
    (2023-nay, co Lam phat co ban + Dich vu y te/giao duc). Ky trung nhau uu tien PX-Web."""
    mo = nm("CPI", metric, pct=True)
    if metric not in CPI_PX:
        return mo
    px = nso_table(CPI_PX[metric], "M")
    if px.empty:
        return mo
    key2name = {nso_key(c): c for c in mo.columns}
    px = px - 100
    px.columns = [key2name.get(nso_key(c), c) for c in px.columns]
    px = px.loc[:, ~px.columns.duplicated()]
    out = px.combine_first(mo)
    # thu tu nhu bang NSO: CPI chung, 11 nhom hang (kem nhom con), vang, USD, lam phat co ban
    thu_tu = ["cpichung", "hanganvadichvuanuong", "luongthuc", "thucpham", "anuongngoaigiadinh", "douongvathuocla",
              "maymacmunongiaydep", "nhaodiennuocchatdotvavlxd", "thietbivadodunggiadinh", "thuocvadichvuyte",
              "dichvuyte", "giaothong", "thongtinvatruyenthong", "giaoduc", "dichvugiaoduc", "vanhoagiaitrivadulich",
              "hanghoavadichvukhac", "chisogiavang", "chisogiadolamy", "lamphatcoban"]
    rank = {k: i for i, k in enumerate(thu_tu)}
    cols = sorted(out.columns, key=lambda c: (rank.get(nso_key(c), 99), c))
    out = out[cols]
    out.index.name = "Ngày"
    return out


def _noi(nso: pd.Series, imf: pd.Series) -> pd.Series:
    """NSO ghi de IMF o ky trung; IMF bu lich su truoc khi co NSO."""
    return nso.combine_first(imf) if len(imf) else nso


def trade_nso() -> pd.DataFrame:
    """XK, NK, can can (trieu USD/thang): IMF (ITG) truoc 2023, Cuc Thong ke (Bieu so lieu thang) tu 2023."""
    imf = mv(group="TRADE")
    x_i = imf["Xuat khau hang hoa FOB"] / 1e6 if "Xuat khau hang hoa FOB" in imf else pd.Series(dtype=float)
    n_i = imf["Nhap khau hang hoa CIF"] / 1e6 if "Nhap khau hang hoa CIF" in imf else pd.Series(dtype=float)
    xk, nk = nm("XK", "LEVEL"), nm("NK", "LEVEL")
    x = _noi(xk["Tổng số"] if "Tổng số" in xk else pd.Series(dtype=float), x_i)
    n = _noi(nk["Tổng số"] if "Tổng số" in nk else pd.Series(dtype=float), n_i)
    out = pd.DataFrame({"Xuất khẩu": x, "Nhập khẩu": n})
    out["Cán cân thương mại"] = out["Xuất khẩu"] - out["Nhập khẩu"]
    out.index.name = "Ngày"
    return out.dropna(how="all")


def trade_nso_yoy() -> pd.DataFrame:
    imf = mv(group="TRADE")
    xk, nk = nm("XK", "YOY", pct=True), nm("NK", "YOY", pct=True)
    g = lambda d, c: d[c] if c in d else pd.Series(dtype=float)        # noqa: E731
    out = pd.DataFrame({"Xuất khẩu": _noi(g(xk, "Tổng số"), g(imf, "Xuat khau hang hoa - YoY %")),
                        "Nhập khẩu": _noi(g(nk, "Tổng số"), g(imf, "Nhap khau hang hoa - YoY %"))})
    out.index.name = "Ngày"
    return out.dropna(how="all")


def iip_nso() -> pd.DataFrame:
    """IIP toan nganh: so cung ky (IMF 2009-2022, NSO 2023-) + so thang truoc (NSO)."""
    imf = mv(group="IIP")
    y = nm("IIP", "YOY", pct=True)
    m = nm("IIP", "MOM", pct=True)
    tn = "Toàn ngành công nghiệp"
    out = pd.DataFrame({"So cùng kỳ": _noi(y[tn] if tn in y else pd.Series(dtype=float),
                                         imf["IIP tang truong YoY"] if "IIP tang truong YoY" in imf else pd.Series(dtype=float)),
                        "So tháng trước": m[tn] if tn in m else pd.Series(dtype=float)})
    out.index.name = "Ngày"
    return out.dropna(how="all")


def nq(group: str, metric: str, items=None, pct: bool = False, section: str = "") -> pd.DataFrame:
    """Chuoi QUY tu bao cao quy NSO (metric _Q): index = thang cuoi quy."""
    return nm(group, metric, items, pct, section)


GDP_NGANH = ["Tổng số", "Nông, lâm nghiệp và thủy sản", "Công nghiệp và xây dựng", "Dịch vụ",
             "Thuế sản phẩm trừ trợ cấp sản phẩm"]


def gdp_nso(kind: str = "YOY_Q", pct: bool = True) -> pd.DataFrame:
    """GDP theo quy (Cuc Thong ke, bao cao quy tu Q1/2023): YOY_Q (tang truong, %), LEVEL_HH_Q (gia hien hanh, ty dong),
    LEVEL_SS_Q (gia so sanh). Cot = nganh kinh te (Tong so + 3 khu vuc + thue + nganh cap 1)."""
    d = nm("GDP", kind, pct=pct and kind.startswith("YOY"))
    if d.empty:
        return d
    ten = {nso_key(x): x for x in GDP_NGANH}
    d.columns = [ten.get(nso_key(c), c) for c in d.columns]       # thuy/thuỷ, thue san pham... -> ten chuan
    d = d.loc[:, ~d.columns.duplicated()]
    order = {k: i for i, k in enumerate(GDP_NGANH)}
    cols = sorted(d.columns, key=lambda c: (order.get(c, 99), c))
    return d[cols]


def gdp_last():
    """(tang truong GDP quy gan nhat %, quy, status)"""
    m = load("nso_monthly")
    s = m[m.series_id == "GDP|YOY_Q|Tổng số"].dropna(subset=["value"]).sort_values("date")
    if not len(s):
        return np.nan, None, ""
    r = s.iloc[-1]
    return r.value - 100, r.date, (r.status if isinstance(r.status, str) else "")


def unemp_nso() -> pd.DataFrame:
    """Ty le that nghiep / thieu viec lam theo quy (%): bang UNEMP (dong = quy) noi voi bang lao dong (LABOR) khi
    NSO gop 2 bang (Q3/2026)."""
    u = nm("UNEMP", "RATE_Q")
    u = u[[c for c in u.columns if c.endswith("- Chung")]].copy()
    u.columns = [c.replace(" - Chung", "") for c in u.columns]
    lb = nm("LABOR", "LEVEL_Q")
    u = u.reindex(u.index.union(lb.index))          # quy moi chi co trong bang lao dong (Q3/2026)
    for c in lb.columns:
        if c.startswith("Tỷ lệ") and not re.search(r"-\s*(Nam|Nữ|Thành thị|Nông thôn)$", c):
            key = nso_key(c)
            match = next((x for x in u.columns if nso_key(x) == key), None)
            if match is not None:
                u[match] = u[match].combine_first(lb[c])
            else:
                u[c] = lb[c]
    u.index.name = "Ngày"
    return u.sort_index()


def fdi_nso() -> pd.DataFrame:
    """FDI dang ky luy ke tu dau nam (trieu USD) + so du an, tong so, theo thang bao cao."""
    cols = {"REG_NEW_YTD": "Vốn đăng ký cấp mới (luỹ kế)", "REG_ADJ_YTD": "Vốn điều chỉnh (luỹ kế)",
            "PROJECTS_YTD": "Số dự án cấp mới (luỹ kế)"}
    out = {}
    for k, ten in cols.items():
        d = nm("FDI", k)
        if "Tổng số" in d:
            out[ten] = d["Tổng số"]
    out = pd.DataFrame(out)
    if {"Vốn đăng ký cấp mới (luỹ kế)", "Vốn điều chỉnh (luỹ kế)"} <= set(out.columns):
        out["Cấp mới + điều chỉnh (luỹ kế)"] = out["Vốn đăng ký cấp mới (luỹ kế)"] + out["Vốn điều chỉnh (luỹ kế)"].fillna(0)
    out.index.name = "Ngày"
    return out


# ------------------------------------------------------ TRAI PHIEU DOANH NGHIEP
BOND_MASTER = os.path.join(BOND, "data", "processed", "bond_master.csv")

# Moc doi chieu VBMA (Bao cao thuong nien TTTP 2022, 2024, 2025 + bao cao thang 6/2026).
# Du lieu HNX CBIS cua app CHI GOM PHAT HANH RIENG LE trong nuoc -> phan cong chung phai
# lay tu VBMA. So danh dau * la SUY RA tu % tang giam VBMA cong bo o bao cao nam sau.
VBMA_NAM = pd.DataFrame([
    # nam, rieng le, cong chung, quoc te (trieu USD), LS BQ toan TT, ghi chu
    (2021, 711_315, 30_339, 1_250, np.nan, "* suy ra từ BC 2022 (riêng lẻ −65,1%, công chúng −30%, quốc tế −~50%)"),
    (2022, 248_249, 21_237, 625, np.nan, "BC thường niên 2022"),
    (2023, 309_229, 39_090, 250, 7.9, "* suy ra từ BC 2024 (+40,9% / −15,8% / −40%; LS −0,6đ%)"),
    (2024, 435_704, 32_914, 150, 7.3, "BC thường niên 2024 (riêng lẻ 7,2% · công chúng 8,5%)"),
    (2025, 574_556, 55_354, 875, 7.39, "BC thường niên 2025 (riêng lẻ 7,4% · công chúng 7,2%)"),
    (2026, 188_000, 24_800, np.nan, np.nan, "6 tháng đầu năm (BC tháng 6/2026)"),
], columns=["Năm", "VBMA riêng lẻ", "VBMA công chúng", "VBMA quốc tế (tr.USD)",
            "VBMA LS BQ (%)", "Nguồn VBMA"])
VBMA_NAM["VBMA tổng trong nước"] = VBMA_NAM["VBMA riêng lẻ"] + VBMA_NAM["VBMA công chúng"]
VBMA_NGANH = {2024: {"Ngân hàng": 64.4, "Bất động sản": 21.9},    # ty trong rieng le (%)
              2025: {"Ngân hàng": 64.1, "Bất động sản": 25.6}}

NGANH_TP = ["Ngân hàng", "Bất động sản", "Xây dựng", "Tài chính", "Chứng khoán",
            "Tiêu dùng", "Công nghiệp", "Năng lượng", "Lĩnh vực khác", "Chưa phân loại"]
# Nganh do PIPELINE gan (bond-pivot/scripts/classify_issuers.py -> cot nganh, ma_ck,
# nganh_chi_tiet, nguon_nganh trong market_issuance_timeline.csv). App KHONG tu phan loai
# nua de chi co 1 nguon su that; sua sai thi sua config/nganh_overrides.csv roi chay lai.


@st.cache_data(show_spinner="Đang nạp dữ liệu trái phiếu ...")
def bonds_clean(mt: float, mt_master: float, mt_fx: float) -> pd.DataFrame:
    """Trai phieu toan thi truong da lam sach:
    - SUA LOI: trai phieu USD (tien_te=USD) bi nhan menh gia USD nhu VND -> gia tri gan 0;
      quy doi lai theo ty gia binh quan thang phat hanh, gan loai 'Quốc tế (USD)'.
    - Nganh lay tu pipeline (cot nganh); thieu cot thi de "Chưa phân loại".
    """
    b = load("bonds")
    m = pd.read_csv(BOND_MASTER, usecols=["ma_tp", "tien_te"])
    b = b.merge(m, on="ma_tp", how="left")
    b["ngay_phat_hanh"] = pd.to_datetime(b["ngay_phat_hanh"], errors="coerce")
    b["ngay_dao_han"] = pd.to_datetime(b["ngay_dao_han"], errors="coerce")

    so = lambda s: pd.to_numeric(s.astype(str).str.replace(",", ""), errors="coerce")  # noqa: E731
    usd = b.tien_te.eq("USD")
    fx = load("fx")
    fxm = (fx[fx.currency == "VND"].assign(ym=lambda d: d.date.dt.to_period("M"))
           .groupby("ym").rate.mean())
    ty_gia = b.loc[usd, "ngay_phat_hanh"].dt.to_period("M").map(fxm).astype(float)
    gt_usd = so(b.loc[usd, "menh_gia"]) * so(b.loc[usd, "kl_phat_hanh"])
    b["gia_tri_usd_tr"] = np.nan
    b.loc[usd, "gia_tri_usd_tr"] = gt_usd / 1e6
    b.loc[usd, "gia_tri_ty"] = gt_usd * ty_gia / 1e9
    lh_usd = so(b.loc[usd, "menh_gia"]) * so(b.loc[usd, "kl_con_luu_hanh"]).fillna(0)
    b.loc[usd, "gia_tri_luu_hanh_ty"] = lh_usd * ty_gia / 1e9
    # SUA LOI: VietBank 2020 (5 lo VIETBANK.L/RL.20.xx) nhap lai suat dang diem co ban
    # (520–800 thay vi 5,2–8,0%) -> keo lai suat BQ ngan hang quy 4/2020 len ~40%
    bp = b.lai_suat > 100
    b.loc[bp, "lai_suat"] = b.loc[bp, "lai_suat"] / 100
    b["loai"] = np.where(usd, "Quốc tế (USD)", "Riêng lẻ trong nước")
    if "nganh" not in b:                    # file timeline cu, chua co cot nganh tu pipeline
        b["nganh"] = "Chưa phân loại"
    b["nganh"] = b["nganh"].fillna("Chưa phân loại")
    for c in ("ma_ck", "nganh_chi_tiet", "nguon_nganh"):
        if c not in b:
            b[c] = ""
    b["quy_ph"] = b.ngay_phat_hanh.dt.to_period("Q").astype(str)
    return b


def bonds_df() -> pd.DataFrame:
    return bonds_clean(_mtime(REGISTRY["bonds"][2]), _mtime(BOND_MASTER), _mtime(REGISTRY["fx"][2]))


# ------------------------------------------------- GIA GD + DUONG CONG LOI SUAT
# Logic dung duong cong nam o bond-pivot/scripts/yield_curve.py (1 nguon su that cho
# pipeline tuan lan app); app chi goi lai de nguoi dung chon as-of / cua so / do tin cay.
_BOND_SCRIPTS = os.path.join(BOND, "scripts")


def _yc():
    if _BOND_SCRIPTS not in sys.path:
        sys.path.insert(0, _BOND_SCRIPTS)
    import yield_curve
    return yield_curve


@st.cache_data(show_spinner="Đang nạp lợi suất trái phiếu ...")
def _yields_prepared(mt: float, tin_cay: str) -> pd.DataFrame:
    y = pd.read_csv(REGISTRY["bond_yields"][2], low_memory=False)
    return _yc().prepare_yields(y, tin_cay)


def bond_yields_prepared(tin_cay="kha") -> pd.DataFrame:
    """bond_yields.csv da loc (co YTM, khong outlier, du tin cay) + cot d, w_lo."""
    return _yields_prepared(_mtime(REGISTRY["bond_yields"][2]), tin_cay)


def duong_cong(as_of, window=30, tin_cay="kha", min_bond=3):
    """-> (long, fit, charts, wide) tai as-of va 4 moc nhin lai 1/3/6/12M."""
    yc = _yc()
    y = bond_yields_prepared(tin_cay)
    tp = load("tpcp_curve")
    tp = tp.assign(ngay=tp.ngay.dt.strftime("%Y-%m-%d")) if len(tp) else None
    long, fit, charts = yc.compute_curves(y, as_of, window, min_bond, tp)
    wide = yc.wide_table(long) if len(long) else pd.DataFrame()
    return long, fit, charts, wide


def lai_suat_bq(d: pd.DataFrame, by) -> pd.Series:
    """Lai suat phat hanh binh quan GIA QUYEN theo gia tri phat hanh."""
    d = d[d.lai_suat.notna() & (d.lai_suat > 0) & (d.gia_tri_ty > 0)]
    if d.empty:
        return pd.Series(dtype=float)
    return d.groupby(by).apply(lambda x: np.average(x.lai_suat, weights=x.gia_tri_ty)).astype(float)


# ------------------------------------------------- EPS / LOI NHUAN THI TRUONG
def yoy(s: pd.Series) -> pd.Series:
    """Tang truong so voi CUNG KY nam truoc theo NGAY (khong dung shift(250) uoc chung)."""
    s = s.dropna()
    if s.empty:
        return s
    prev = s.reindex(s.index - pd.DateOffset(years=1), method="ffill")
    prev.index = s.index
    return (s / prev - 1) * 100


def despike(s: pd.Series, tol=0.3) -> pd.Series:
    """Bo diem GAI don le cua nguon (vd VN30 ln_ttm 03/11/2025 vot gap doi roi ve ngay).

    Chi bo khi lech > tol so voi CA diem truoc VA diem sau -> buoc nhay that (doi ro,
    doi cach tinh) van giu nguyen vi no keo dai."""
    s = s.dropna()
    if len(s) < 3:
        return s
    a = (s / s.shift() - 1).abs()
    b = (s / s.shift(-1) - 1).abs()
    return s.mask((a > tol) & (b > tol))


@st.cache_data(show_spinner=False)
def eps_market(mt: float) -> pd.DataFrame:
    """EPS & loi nhuan TTM toan thi truong theo ro (VNDirect tinh san).

    eps_index = EPS quy ve DIEM chi so (chia theo divisor) -> so sanh duoc theo thoi gian;
    ln_ttm = tong loi nhuan sau thue 12 thang cua ro (ty VND sau khi /1e9)."""
    v = load("valuation_wide")
    out = []
    for code, d in v.groupby("code"):
        d = d.set_index("date").sort_index()
        eps = despike(d.eps_index.dropna())
        ln = despike(d.ln_ttm.dropna() / 1e9)
        out.append(pd.DataFrame({"code": code, "EPS TTM (điểm)": eps,
                                 "Tăng trưởng EPS (%)": yoy(eps),
                                 "LN TTM (tỷ)": ln, "Tăng trưởng LN TTM (%)": yoy(ln),
                                 "P/E": d.pe, "P/B": d.pb, "ROE TTM (%)": d.roe_ttm * 100}))
    r = pd.concat(out).reset_index().rename(columns={"index": "date", "date": "Ngày"})
    if "Ngày" not in r:
        r = r.rename(columns={r.columns[0]: "Ngày"})
    return r


def eps_market_df():
    return eps_market(_mtime(REGISTRY["valuation_wide"][2]))


@st.cache_data(show_spinner=False)
def eps_sector(mt: float, cap: int, end: str) -> pd.DataFrame:
    """Loi nhuan TTM + tang truong theo NGANH ICB (cap 2 hoac 3), tinh den ngay `end`."""
    s = load("sectors_wide")
    s = s[(s.cap_icb == cap) & (s.date <= pd.Timestamp(end))]
    rows = []
    for _, d in s.groupby("code"):
        d = d.set_index("date").sort_index()
        ln = despike(d.ln_ttm.dropna() / 1e9).dropna()
        if ln.empty:
            continue
        g = yoy(ln)
        rows.append({"Ngành": d.ten_nganh.iloc[-1],
                     "LN TTM (tỷ)": ln.iloc[-1],
                     "Tăng trưởng LN TTM (%)": g.iloc[-1] if len(g) else np.nan,
                     "P/E": d.pe.dropna().iloc[-1] if d.pe.notna().any() else np.nan,
                     "P/B": d.pb.dropna().iloc[-1] if d.pb.notna().any() else np.nan,
                     "ROE TTM (%)": d.roe_ttm.dropna().iloc[-1] * 100 if d.roe_ttm.notna().any() else np.nan,
                     "Vốn hoá (nghìn tỷ)": d.marketcap.dropna().iloc[-1] / 1e12
                     if d.marketcap.notna().any() else np.nan})
    return (pd.DataFrame(rows).dropna(subset=["Tăng trưởng LN TTM (%)"])
            .sort_values("Tăng trưởng LN TTM (%)", ascending=False)
            .set_index("Ngành").round(1))


def eps_sector_df(cap=2, end=None):
    return eps_sector(_mtime(REGISTRY["sectors_wide"][2]), cap,
                      end or str(pd.Timestamp.now().date()))


FS_GLOB = os.path.join(RAW, "vn_foreign_stocks_20*.csv")
FS_VCI = os.path.join(RAW, "vn_foreign_stocks_vci.parquet")


@st.cache_data(show_spinner=False)
def _foreign_stocks(mt: float) -> pd.DataFrame:
    """Khoi ngoai mua/ban theo tung ma + nhom nganh.

    Nguon chinh Vietcap (index-fetcher\\fetch_foreign_vci.py, tu 2000, tach thoa thuan);
    VNDirect (fetch_foreign_stocks.py, chi co tu 08/2018) chi bu cac phien MOI HON phien cuoi
    cua Vietcap. netMatch = rong khop lenh (bo thoa thuan), NaN o phien lay tu VNDirect."""
    cot = ["code", "date", "floor", "buyVal", "sellVal", "netVal", "netMatch", "totalValue", "owned_pct"]
    parts = []
    if os.path.exists(FS_VCI):
        v = pd.read_parquet(FS_VCI, columns=["code", "date", "buyVal", "sellVal", "netVal",
                                             "buyVal_deal", "sellVal_deal", "totalValue", "owned_pct"])
        v["netMatch"] = v.netVal - (v.buyVal_deal.fillna(0) - v.sellVal_deal.fillna(0))
        parts.append(v.drop(columns=["buyVal_deal", "sellVal_deal"]))
    san_vnd = pd.Series(dtype=object)
    fs = sorted(glob.glob(FS_GLOB))
    if fs:
        n = pd.concat([pd.read_csv(f, usecols=["code", "date", "floor", "buyVal", "sellVal", "netVal"],
                                   dtype={"date": str, "code": str}) for f in fs], ignore_index=True)
        n["date"] = pd.to_datetime(n.date)
        san_vnd = n.sort_values("date").groupby("code").floor.last()
        if parts:
            n = n[n.date > parts[0].date.max()]
        parts.append(n.drop(columns=["floor"]))
    if not parts:
        return pd.DataFrame(columns=cot)
    d = pd.concat(parts, ignore_index=True).drop_duplicates(["code", "date"], keep="first")
    m = meta()
    # san: meta TradingView hien hanh -> san VNDirect gan nhat -> san Vietcap (ma huy niem yet: OTC/OTHER)
    d["floor"] = (d.code.map(m.exchange).fillna(d.code.map(san_vnd))
                  .fillna(d.code.map(icb().get("floor", pd.Series(dtype=str)))).fillna("KHAC"))
    d["floor"] = d.floor.astype("category")
    for c in ("totalValue", "owned_pct"):   # phien bu tu VNDirect khong co 2 cot nay
        if c not in d:
            d[c] = np.nan
    return d[cot]


def foreign_stocks() -> pd.DataFrame:
    fs = glob.glob(FS_GLOB) + ([FS_VCI] if os.path.exists(FS_VCI) else [])
    return _foreign_stocks(max((_mtime(f) for f in fs), default=0.0))


ICB_CSV = os.path.join(RAW, "vn_icb_vci.csv")
VIN = ("VIC", "VHM", "VRE", "VPL")
PN_MAC_DINH = (2, True)   # (cap ICB 1-4, tach rieng nhom Vingroup)
PN = PN_MAC_DINH          # phan nganh dang dung - app dat lai o thanh ben moi lan chay (set_pn)


def set_pn(pn):
    """Dat phan nganh cho TOAN app (app 1 nguoi dung, chay local)."""
    global PN
    PN = tuple(pn)


@st.cache_data(show_spinner=False)
def _icb(mt: float) -> pd.DataFrame:
    """Phan nganh ICB 4 cap tieng Viet cua Vietcap (index-fetcher\fetch_icb_vci.py)."""
    if not os.path.exists(ICB_CSV):
        return pd.DataFrame(columns=["icb1", "icb2", "icb3", "icb4"])
    return pd.read_csv(ICB_CSV, dtype=str).drop_duplicates("code").set_index("code")


def icb() -> pd.DataFrame:
    return _icb(_mtime(ICB_CSV) if os.path.exists(ICB_CSV) else 0.0)


def nganh_ma(pn=None) -> pd.Series:
    """code -> ten nganh theo pn = (cap ICB, tach Vingroup). Thieu cap chi tiet thi lui ve cap tren."""
    cap, vin = pn or PN
    t = icb()
    s = t.get(f"icb{cap}", pd.Series(dtype=str))
    for c in range(cap - 1, 0, -1):
        s = s.fillna(t.get(f"icb{c}", pd.Series(dtype=str)))
    if vin:
        s = s.copy()
        s.loc[s.index.isin(VIN)] = "Vingroup"
    return s


def ds_nganh(pn=None) -> list:
    return sorted(set(nganh_ma(pn).dropna())) + ["Chưa phân ngành"]


def _gan_nganh(d, pn):
    return d.assign(nhom=d.code.map(nganh_ma(pn)).fillna("Chưa phân ngành"))


BA_SAN = {"HOSE", "HNX", "UPCOM"}


def _loc_ngay_san(d, start, end, san):
    """Loc ky + san. Chon du 3 san = khong loc san (giu ca ma da huy niem yet cho so lieu cu)."""
    d = d[(d.date >= pd.Timestamp(start)) & (d.date <= pd.Timestamp(end))]
    return d if BA_SAN <= set(san) else d[d.floor.isin(san)]


def _loc_fs(start, end, san, pn=None):
    d = _loc_ngay_san(foreign_stocks(), start, end, san)
    return _gan_nganh(d, pn)


def flows_sector(start, end, san=("HOSE", "HNX", "UPCOM"), tan_suat="W", chi_tieu="netVal", pn=None):
    """(bang ky x nhom nganh, tong ky theo nhom) - ty VND. tan_suat: D / W / M."""
    d = _loc_fs(start, end, san, pn)
    if d.empty:
        return pd.DataFrame(), pd.DataFrame()
    rule = {"D": "D", "W": "W-FRI", "M": "ME"}[tan_suat]
    piv = (d.pivot_table(index="date", columns="nhom", values=chi_tieu, aggfunc="sum", observed=True)
           .resample(rule).sum(min_count=1).dropna(how="all") / 1e9)
    tong = d.groupby("nhom", observed=True)[["buyVal", "sellVal", "netVal", "netMatch"]].sum() / 1e9
    tong.columns = ["Mua (tỷ)", "Bán (tỷ)", "Ròng (tỷ)", "Ròng khớp lệnh (tỷ)"]
    tong["Số mã có GD"] = d[(d.buyVal > 0) | (d.sellVal > 0)].groupby("nhom", observed=True).code.nunique()
    return piv.round(1), tong.sort_values("Ròng (tỷ)", ascending=False).round(1)


def flows_top_ma(start, end, nhom=None, san=("HOSE", "HNX", "UPCOM"), n=10,
                 chi_tieu="netVal", pn=None) -> pd.DataFrame:
    """Top mua rong / ban rong theo ma trong ky (loc theo nhom nganh neu co), ty VND."""
    d = _loc_fs(start, end, san, pn)
    if nhom:
        d = d[d.nhom.isin(nhom)]
    g = d.groupby("code", observed=True).agg(nhom=("nhom", "last"), rong=(chi_tieu, "sum"))
    g["rong"] = g.rong / 1e9
    g = g[g.rong != 0].sort_values("rong")
    return pd.concat([g.tail(n).iloc[::-1], g.head(n)]).rename(
        columns={"nhom": "Nhóm ngành", "rong": "Ròng (tỷ)"}).round(1)


def flows_nhom_ma(start, end, nhom, san=("HOSE", "HNX", "UPCOM"), chi_tieu="netVal", pn=None):
    """Soi 1 nhom nganh: (bang tung ma trong ky, chuoi ngay x ma cua chi_tieu) - ty VND.

    Bang: mua, ban, rong, rong khop lenh, GTGD, % KN trong GTGD, so phien mua/ban rong,
    so huu NN dau -> cuoi ky. Sap xep theo chi_tieu giam dan (mua rong nhieu nhat o tren)."""
    d = _loc_fs(start, end, san, pn)
    d = d[d.nhom == nhom]
    if d.empty:
        return pd.DataFrame(), pd.DataFrame()
    d = d.sort_values("date")
    g = d.groupby("code", observed=True)
    b = pd.DataFrame({
        "Sàn": g.floor.last().astype(str),
        "Mua (tỷ)": g.buyVal.sum() / 1e9,
        "Bán (tỷ)": g.sellVal.sum() / 1e9,
        "Ròng (tỷ)": g.netVal.sum() / 1e9,
        "Ròng khớp lệnh (tỷ)": g.netMatch.sum(min_count=1) / 1e9,
        "GTGD (tỷ)": g.totalValue.sum(min_count=1) / 1e9,
        "Phiên mua ròng": g.netVal.apply(lambda x: int((x > 0).sum())),
        "Phiên bán ròng": g.netVal.apply(lambda x: int((x < 0).sum())),
        "Sở hữu NN đầu kỳ (%)": g.owned_pct.apply(lambda x: x.dropna().iloc[0] * 100 if x.notna().any() else np.nan),
        "Sở hữu NN cuối kỳ (%)": g.owned_pct.apply(lambda x: x.dropna().iloc[-1] * 100 if x.notna().any() else np.nan),
    })
    b["KN / GTGD (%)"] = (b["Mua (tỷ)"] + b["Bán (tỷ)"]) / 2 / b["GTGD (tỷ)"] * 100
    b["Thay đổi sở hữu (điểm %)"] = b["Sở hữu NN cuối kỳ (%)"] - b["Sở hữu NN đầu kỳ (%)"]
    b = b[(b["Mua (tỷ)"] > 0) | (b["Bán (tỷ)"] > 0)]
    khoa = {"netVal": "Ròng (tỷ)", "netMatch": "Ròng khớp lệnh (tỷ)"}.get(chi_tieu, "Ròng (tỷ)")
    b = b.sort_values(khoa, ascending=False).round(2)
    b.index.name = "Mã"
    ts = (d[d.code.isin(b.index)].pivot_table(index="date", columns="code", values=chi_tieu,
                                              aggfunc="sum", observed=True) / 1e9)
    return b, ts


PROP_CSV = os.path.join(RAW, "vn_prop_stocks.csv")


@st.cache_data(show_spinner=False)
def _prop_stocks(mt: float) -> pd.DataFrame:
    """Tu doanh CTCK theo tung ma (VNDirect, index-fetcher\fetch_prop_stocks.py, tu 17/05/2022).
    Chi co dong cho ma CO giao dich tu doanh trong phien."""
    cot = ["code", "date", "floor", "buyVal", "sellVal", "netVal"]
    if not os.path.exists(PROP_CSV):
        return pd.DataFrame(columns=cot)
    d = pd.read_csv(PROP_CSV, usecols=["code", "date", "floor", "buyingVal", "sellingVal", "netVal"],
                    dtype={"date": str, "code": str}).rename(columns={"buyingVal": "buyVal",
                                                                      "sellingVal": "sellVal"})
    d["date"] = pd.to_datetime(d.date)
    d["floor"] = d.floor.astype("category")
    return d[cot]


def prop_stocks() -> pd.DataFrame:
    return _prop_stocks(_mtime(PROP_CSV) if os.path.exists(PROP_CSV) else 0.0)


def _loc_prop(start, end, san, pn=None):
    d = _loc_ngay_san(prop_stocks(), start, end, san)
    return _gan_nganh(d, pn)


def prop_sector(start, end, san=("HOSE", "HNX", "UPCOM"), pn=None) -> pd.Series:
    """Tu doanh rong ca ky theo nhom nganh, ty VND."""
    d = _loc_prop(start, end, san, pn)
    return (d.groupby("nhom", observed=True).netVal.sum() / 1e9).sort_values(ascending=False).round(1)


def prop_nhom_ma(start, end, nhom, san=("HOSE", "HNX", "UPCOM"), pn=None):
    """Soi tu doanh trong 1 nhom nganh: (bang tung ma, chuoi ngay x ma netVal) - ty VND."""
    d = _loc_prop(start, end, san, pn)
    d = d[d.nhom == nhom].sort_values("date")
    if d.empty:
        return pd.DataFrame(), pd.DataFrame()
    g = d.groupby("code", observed=True)
    b = pd.DataFrame({
        "Sàn": g.floor.last().astype(str),
        "Mua (tỷ)": g.buyVal.sum() / 1e9,
        "Bán (tỷ)": g.sellVal.sum() / 1e9,
        "Ròng (tỷ)": g.netVal.sum() / 1e9,
        "Phiên mua ròng": g.netVal.apply(lambda x: int((x > 0).sum())),
        "Phiên bán ròng": g.netVal.apply(lambda x: int((x < 0).sum())),
    })
    # GTGD toan phien lay tu bo Vietcap (cung ky, cung ma) de tinh ty trong tu doanh
    f = _loc_fs(start, end, san)
    gt = f[f.code.isin(b.index)].groupby("code", observed=True).totalValue.sum(min_count=1) / 1e9
    b["TD / GTGD (%)"] = (b["Mua (tỷ)"] + b["Bán (tỷ)"]) / 2 / gt.reindex(b.index) * 100
    b = b[(b["Mua (tỷ)"] > 0) | (b["Bán (tỷ)"] > 0)].sort_values("Ròng (tỷ)", ascending=False).round(2)
    b.index.name = "Mã"
    ts = d.pivot_table(index="date", columns="code", values="netVal", aggfunc="sum", observed=True) / 1e9
    return b, ts


def dong_tien_nhom(start, end, nhom, san=("HOSE", "HNX", "UPCOM"), ct_kn="netVal", pn=None):
    """Ghep khoi ngoai + tu doanh tung ma trong 1 nhom: (bang, chuoi ngay tong KN+TD)."""
    kn, ts_kn = flows_nhom_ma(start, end, nhom, san, ct_kn, pn)
    td, ts_td = prop_nhom_ma(start, end, nhom, san, pn)
    cot_kn = "Ròng (tỷ)" if ct_kn == "netVal" else "Ròng khớp lệnh (tỷ)"
    b = pd.DataFrame({"Khối ngoại ròng (tỷ)": kn[cot_kn] if len(kn) else pd.Series(dtype=float),
                      "Tự doanh ròng (tỷ)": td["Ròng (tỷ)"] if len(td) else pd.Series(dtype=float)})
    if b.empty:
        return b, pd.DataFrame()
    b = b.fillna(0)
    b["Tổng KN + TD (tỷ)"] = b.sum(axis=1)
    dau = np.sign(b["Khối ngoại ròng (tỷ)"]) * np.sign(b["Tự doanh ròng (tỷ)"])
    b["Chiều"] = np.select([dau > 0, dau < 0], ["Cùng chiều", "Ngược chiều"], "Một phía")
    b = b.sort_values("Tổng KN + TD (tỷ)", ascending=False).round(2)
    b.index.name = "Mã"
    ts = ts_kn.add(ts_td, fill_value=0) if len(ts_td) else ts_kn
    return b, ts


SV_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "cache", "stock_val")
SV_RATIOS = {"PRICE_TO_EARNINGS": "pe", "PRICE_TO_BOOK": "pb", "PRICE_TO_SALES": "ps",
             "DIVIDEND_YIELD": "div_yield", "MARKETCAP": "marketcap", "BVPS_CR": "bvps"}


@st.cache_data(show_spinner="Đang kéo định giá mã này từ VNDirect (lần đầu ~15 giây)...")
def _stock_val_live(code: str, ngay: str) -> pd.DataFrame:
    """Dinh gia 1 ma (VNDirect v4/ratios) cho ma KHONG co trong stocks-wide.csv.
    Cache dia cache/stock_val/<MA>.parquet (long), moi ngay chi keo phan moi."""
    import sys
    if VALD not in sys.path:
        sys.path.insert(0, VALD)
    from fetch_valuation import fetch_pair
    os.makedirs(SV_DIR, exist_ok=True)
    f = os.path.join(SV_DIR, f"{code}.parquet")
    old = pd.read_parquet(f) if os.path.exists(f) else pd.DataFrame(columns=["code", "date", "ratio", "value"])
    rows = []
    try:
        for rc in SV_RATIOS:
            tu = old[old.ratio == rc].date.max() if len(old) else None
            rows += fetch_pair(code, rc, tu if isinstance(tu, str) else None)
    except Exception:   # mat mang -> dung cache cu
        rows = []
    d = pd.concat([x for x in (old, pd.DataFrame(rows)) if len(x)], ignore_index=True) if (len(old) or rows) \
        else old
    if d.empty:
        return pd.DataFrame()
    d["date"] = d.date.astype(str).str[:10]
    d = d.drop_duplicates(["code", "date", "ratio"], keep="last")
    d.to_parquet(f, index=False)
    w = d.pivot_table(index="date", columns="ratio", values="value").rename(columns=SV_RATIOS)
    for c in SV_RATIOS.values():
        if c not in w:
            w[c] = np.nan
    w["ln_ttm"] = w.marketcap / w.pe
    w["roe_ttm"] = w.pb / w.pe
    w.index = pd.to_datetime(w.index)
    w.index.name = "date"
    return w.sort_index()


def stock_val(code: str) -> pd.DataFrame:
    """Dinh gia theo ngay cua 1 ma: stocks-wide.csv neu co, khong thi tu keo VNDirect."""
    sw = load("stocks_wide")
    s = sw[sw.code == code]
    if not s.empty:
        return s.set_index("date").sort_index()
    return _stock_val_live(code, str(pd.Timestamp.now().date()))


def flows_vn(since=None) -> pd.DataFrame:
    fl = load("flows")
    d = fl[fl.index_code.isin(["VNINDEX", "HNXINDEX", "UPCOM"])]
    if since:
        d = d[d.date >= since]
    piv = d.pivot_table(index="date", columns=["flow_type", "index_code"], values="net_val",
                        aggfunc="sum") / 1e9
    ren = {("foreign", "VNINDEX"): "KN ròng HOSE", ("foreign", "HNXINDEX"): "KN ròng HNX",
           ("foreign", "UPCOM"): "KN ròng UPCoM", ("prop", "VNINDEX"): "Tự doanh HOSE",
           ("prop", "HNXINDEX"): "Tự doanh HNX", ("prop", "UPCOM"): "Tự doanh UPCoM"}
    piv.columns = [ren.get(tuple(c), "|".join(map(str, c))) for c in piv.columns]
    piv = piv.reindex(columns=[v for v in ren.values() if v in piv.columns])
    kn = [c for c in piv.columns if c.startswith("KN")]
    piv["KN ròng toàn TT"] = piv[kn].sum(axis=1, min_count=1)
    piv["Luỹ kế KN toàn TT"] = piv["KN ròng toàn TT"].fillna(0).cumsum()
    piv.index.name = "Ngày"
    return piv.round(2)


# ------------------------------------------------------------- TIEN ICH
def last_session(exclude_today=True) -> pd.Timestamp:
    _, px, _ = prices()          # dung chung cache voi cac ham khac
    end = px.index.max()
    now = pd.Timestamp.now()
    if exclude_today and end.normalize() == now.normalize() and now.hour < 15:
        end = px.index[px.index < now.normalize()].max()
    return end


def sheet_name(ten: str, dung_roi=()) -> str:
    """Excel cam \\ / * ? : [ ] trong ten sheet va gioi han 31 ky tu."""
    s = str(ten)
    for ch in "\\/*?:[]":
        s = s.replace(ch, "-")
    s = s.strip()[:31] or "Sheet"
    goc, i = s, 2
    while s in dung_roi:
        hau = f"_{i}"
        s = goc[:31 - len(hau)] + hau
        i += 1
    return s


def to_excel(sheets: dict) -> bytes:
    """dict {ten_sheet: DataFrame} -> bytes xlsx (co freeze pane + do rong cot)."""
    buf = io.BytesIO()
    dung = []
    with pd.ExcelWriter(buf, engine="openpyxl", datetime_format="dd/mm/yyyy") as xw:
        for name, df in sheets.items():
            sn = sheet_name(name, dung)
            dung.append(sn)
            df.to_excel(xw, sheet_name=sn)
    buf.seek(0)
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font
    wb = load_workbook(buf)
    for ws in wb.worksheets:
        ws.freeze_panes = "B2"
        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[1].height = 30
        for col in ws.columns:
            w = max((len(str(c.value)) for c in col[:60] if c.value is not None), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max(11, w + 1), 36)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@st.cache_data(show_spinner=False)
def freshness(stamp: float) -> pd.DataFrame:
    """Bang do tuoi cua moi dataset - cot quan trong nhat cua trang Tong quan."""
    rows = []
    today = pd.Timestamp.now().normalize()
    for key, (ten, nhom, path, dcol, ecol, mota) in REGISTRY.items():
        if not os.path.exists(path):
            rows.append({"Dataset": ten, "Nhóm": nhom, "Trạng thái": "THIẾU FILE",
                         "Dữ liệu đến": None, "Trễ (ngày)": None, "Số dòng": 0,
                         "Cập nhật lúc": None, "File": path, "Mô tả": mota})
            continue
        try:
            if key == "tv_history":
                d = tv()
                last, n = d.date.max(), len(d)
            else:
                d = load(key)
                last, n = pd.to_datetime(d[dcol], errors="coerce").max(), len(d)
        except Exception as e:  # noqa: BLE001
            rows.append({"Dataset": ten, "Nhóm": nhom, "Trạng thái": f"LỖI: {e}",
                         "Dữ liệu đến": None, "Trễ (ngày)": None, "Số dòng": 0,
                         "Cập nhật lúc": None, "File": path, "Mô tả": mota})
            continue
        tre = (today - last.normalize()).days if pd.notna(last) else None
        # du lieu thang/quy tre tu nhien -> nguong rong hon
        gioi_han = 400 if key == "nso" else (95 if nhom.startswith("Vĩ mô") else (45 if key in ("vsdc", "bonds") else 8))
        tt = "OK" if tre is not None and tre <= gioi_han else "CŨ"
        rows.append({"Dataset": ten, "Nhóm": nhom, "Trạng thái": tt,
                     "Dữ liệu đến": last, "Trễ (ngày)": tre, "Số dòng": n,
                     "Cập nhật lúc": datetime.fromtimestamp(_mtime(path)),
                     "File": path, "Mô tả": mota})
    return pd.DataFrame(rows)


def freshness_df():
    return freshness(sum(_mtime(v[2]) for v in REGISTRY.values()))
