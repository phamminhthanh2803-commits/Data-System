# -*- coding: utf-8 -*-
r"""pull_vsdc_accounts.py - SO TAI KHOAN GIAO DICH NHA DAU TU (VSDC) - thay cho copy tay vao sheet 'So TK mo moi'.

Nguon: https://vsdc.vn/vi/tra-cuu-thong-ke/TK_SL_TKGD_NDT?tab=7  (bang "So luong tai khoan giao dich cua NDT", tich luy cuoi thang)
Co che (doc tu js/site.js): trang gan header anti-forgery "__VPToken" = <meta name="__VPToken"> vao moi $.ajax;
   thieu header nay -> HTTP 400 rong. Cookie __VPToken di kem (cung phien GET).
   POST /thongke-tkgd_ndt/search  body JSON {SearchKey:"<nam>|VI", CurrentPage:1, RecordOnPage:50, OrderBy:"", OrderType:""}
   (server co dinh 10 dong/trang, lat CurrentPage toi khi du N ban ghi)
   -> tra HTML fragment <table> 6 cot: Thoi gian | trong nuoc CN | TC | nuoc ngoai CN | TC | Tong. Co du lieu tu 2015.
Chuan hoa ngay: VSDC ghi NGAY BAO CAO (vd 07/05/2026 = so lieu cuoi thang 4) -> quy ve cuoi thang so lieu (ngay <=15 -> thang truoc).
   Khop voi cach file IB&Brokerage dang nhap (07/05/2026 -> 30/04/2026).

Dau ra (cung thu muc):
   vsdc_tk_ndt.csv        long-format master (merge/dedup theo thang), co ngay_bao_cao goc
   vsdc_tk_ndt_table2.csv dung layout Table2 cua sheet 'So TK mo moi' (Date, Ca nhan, To chuc, Ca nhan NN, To chuc NN, Tong,
                          Ca nhan mo them, To chuc mo them, Tang truong, Nam, Quy, Time) -> dan thang vao workbook
   vsdc_tk_ndt.xlsx       ban Excel cua table2 (bang Excel 'Table2')
   (khong chep sang OneDrive: build_presentation.py nganh CK tu doc file table2 -> sheet 'So TK mo moi' + Data_TK trong IB&Brokerage_Nganh.xlsx)
Chay:  python pull_vsdc_accounts.py            (tat ca nam 2015..nay)
       python pull_vsdc_accounts.py --years 2025,2026
"""
from __future__ import annotations

import argparse
import datetime as dt
import html
import json
import os
import re
import sys

import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
from mdlib import log, new_session, safe_to_csv, setup_stdout  # noqa: E402

PAGE = "https://vsdc.vn/vi/tra-cuu-thong-ke/TK_SL_TKGD_NDT?tab=7"
API = "https://vsdc.vn/thongke-tkgd_ndt/search"
OUT_LONG = os.path.join(HERE, "vsdc_tk_ndt.csv")
OUT_T2 = os.path.join(HERE, "vsdc_tk_ndt_table2.csv")
OUT_XLSX = os.path.join(HERE, "vsdc_tk_ndt.xlsx")
COLS = ["ca_nhan", "to_chuc", "ca_nhan_nn", "to_chuc_nn", "tong"]
T2_HDR = ["Date", "Cá nhân", "Tổ chức", "Cá nhân NN", "Tổ chức NN", "Tổng", "Cá nhân mở thêm  (TK)",
          "Tổ chức mở thêm (TK)", "Tăng trưởng số TK theo tháng (phải)", "Năm", "Quý", "Time", "Ngày báo cáo"]

ROW_RE = re.compile(r"<tr>\s*<td[^>]*>\s*(\d{2}/\d{2}/\d{4})\s*</td>(.*?)</tr>", re.S)
CELL_RE = re.compile(r"<td[^>]*>\s*([^<]*?)\s*</td>")
TOK_RE = re.compile(r'<meta name="__VPToken" content="([^"]+)"')
TOTAL_RE = re.compile(r"/\s*(\d+)\s*b")            # "Hien thi: 1 - 8 / 8 ban ghi"


def month_end(report_date: dt.date) -> dt.date:
    """Ngay bao cao -> cuoi thang so lieu. Ngay <= 15: so lieu thang truoc; ngay > 15: chinh thang do."""
    y, m = report_date.year, report_date.month
    if report_date.day <= 15:
        y, m = (y - 1, 12) if m == 1 else (y, m - 1)
    nxt = dt.date(y + 1, 1, 1) if m == 12 else dt.date(y, m + 1, 1)
    return nxt - dt.timedelta(days=1)


def to_int(s: str) -> int | None:
    s = html.unescape(s).replace(".", "").replace(",", "").strip()
    return int(s) if s.lstrip("-").isdigit() else None


def open_session():
    """GET trang de lay cookie + meta token (token va cookie phai cung phien)."""
    for i in range(3):
        s = new_session("browser")
        try:
            r = s.get(PAGE, timeout=60)
        except Exception as e:                                    # noqa: BLE001
            log(f"  ! GET trang loi ({e}) - thu lai {i + 1}/3"); continue
        m = TOK_RE.search(r.text) if r.status_code == 200 else None
        if m:
            return s, m.group(1)
        log(f"  ! GET trang HTTP {r.status_code}, khong thay __VPToken - thu lai {i + 1}/3")
    raise SystemExit("LOI: khong lay duoc token VSDC")


def fetch_year(s, token: str, year: int) -> list[dict]:
    """Server co dinh 10 dong/trang (bo qua RecordOnPage) -> lat trang toi khi du 'x / N ban ghi'."""
    hdr = {"Content-Type": "application/json;charset=utf-8", "__VPToken": token,
           "X-Requested-With": "XMLHttpRequest", "Referer": PAGE, "Accept": "*/*"}
    out, page, total = [], 1, None
    while page <= 10:
        body = {"SearchKey": f"{year}|VI", "CurrentPage": page, "RecordOnPage": 50, "OrderBy": "", "OrderType": ""}
        r = s.post(API, data=json.dumps(body), headers=hdr, timeout=60)
        if r.status_code != 200:
            log(f"  ! {year} trang {page}: HTTP {r.status_code} (body {len(r.text)} ky tu)")
            break
        got = 0
        for d, rest in ROW_RE.findall(r.text):
            nums = [to_int(x) for x in CELL_RE.findall(rest)]
            got += 1
            if len(nums) < 5 or any(v is None for v in nums[:5]):
                log(f"  ! {year}: dong {d} khong doc duoc so: {nums}"); continue
            rd = dt.datetime.strptime(d, "%d/%m/%Y").date()
            rec = {"date": month_end(rd).isoformat(), "ngay_bao_cao": rd.isoformat()}
            rec.update(dict(zip(COLS, nums[:5])))
            out.append(rec)
        tot = TOTAL_RE.search(r.text)
        total = int(tot.group(1)) if tot else total
        if got == 0 or total is None or len(out) >= total:
            break
        page += 1
    if total is not None and len(out) < total:
        log(f"  ! {year}: chi lay duoc {len(out)}/{total} dong")
    return out


def build_table2(df: pd.DataFrame) -> pd.DataFrame:
    d = df.sort_values("date").reset_index(drop=True)
    t = pd.DataFrame({
        "Date": pd.to_datetime(d["date"]), "Cá nhân": d.ca_nhan, "Tổ chức": d.to_chuc,
        "Cá nhân NN": d.ca_nhan_nn, "Tổ chức NN": d.to_chuc_nn, "Tổng": d.tong})
    t["Cá nhân mở thêm  (TK)"] = (t["Cá nhân"] + t["Cá nhân NN"]).diff()
    t["Tổ chức mở thêm (TK)"] = (t["Tổ chức"] + t["Tổ chức NN"]).diff()
    t["Tăng trưởng số TK theo tháng (phải)"] = t["Tổng"].astype(float).pct_change()
    t["Năm"] = t.Date.dt.year
    t["Quý"] = (t.Date.dt.month - 1) // 3 + 1
    t["Time"] = "Q" + t["Quý"].astype(str) + "-" + t["Năm"].astype(str)
    t["Ngày báo cáo"] = pd.to_datetime(d["ngay_bao_cao"])
    return t[T2_HDR]


def write_xlsx(t: pd.DataFrame, path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = "Số TK mở mới"
    ws.append(T2_HDR)
    for row in t.itertuples(index=False):
        ws.append([v.to_pydatetime() if hasattr(v, "to_pydatetime") else (None if pd.isna(v) else v) for v in row])
    n = len(t) + 1
    for r in range(2, n + 1):
        ws.cell(r, 1).number_format = "dd/mm/yyyy"; ws.cell(r, 13).number_format = "dd/mm/yyyy"
        ws.cell(r, 9).number_format = "0.00%"
        for c in range(2, 9):
            ws.cell(r, c).number_format = "#,##0"
    tab = Table(displayName="Table2", ref=f"A1:{get_column_letter(len(T2_HDR))}{n}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)
    ws.column_dimensions["A"].width = 12; ws.column_dimensions["M"].width = 12
    for c in range(2, 13):
        ws.column_dimensions[get_column_letter(c)].width = 14
    try:
        wb.save(path)
    except PermissionError:
        log(f"  X {os.path.basename(path)} dang mo trong Excel - bo qua ban xlsx")


def main() -> int:
    setup_stdout()
    ap = argparse.ArgumentParser()
    ap.add_argument("--years", default="", help="vd 2025,2026 (mac dinh 2015..nam nay)")
    a = ap.parse_args()
    this_year = dt.date.today().year
    years = [int(x) for x in a.years.split(",") if x.strip()] or list(range(2015, this_year + 1))
    log(f"VSDC so TK giao dich NDT: {years[0]}..{years[-1]}")
    s, tok = open_session()
    rows = []
    for y in years:
        got = fetch_year(s, tok, y)
        log(f"  {y}: {len(got)} thang" + (f"  (moi nhat {max(r['date'] for r in got)})" if got else ""))
        rows += got
    if not rows:
        log("LOI: khong lay duoc dong nao"); return 2
    new = pd.DataFrame(rows)
    old = (pd.read_csv(OUT_LONG, dtype={"date": str, "ngay_bao_cao": str}, encoding="utf-8-sig")
           if os.path.exists(OUT_LONG) else pd.DataFrame(columns=new.columns))
    n_old = len(old)
    df = pd.concat([old, new], ignore_index=True).drop_duplicates("date", keep="last").sort_values("date").reset_index(drop=True)
    for c in COLS:
        df[c] = pd.to_numeric(df[c], errors="coerce").astype("Int64")
    # canh bao ngay bao cao giua thang de nguoi dung biet cach quy thang
    odd = new[pd.to_datetime(new.ngay_bao_cao).dt.day <= 15]
    if len(odd):
        log("  Ngay bao cao giua thang -> quy ve thang truoc: " + ", ".join(f"{r.ngay_bao_cao}->{r.date}" for r in odd.itertuples()))
    safe_to_csv(df, OUT_LONG)
    t2 = build_table2(df)
    safe_to_csv(t2, OUT_T2, date_format="%Y-%m-%d")
    write_xlsx(t2, OUT_XLSX)
    last = df.iloc[-1]
    log(f"Da luu {len(df)} thang ({df.date.min()}..{df.date.max()}), them moi {len(df) - n_old}; "
        f"moi nhat {last.date}: tong {int(last.tong):,} TK -> {os.path.basename(OUT_LONG)}, "
        f"{os.path.basename(OUT_T2)}, {os.path.basename(OUT_XLSX)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
