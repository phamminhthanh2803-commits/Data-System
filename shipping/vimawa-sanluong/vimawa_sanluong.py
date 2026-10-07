# -*- coding: utf-8 -*-
"""
vimawa_sanluong.py - Kéo + parse BIỂU THỐNG KÊ KHỐI LƯỢNG HÀNG HÓA THÔNG QUA CẢNG BIỂN (Biểu 28-T / 07-T)
                     của Cục Hàng hải và Đường thủy VN, gộp thành 1 CSV tổng.
    https://vimawa.gov.vn/vi/thong-ke  (phân trang ?page=N, ~17 bài/trang, mới nhất trước)
    Mỗi bài /vi/noi-dung/<slug> đính 1 file, định dạng đổi theo thời kỳ nhưng CÙNG 1 MẪU BIỂU:
        2024 -> nay : .docx (1 bảng)                    2021-2023 : .pdf (text; T1-T6/2021 + T4-T11/2023 là ảnh scan -> OCR)
        2020, 2018  : .xlsx 1 sheet                     2017-6/2018: bảng HTML ngay trong bài (không có file)
        "Biểu 28-20xx BC Bộ": .xlsx nhiều sheet T1..T12 + Năm  (bài 5/2025, 6/2026)
        2015 - 6/2016: bài rỗng trên web (không có dữ liệu).  Bài "Tàu thuyền ra, vào cảng" / "Năng lực..." là biểu khác -> bỏ qua.

Mẫu biểu (đơn vị 1000 tấn / 1000 TEU), cột:
    1 Kế hoạch năm | 2 Lũy kế đến hết tháng trước | 3 ƯỚC thực hiện tháng báo cáo | 4 Lũy kế đến hết tháng báo cáo
    5 Lũy kế cùng kỳ năm trước | 6=4/5 (%) | 7=4/1 (%)
    LƯU Ý: cột 3 chỉ là ƯỚC (báo cáo lập ngày 15, thường = cột 2 / số tháng đã qua). Số thực của tháng m nằm ở
    cột 2 của kỳ m+1  ->  cột thang_thuc_te = luy_ke_thang_truoc(m+1) - luy_ke_thang_truoc(m); tháng 12 = số năm - cột 2 T12.

Kết quả data/:
    vimawa_sanluong.csv      1 dòng = kỳ x chỉ tiêu x đơn vị, mỗi kỳ chỉ giữ 1 phiên bản (file riêng của tháng > sheet Biểu 28)
    vimawa_sanluong_all.csv  mọi phiên bản đã parse (đối chiếu)
    vimawa_index.csv         danh sách bài + file đã tải
raw/  file gốc (<slug>__<tên file>), bảng HTML (<slug>__inline.html), cache OCR (ocr/<file>.json)

Chạy:
    python vimawa_sanluong.py              # dò bài mới (dừng khi gặp trang toàn bài đã có), tải, parse lại toàn bộ, ghi CSV
    python vimawa_sanluong.py --full       # duyệt hết các trang danh sách
    python vimawa_sanluong.py --force      # tải lại file + OCR lại
    python vimawa_sanluong.py --no-fetch   # chỉ parse lại từ raw/
    python vimawa_sanluong.py --ocr off    # bỏ qua PDF scan
"""
import argparse, csv, json, os, re, sys, time, unicodedata
from collections import defaultdict
from datetime import datetime
from urllib.parse import unquote, urljoin

import requests
import lxml.html as LH

ROOT = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(ROOT, 'data'); RAW = os.path.join(ROOT, 'raw'); OCR_DIR = os.path.join(RAW, 'ocr')
MASTER_CSV = os.path.join(DATA, 'vimawa_sanluong.csv'); ALL_CSV = os.path.join(DATA, 'vimawa_sanluong_all.csv')
INDEX_CSV = os.path.join(DATA, 'vimawa_index.csv')
OVERRIDE_CSV = os.path.join(ROOT, 'period_overrides.csv')     # slug,period: bài không ghi năm / ghi sai kỳ
BASE = 'https://vimawa.gov.vn'; LIST_URL = BASE + '/vi/thong-ke'
HEADERS = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/128 vimawa-sanluong', 'Accept-Language': 'vi'}
TESSDATA = os.path.join(os.path.normpath(os.environ.get("BCTC_ROOT", "D:/bctc")), 'pdf-detector', 'markitdown-tool', 'tessdata')   # vie/eng traineddata đi kèm tool markitdown
_TESS_WIN = os.path.normpath('C:/Program Files/Tesseract-OCR/tesseract.exe')
TESSERACT = os.environ.get('TESSERACT_CMD') or (_TESS_WIN if os.path.exists(_TESS_WIN) else 'tesseract')

INDEX_FIELDS = ['slug', 'title', 'url', 'kind', 'files', 'file_urls', 'first_seen']
METRICS = ['ke_hoach_nam', 'luy_ke_thang_truoc', 'thang_bao_cao', 'luy_ke', 'luy_ke_cung_ky', 'yoy_pct', 'pct_ke_hoach']
FIELDS = ['period', 'ky_loai', 'nam', 'thang', 'ngay_bao_cao', 'nhom', 'chieu', 'chi_tieu', 'don_vi'] + METRICS + \
         ['thang_thuc_te', 'nhan_goc', 'nguon_loai', 'nguon_file', 'sheet', 'url', 'ghi_chu', 'canh_bao']
ALL_FIELDS = FIELDS + ['chon']
STD_MONTH = METRICS                                                           # thứ tự cột chuẩn của biểu tháng
STD_YEAR = ['ke_hoach_nam', 'luy_ke', 'luy_ke_cung_ky', 'yoy_pct', 'pct_ke_hoach']   # biểu năm

NHOM_ORDER = ['tong', 'container', 'hang_long', 'hang_kho', 'qua_canh']
CHIEU_ORDER = ['tong', 'xuat_khau', 'nhap_khau', 'noi_dia', 'qua_canh_boc_do']
NHOM_NAME = {'tong': 'Tổng số', 'container': 'Container', 'hang_long': 'Hàng lỏng', 'hang_kho': 'Hàng khô',
             'qua_canh': 'Hàng quá cảnh'}
CHIEU_NAME = {'xuat_khau': 'xuất khẩu', 'nhap_khau': 'nhập khẩu', 'noi_dia': 'nội địa', 'qua_canh_boc_do': 'quá cảnh bốc dỡ'}


def log(msg):
    print(f"[{datetime.now():%H:%M:%S}] {msg}", flush=True)


# ---------------------------------------------------------------- tiện ích chuỗi / số
def nfc(v):
    return ' '.join(unicodedata.normalize('NFC', str(v)).split()) if v is not None else ''


def fold(v):
    """bỏ dấu + thường hoá để so khớp nhãn (chịu được lỗi OCR dấu, 'Nhâp khẩu'...)"""
    s = unicodedata.normalize('NFD', nfc(v))
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn').replace('đ', 'd').replace('Đ', 'D')
    return ' '.join(s.lower().split())


def num(v, pct=False):
    """'1.300.382' / '784,826' (cả . và , đều là phân cách nghìn), số xlsx, '118%' -> float; rỗng/'-' -> None"""
    if v is None or isinstance(v, bool): return None
    if isinstance(v, (int, float)):
        x = float(v)
        if pct: x = x * 100 if abs(x) < 20 else x            # xlsx lưu 1.07 = 107%
        return round(x, 3 if not pct else 1)
    s = nfc(v).replace(' ', '')
    if not s or s in '-–—' or s.startswith('#'): return None
    had_pct = '%' in s
    s = s.replace('%', '')
    neg = s.startswith('-'); s = s.lstrip('-')
    if not re.fullmatch(r'[\d.,]+', s) or not re.search(r'\d', s): return None
    if re.fullmatch(r'\d{1,3}([.,]\d{3})+', s): x = float(re.sub(r'[.,]', '', s))
    elif re.fullmatch(r'\d{1,3}([.,]\d{3})+[.,]\d{1,2}', s):                 # 1.234.567,8
        x = float(re.sub(r'[.,]', '', s[:s.rfind(s[-2] if s[-2] in '.,' else s[-3])]) + '.' + re.split(r'[.,]', s)[-1])
    elif re.fullmatch(r'\d+[.,]\d+', s): x = float(s.replace(',', '.'))
    elif s.isdigit(): x = float(s)
    else: return None
    if neg: x = -x
    if pct: return round(x * 100 if (not had_pct and abs(x) < 20) else x, 1)
    return round(x, 3)


def fmt(x):
    if x is None or x == '': return ''
    if isinstance(x, float): return str(int(x)) if x == int(x) else f"{x:.3f}".rstrip('0').rstrip('.')
    return str(x)


def read_csv(path):
    if not os.path.exists(path): return []
    with open(path, encoding='utf-8-sig', newline='') as f: return list(csv.DictReader(f))


def write_csv(path, rows, fields):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    tmp = path + '.tmp'
    with open(tmp, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction='ignore'); w.writeheader()
        for r in rows: w.writerow({k: fmt(r.get(k)) for k in fields})
    os.replace(tmp, path)


# ---------------------------------------------------------------- kỳ báo cáo
def find_period(lines, fallback=''):
    """-> (ky_loai, nam, thang). Ưu tiên dòng tiêu đề trong văn bản ('Tháng 8/2026', 'Năm 2023'), sau đó tiêu đề bài."""
    pats = [(r'thang\s*(\d{1,2})\s*(?:[/.\-]|nam|\s)\s*(?:nam\s*)?((?:19|20)\d{2})', 'thang'),
            (r'(\d{1,2})\s*thang\s*(?:dau\s*)?nam\s*((?:19|20)\d{2})', 'thang'),
            (r'\bt\s*\.?\s*(\d{1,2})\s*[/.\-\s]\s*((?:19|20)\d{2})', 'thang'),
            (r'()\bnam\s*((?:19|20)\d{2})', 'nam')]
    folded = [fold(x) for x in lines if x]
    for full in (True, False):
        for t in folded:
            if 'ngay bao cao' in t or len(t) > 60 and full: continue
            for rx, kind in pats:
                m = re.fullmatch(rx, t) if full else re.search(rx, t)
                if m:
                    mm = int(m.group(1)) if m.group(1) else None
                    if kind == 'thang' and not (1 <= mm <= 12): continue
                    return kind, int(m.group(2)), mm
    t = fold(fallback)
    for rx, kind in pats:
        m = re.search(rx, t)
        if m and (kind == 'nam' or 1 <= int(m.group(1)) <= 12):
            return kind, int(m.group(2)), int(m.group(1)) if m.group(1) else None
    return None, None, None


def find_report_date(lines):
    for t in lines:
        m = re.search(r'ngay bao cao\s*:?\s*(\d{1,2})\s*/\s*(\d{1,2})\s*/\s*(\d{4})', fold(t))
        if m:
            try: return f"{int(m.group(3)):04d}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
            except ValueError: pass
    return ''


# ---------------------------------------------------------------- parse 1 lưới (grid) -> bản ghi
def header_key(text):
    t = fold(text)
    if 'so sanh' in t: return 'pct_ke_hoach' if 'ke hoach' in t else ('yoy_pct' if 'cung ky' in t else None)
    if 'thuc hien nam bao cao' in t: return 'luy_ke'
    if 'thuc hien nam truoc' in t: return 'luy_ke_cung_ky'
    if 'het thang truoc' in t: return 'luy_ke_thang_truoc'
    if 'cung ky' in t: return 'luy_ke_cung_ky'
    if 'luy ke' in t or 'het thang bao cao' in t: return 'luy_ke'
    if 'uoc thuc hien' in t or 'thang bao cao' in t: return 'thang_bao_cao'
    if 'ke hoach nam' in t: return 'ke_hoach_nam'
    return None


def unit_of(text):
    t = fold(text)
    if 'teu' in t: return '1000 TEU'
    if 'tan' in t or '1000' in t: return '1000 tấn'
    return ''


def parse_grid(grid, ctx_lines=(), title='', force_period=None, prefer_title=False):
    """grid: list[list[cell]] (cell = str | số | None). -> dict(ky_loai, nam, thang, ngay_bao_cao, ghi_chu, rows[]) | None"""
    grid = [list(r) for r in grid if r is not None]
    pos = None
    for ri, row in enumerate(grid):
        for ci, c in enumerate(row):
            if fold(c) in ('tong so', 'tong cong'): pos = (ri, ci); break
        if pos: break
    if not pos: return None
    r0, lc = pos
    uc = lc + 1
    ncol = max(len(r) for r in grid)
    h0 = next((i for i in range(r0) if any('danh muc' in fold(c) for c in grid[i])), None)
    hdr = {}
    if h0 is not None:
        for ci in range(uc + 1, ncol):
            txt = ' '.join(dict.fromkeys(nfc(grid[i][ci]) for i in range(h0, r0) if ci < len(grid[i]) and grid[i][ci] is not None))
            k = header_key(txt)
            if k and k not in hdr.values(): hdr[ci] = k
    pre = [' '.join(dict.fromkeys(nfc(c) for c in grid[i] if nfc(c))) for i in range(h0 if h0 is not None else r0)]
    ky, nam, thang = find_period(pre + list(ctx_lines), title)
    if prefer_title and find_period([], title)[1]: ky, nam, thang = find_period([], title)   # OCR: số trong tiêu đề scan dễ sai
    if force_period:
        m = re.fullmatch(r'(\d{4})(?:-(\d{1,2}))?', force_period.strip())
        if m: nam, thang = int(m.group(1)), int(m.group(2)) if m.group(2) else None; ky = 'thang' if thang else 'nam'
    seq = [hdr[c] for c in sorted(hdr)]
    if seq in (STD_MONTH, STD_YEAR): colmap, std = hdr, seq
    else:                                                     # header vỡ / OCR -> theo vị trí chuẩn, đếm số cột số của dòng Tổng số
        nval = sum(1 for c in grid[r0][uc + 1:] if num(c) is not None)
        std = STD_YEAR if nval <= 5 else STD_MONTH
        colmap = {uc + 1 + i: k for i, k in enumerate(std)}
    if std == STD_YEAR: ky, thang = 'nam', None   # biểu năm (kể cả khi tiêu đề ghi "Tháng 12/20xx")
    if not nam: return None

    out, note, cur, last = [], '', 'tong', None
    for row in grid[r0:]:
        cells = row + [None] * (ncol - len(row))
        label = nfc(cells[lc]); f = fold(label)
        joined = fold(' '.join(nfc(c) for c in cells if nfc(c)))
        if joined.startswith('ghi chu'):
            note = ' '.join(dict.fromkeys(nfc(c) for c in cells if nfc(c))); continue
        if 'chia ra' in f: continue
        unit = unit_of(cells[uc]) if uc < len(cells) else ''
        vals = {k: num(cells[ci], pct=k.endswith('pct') or k.startswith('pct')) for ci, k in colmap.items() if ci < len(cells)}
        if not any(v is not None for k, v in vals.items() if 'pct' not in k):
            if f and 'container' in f: cur = 'container'; last = ('container', 'tong', label)
            continue
        if not f:
            if not (unit and last): continue
            nhom, chieu, label = last
        elif 'tong so' in f or 'tong cong' in f: nhom, chieu = 'tong', 'tong'
        elif 'container' in f: cur = 'container'; nhom, chieu = cur, 'tong'
        elif 'hang long' in f: cur = 'hang_long'; nhom, chieu = cur, 'tong'
        elif 'hang kho' in f: cur = 'hang_kho'; nhom, chieu = cur, 'tong'
        elif 'qua canh' in f:
            if cur == 'tong': nhom, chieu = 'tong', 'qua_canh_boc_do'
            else: cur = 'qua_canh'; nhom, chieu = cur, 'tong'
        elif 'xuat khau' in f: nhom, chieu = cur, 'xuat_khau'
        elif 'nhap khau' in f: nhom, chieu = cur, 'nhap_khau'
        elif 'noi dia' in f: nhom, chieu = cur, 'noi_dia'
        else: continue
        last = (nhom, chieu, label)
        name = NHOM_NAME[nhom] if chieu == 'tong' else (
            f"Hàng {CHIEU_NAME[chieu]}" if nhom == 'tong' else f"{NHOM_NAME[nhom]} - {CHIEU_NAME[chieu]}")
        rec = {'nhom': nhom, 'chieu': chieu, 'chi_tieu': name, 'don_vi': unit or '1000 tấn', 'nhan_goc': label}
        rec.update({k: vals.get(k) for k in METRICS})
        out.append(rec)
    if not out: return None
    seen, uniq = set(), []
    for r in out:                                            # nhãn lặp (OCR đọc 2 lần...) -> giữ dòng đầu
        k = (r['nhom'], r['chieu'], r['don_vi'])
        if k in seen: continue
        seen.add(k); uniq.append(r)
    return {'ky_loai': ky or ('thang' if thang else 'nam'), 'nam': nam, 'thang': thang,
            'ngay_bao_cao': find_report_date(pre + list(ctx_lines)), 'ghi_chu': note, 'rows': uniq}


def validate(rows, tol=0.015):
    """kiểm tra số học trong biểu -> ghi canh_bao vào dòng lệch. Trả về số cảnh báo."""
    idx = {(r['nhom'], r['chieu'], r['don_vi']): r for r in rows}
    n = 0

    def warn(r, msg):
        nonlocal n
        r['canh_bao'] = (r.get('canh_bao', '') + '; ' + msg).strip('; '); n += 1

    def close(a, b): return abs(a - b) <= max(3.0, tol * max(abs(a), abs(b)))
    for r in rows:
        a, b, c = r.get('luy_ke_thang_truoc'), r.get('thang_bao_cao'), r.get('luy_ke')
        if None not in (a, b, c) and not close(a + b, c): warn(r, f"luy_ke != thang_truoc + thang ({fmt(a + b)} vs {fmt(c)})")
    for nhom in NHOM_ORDER:
        for unit in ('1000 tấn', '1000 TEU'):
            tot = idx.get((nhom, 'tong', unit))
            if not tot: continue
            for k in ('luy_ke_thang_truoc', 'thang_bao_cao', 'luy_ke', 'luy_ke_cung_ky'):
                parts = [idx[(nhom, ch, unit)].get(k) for ch in CHIEU_ORDER[1:] if (nhom, ch, unit) in idx]
                if len(parts) < 3 or None in parts or tot.get(k) is None: continue
                if not close(sum(parts), tot[k]): warn(tot, f"{k} != tong cac chieu ({fmt(sum(parts))} vs {fmt(tot[k])})")
    return n


# ---------------------------------------------------------------- trích lưới theo định dạng
def grids_docx(path):
    import docx
    d = docx.Document(path)
    paras = [nfc(p.text) for p in d.paragraphs if p.text.strip()]
    for tb in d.tables:
        yield '', [[c.text for c in row.cells] for row in tb.rows], paras


def grids_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        if ws.sheet_state != 'visible' or ws.max_row < 5: continue
        yield ws.title, [list(r) for r in ws.iter_rows(values_only=True)], []


def table_to_grid(tb):
    """bảng HTML -> lưới, trải colspan/rowspan"""
    grid, span = [], {}
    for tr in tb.xpath('.//tr'):
        row, ci = [], 0
        cells = tr.xpath('./td|./th')
        it = iter(cells)
        cell = next(it, None)
        while cell is not None or ci in span:
            if ci in span:
                txt, left = span[ci]
                row.append(txt)
                if left <= 1: del span[ci]
                else: span[ci] = (txt, left - 1)
                ci += 1; continue
            txt = nfc(cell.text_content())
            cs = int(re.sub(r'\D', '', cell.get('colspan') or '') or 1); rs = int(re.sub(r'\D', '', cell.get('rowspan') or '') or 1)
            for k in range(cs):
                row.append(txt)
                if rs > 1: span[ci] = (txt, rs - 1)
                ci += 1
            cell = next(it, None)
        grid.append(row)
    return grid


def grids_html(path):
    with open(path, 'rb') as f: doc = LH.fromstring(f.read().decode('utf-8', 'replace'))
    for tb in doc.xpath('//table'):
        yield '', table_to_grid(tb), []


def words_to_grid(words, bounds):
    """words: [(x0, y0, x1, y1, text)], bounds: biên cột (x tăng dần) -> lưới theo dòng (gom theo y) x cột"""
    if not words: return []
    hs = sorted(w[3] - w[1] for w in words); h = hs[len(hs) // 2] or 10
    rows, cur, cy = [], [], None
    for w in sorted(words, key=lambda w: (w[1] + w[3]) / 2):
        y = (w[1] + w[3]) / 2
        if cy is not None and y - cy > 0.6 * h: rows.append(cur); cur = []
        cur.append(w); cy = y if not cur[:-1] else (cy * (len(cur) - 1) + y) / len(cur)
    if cur: rows.append(cur)
    grid = []
    for r in rows:
        cells = [[] for _ in range(len(bounds) - 1)]
        for w in sorted(r, key=lambda w: w[0]):
            xc = (w[0] + w[2]) / 2
            for i in range(len(bounds) - 1):
                if bounds[i] <= xc < bounds[i + 1]: cells[i].append(w[4]); break
        grid.append([' '.join(c) for c in cells])
    return grid


def grids_pdf(path, ocr='auto', force=False):
    """-> (loại, sheet, grid, ctx). 3 tầng: bảng pdfplumber -> chữ lớp text + biên cột từ đường kẻ -> OCR Tesseract (cache raw/ocr)."""
    import pdfplumber, pymupdf
    with pdfplumber.open(path) as pdf:
        for pi, page in enumerate(pdf.pages):
            text = unicodedata.normalize('NFC', page.extract_text() or '')
            if 'tong so' not in fold(text): continue
            for tb in page.extract_tables():
                if any(fold(c) == 'tong so' for r in tb for c in r if c):
                    yield 'pdf', f'p{pi + 1}', tb, text.splitlines(); return
    doc = pymupdf.open(path)
    for pi, page in enumerate(doc):                           # có lớp text nhưng pdfplumber không dựng được bảng
        if 'tong so' not in fold(page.get_text()): continue
        r = page_grid(page, text_layer=True)
        if r and any('container' in fold(c) for row in r[0] for c in row):
            yield 'pdf', f'p{pi + 1}', r[0], r[1]; return
    if ocr == 'off': return
    res = ocr_pdf(path, force)
    if res: yield 'pdf-ocr', f"p{res['page']}", res['grid'], res['lines']


# ---------------------------------------------------------------- OCR PDF scan
def _long_runs(dark, min_len):
    """mask các đoạn điểm tối liên tiếp theo chiều ngang dài >= min_len (đường kẻ bảng, gạch chân)"""
    import numpy as np
    mask = np.zeros_like(dark)
    for i in np.flatnonzero(dark.sum(axis=1) >= min_len):
        d = np.diff(np.concatenate(([0], dark[i].astype(np.int8), [0])))
        for s, e in zip(np.flatnonzero(d == 1), np.flatnonzero(d == -1)):
            if e - s >= min_len: mask[i, s:e] = True
    return mask


def _cluster(xs, gap):
    out, cur = [], []
    for x in sorted(xs):
        if cur and x - cur[-1] > gap: out.append(cur); cur = []
        cur.append(x)
    if cur: out.append(cur)
    return [sum(c) / len(c) for c in out]


def page_grid(page, dpi=300, text_layer=False):
    """1 trang PDF -> (grid, lines) | None. Dò đường kẻ bảng trên ảnh render (đường dọc = biên cột), xoá đường kẻ,
    rồi lấy chữ từ lớp text (text_layer=True) hoặc OCR Tesseract 1 lượt; xếp chữ vào lưới dòng x cột."""
    import numpy as np
    pix = page.get_pixmap(dpi=dpi, colorspace='gray')
    img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width).copy()
    dark = img < 160
    hmask = _long_runs(dark, int(dpi * 0.5))
    vmask = _long_runs(dark.T, int(dpi * 0.3)).T
    colsum = vmask.sum(axis=0)
    if colsum.max() < dpi: return None                                   # không có bảng kẻ
    ys = np.flatnonzero(vmask.sum(axis=1) > 0)
    bounds = _cluster(np.flatnonzero(colsum >= 0.45 * colsum.max()), dpi * 0.05)
    if len(bounds) < 9: return None
    if text_layer:
        k = dpi / 72
        words = [(w[0] * k, w[1] * k, w[2] * k, w[3] * k, w[4]) for w in page.get_text('words')]
    else:
        import pytesseract
        from PIL import Image
        mask = hmask | vmask
        for sh in (1, 2):                                                # nới mask cho sạch viền đường kẻ
            mask[sh:, :] |= mask[:-sh, :].copy(); mask[:-sh, :] |= mask[sh:, :].copy()
            mask[:, sh:] |= mask[:, :-sh].copy(); mask[:, :-sh] |= mask[:, sh:].copy()
        img[mask] = 255
        d = pytesseract.image_to_data(Image.fromarray(img), lang='vie', config='--psm 6', output_type=pytesseract.Output.DICT)
        words = []
        for i, t in enumerate(d['text']):
            t = t.strip().strip('|[]{}_')
            if not t or float(d['conf'][i]) < 0: continue
            words.append((d['left'][i], d['top'][i], d['left'][i] + d['width'][i], d['top'][i] + d['height'][i], t))
    top = ys.min()
    above = [w for w in words if w[3] < top]                             # chữ phía trên bảng: tiêu đề, kỳ, ngày báo cáo
    lines = [' '.join(c for c in r if c) for r in words_to_grid(above, [0, pix.width])]
    inside = [w for w in words if w[3] >= top and bounds[0] - 5 <= (w[0] + w[2]) / 2 <= bounds[-1] + 5]
    grid = words_to_grid(inside, bounds)
    lc = next((ci for r in grid for ci, c in enumerate(r) if fold(c) == 'tong so'), None)
    if lc is None: return None
    if not text_layer:
        for r in grid:                                                   # làm sạch ô số OCR
            for ci in range(lc + 2, len(r)):
                v = r[ci].replace(' ', '').translate(str.maketrans('OoQlI|SsB', '000111558'))
                r[ci] = v if re.fullmatch(r'[\d.,]*%?', v) else re.sub(r'[^\d.,%]', '', v)
    return grid, lines


TEMPLATE = [('Tổng số', '1000 tấn'), ('Hàng xuất khẩu', '1000 tấn'), ('Hàng nhập khẩu', '1000 tấn'), ('Hàng nội địa', '1000 tấn'),
            ('Hàng quá cảnh', '1000 tấn'), ('Chia ra', ''), ('Container', '1000 tấn'), ('', '1000 Teus'),
            ('Xuất khẩu', '1000 tấn'), ('', '1000 Teus'), ('Nhập khẩu', '1000 tấn'), ('', '1000 Teus'),
            ('Nội địa', '1000 tấn'), ('', '1000 Teus'), ('Hàng lỏng', '1000 tấn'), ('Xuất khẩu', '1000 tấn'),
            ('Nhập khẩu', '1000 tấn'), ('Nội địa', '1000 tấn'), ('Hàng khô', '1000 tấn'), ('Xuất khẩu', '1000 tấn'),
            ('Nhập khẩu', '1000 tấn'), ('Nội địa', '1000 tấn'), ('Hàng quá cảnh', '1000 tấn')]   # thứ tự dòng cố định của biểu


# nhóm dòng của TEMPLATE để ràng buộc cột (dòng tổng, các dòng thành phần)
TPL_GROUPS = [(0, (1, 2, 3, 4)), (6, (8, 10, 12)), (7, (9, 11, 13)), (14, (15, 16, 17)), (18, (19, 20, 21))]
OCR_VARIANTS = [(28, 'lanczos', False), (28, 'lanczos', True), (36, 'lanczos', False), (28, 'bicubic', False), (36, 'lanczos', True)]


def _ocr_cell(img, box, numeric, dpi):
    """OCR 1 ô (psm 7) -> danh sách chuỗi ứng viên (nhiều biến thể phóng/lọc; ô chữ chỉ 1 biến thể). Ô trống -> [].
    Cắt từ ảnh GỐC rồi phóng sao cho chữ cao ~28-36px (phóng quá tay làm Tesseract nhầm 5/3, 8/3)."""
    import numpy as np, pytesseract
    from PIL import Image, ImageFilter
    x0, y0, x1, y1 = box
    crop = img[y0:y1, x0:x1].copy()
    if crop.size == 0: return []
    dark = crop < 160
    h, w = dark.shape
    kill = _long_runs(dark, max(int(dpi * 0.15), int(0.3 * w))) | _long_runs(dark.T, max(3, int(0.85 * h))).T
    kill[1:, :] |= kill[:-1, :].copy(); kill[:-1, :] |= kill[1:, :].copy()      # đường kẻ sót + gạch chân
    crop[kill] = 255
    dark = crop < 160
    if dark.sum() < max(8, 0.0003 * dpi * dpi): return []
    ys, xs = np.flatnonzero(dark.any(axis=1)), np.flatnonzero(dark.any(axis=0))
    crop = crop[max(ys[0] - 2, 0):ys[-1] + 3, max(xs[0] - 2, 0):xs[-1] + 3]
    th = max(ys[-1] - ys[0] + 1, 4)
    base = Image.fromarray(crop)
    out = []
    for target, rs, sharp in (OCR_VARIANTS if numeric else OCR_VARIANTS[:1]):
        k = max(target / th, 0.5) if numeric else max(34 / th, 0.5)
        im = base.resize((max(8, int(base.width * k)), max(8, int(base.height * k))),
                         Image.LANCZOS if rs == 'lanczos' else Image.BICUBIC)
        if sharp: im = im.filter(ImageFilter.UnsharpMask(radius=2, percent=150))
        pad = Image.new('L', (im.width + 60, im.height + 40), 255); pad.paste(im, (30, 20))
        if numeric:
            t = pytesseract.image_to_string(pad, lang='eng', config='--psm 7 -c tessedit_char_whitelist=0123456789.,%')
            out.append(re.sub(r'[^\d.,%]', '', t).strip('.,'))
        else:
            out.append(nfc(pytesseract.image_to_string(pad, lang='vie', config='--psm 7')).strip('|[]{}_ '))
    return out


def page_image(doc, page, dpi=300):
    """-> (ảnh xám, dpi tương đương). Trang chỉ chứa 1 ảnh chụp nhỏ (2023: ~600px) -> dùng ảnh gốc, không render phóng."""
    import io, numpy as np
    from PIL import Image
    ims = page.get_images(full=True)
    if len(ims) == 1:
        try:
            pil = Image.open(io.BytesIO(doc.extract_image(ims[0][0])['image'])).convert('L')
            if pil.width < 0.6 * page.rect.width / 72 * dpi:
                return np.array(pil), pil.width / (page.rect.width / 72)
        except Exception: pass
    pix = page.get_pixmap(dpi=dpi, colorspace='gray')
    return np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width).copy(), dpi


def deskew(img, dpi):
    """xoay thẳng ảnh scan: thử góc -2..2 độ trên mask đường kẻ ngang, chọn góc cho profile ngang nhọn nhất"""
    import numpy as np
    from PIL import Image
    hm = _long_runs(img < 160, int(dpi * 0.5))
    if hm.sum() == 0: return img
    k = max(1, int(dpi / 100))
    small = Image.fromarray((hm * 255).astype(np.uint8)).resize((img.shape[1] // k, img.shape[0] // k), Image.BILINEAR)
    best, ang = -1, 0.0
    for a in [x / 10 for x in range(-20, 21)]:
        prof = np.asarray(small.rotate(a, resample=Image.BILINEAR), dtype=np.float64).sum(axis=1)
        sc = float((prof ** 2).sum())
        if sc > best: best, ang = sc, a
    if abs(ang) < 0.05: return img
    return np.array(Image.fromarray(img).rotate(ang, resample=Image.BICUBIC, fillcolor=255))


def _rank(cands):
    """ứng viên chuỗi -> [(giá trị số, số phiếu)] giảm dần theo phiếu"""
    votes = defaultdict(int)
    for c in cands:
        v = num(c)
        if v is not None: votes[v] += 1
    return sorted(votes.items(), key=lambda kv: -kv[1])


def resolve_cells(cand, nrow, use_tpl, tol=1.5):
    """chọn giá trị cho từng ô số từ các ứng viên OCR. cand[(dòng, j)] j=0..6 = kế hoạch, A, B, C, D, %, %.
    Ưu tiên bộ (A, B, C) thoả A + B = C; cột D (cùng kỳ) + dòng tổng chọn theo tổng nhóm = các thành phần."""
    from itertools import product
    val = {}
    for (r, j), c in cand.items():
        rk = _rank(c)
        val[(r, j)] = rk[0][0] if rk else None
    for r in range(nrow):
        opts = [_rank(cand.get((r, j), [])) for j in (1, 2, 3)]
        if not all(opts): continue
        best = None
        for (a, va), (b, vb), (c, vc) in product(*opts):
            if abs(a + b - c) <= tol and (best is None or va + vb + vc > best[0]): best = (va + vb + vc, a, b, c)
        if best: val[(r, 1)], val[(r, 2)], val[(r, 3)] = best[1:]
    if use_tpl:
        for tot, parts in TPL_GROUPS:
            for j in (1, 2, 3, 4):
                cells = (tot,) + parts
                if any(val.get((r, j)) is None for r in cells): continue
                if abs(val[(tot, j)] - sum(val[(r, j)] for r in parts)) <= tol * 2: continue
                opts = [_rank(cand.get((r, j), [])) for r in cells]
                best = None
                for combo in product(*opts):
                    if abs(combo[0][0] - sum(x[0] for x in combo[1:])) <= tol * 2:
                        sc = sum(x[1] for x in combo)
                        if best is None or sc > best[0]: best = (sc, [x[0] for x in combo])
                if best and j == 4:                              # A, B, C đã có ràng buộc dòng; chỉ đổi D theo ràng buộc cột
                    for r, v in zip(cells, best[1]): val[(r, j)] = v
    return val


def ocr_table(doc, page, dpi=300):
    """Trang scan -> (grid, lines) | None. Lưới ô lấy từ đường kẻ bảng; mỗi ô số OCR riêng (chỉ nhận 0-9 . , %)."""
    import numpy as np, pytesseract
    from PIL import Image
    from concurrent.futures import ThreadPoolExecutor
    img, dpi = page_image(doc, page, dpi)
    img = deskew(img, dpi)
    dark = img < 160
    hmask = _long_runs(dark, int(dpi * 0.5)); vmask = _long_runs(dark.T, int(dpi * 0.3)).T
    colsum = vmask.sum(axis=0)
    if colsum.max() < dpi: return None
    gap = max(3, dpi * 0.05)
    xb = [int(round(x)) for x in _cluster(np.flatnonzero(colsum >= 0.45 * colsum.max()), gap)]
    if len(xb) < 8: return None
    rowsum = hmask[:, xb[0]:xb[-1]].sum(axis=1)
    yb = [int(round(y)) for y in _cluster(np.flatnonzero(rowsum >= 0.5 * (xb[-1] - xb[0])), gap)]
    mask = hmask | vmask
    for sh in ((1, 2) if dpi >= 200 else (1,)):
        mask[sh:, :] |= mask[:-sh, :].copy(); mask[:-sh, :] |= mask[sh:, :].copy()
        mask[:, sh:] |= mask[:, :-sh].copy(); mask[:, :-sh] |= mask[:, sh:].copy()
    img[mask] = 255
    ncol = len(xb) - 1
    nnum = 7                                                             # biểu tháng: 7 cột số (biểu năm chưa gặp ở dạng scan)
    lc = 0 if xb[1] - xb[0] >= xb[2] - xb[1] else 1                      # cột nhãn = cột rộng nhất trong 2 cột đầu (cột TT có thể mất viền trái)
    uc = lc + 1
    if ncol < uc + 1 + nnum: return None                                 # cột thừa bên phải biểu (T7, T8/2023) bị bỏ qua
    ncol = uc + 1 + nnum
    bands = [(yb[i], yb[i + 1]) for i in range(len(yb) - 1) if 0.1 * dpi <= yb[i + 1] - yb[i] <= 0.6 * dpi]
    m = max(2, int(dpi * 0.02))                                          # lề trong ô, tránh dính đường kẻ
    cell = lambda bi, ci: (xb[ci] + m, bands[bi][0] + m, xb[ci + 1] - m, bands[bi][1] - m)
    first = lambda c: c[0] if c else ''
    with ThreadPoolExecutor(max_workers=8) as ex:
        labels = [first(c) for c in ex.map(lambda bi: _ocr_cell(img, cell(bi, lc), False, dpi), range(len(bands)))]
        fl = [fold(t) for t in labels]
        # dòng "Tổng số" chữ đậm hay bị đọc sai -> neo theo 2 dòng ngay dưới (Hàng xuất khẩu, Hàng nhập khẩu)
        #   hoặc theo dòng đánh dấu cột "A B C" ngay trên; lấy vị trí sớm nhất trong các dấu hiệu
        start = next((i for i in range(len(fl)) if fl[i].endswith('ng so') or (i > 0 and fl[i - 1] == 'b')
                      or (i + 2 < len(fl) and 'xuat' in fl[i + 1] and 'nhap' in fl[i + 2])), None)
        if start is None: return None
        rows = list(range(start, min(start + len(TEMPLATE), len(bands))))
        hit = sum(1 for k, bi in enumerate(rows) if TEMPLATE[k][0] and fold(TEMPLATE[k][0]).split()[0] in fl[bi])
        use_tpl = len(rows) == len(TEMPLATE) and hit >= 0.5 * sum(1 for t in TEMPLATE if t[0])
        jobs = [(k, j) for k in range(len(rows)) for j in range(nnum)]
        cand = dict(zip(jobs, ex.map(lambda kj: _ocr_cell(img, cell(rows[kj[0]], uc + 1 + kj[1]), True, dpi), jobs)))
        units = {} if use_tpl else dict(zip(rows, (first(c) for c in ex.map(lambda bi: _ocr_cell(img, cell(bi, uc), False, dpi), rows))))
    val = resolve_cells({k: v for k, v in cand.items() if k[1] < 5}, len(rows), use_tpl)
    grid = []
    for k, bi in enumerate(rows):
        lab, unit = TEMPLATE[k] if use_tpl else (labels[bi], units[bi])
        pcts = []
        for j in (5, 6):
            c = [x for x in cand[(k, j)] if x]
            pcts.append(max(set(c), key=c.count) if c else '')
        grid.append([''] * lc + [lab, unit] + [val.get((k, j)) for j in range(5)] + pcts)
    hd = Image.fromarray(img[:bands[start][0], :])
    z = 300 / dpi
    if z > 1.3: hd = hd.resize((int(hd.width * z), int(hd.height * z)), Image.LANCZOS)
    txt = pytesseract.image_to_string(hd, lang='vie', config='--psm 6')
    return grid, [nfc(l) for l in txt.splitlines() if l.strip()]


def ocr_pdf(path, force=False):
    import pymupdf, pytesseract
    os.makedirs(OCR_DIR, exist_ok=True)
    cache = os.path.join(OCR_DIR, os.path.basename(path) + '.json')
    if os.path.exists(cache) and not force:
        with open(cache, encoding='utf-8') as f: return json.load(f) or None
    if os.path.exists(TESSERACT): pytesseract.pytesseract.tesseract_cmd = TESSERACT
    if os.path.isdir(TESSDATA): os.environ['TESSDATA_PREFIX'] = TESSDATA
    os.environ['OMP_THREAD_LIMIT'] = '1'
    doc = pymupdf.open(path)
    order = sorted(range(len(doc)), key=lambda i: (not re.search(r'phu luc i\b|bieu so|teus', fold(doc[i].get_text())), i))
    res = None
    for pi in order[:12]:
        try: r = ocr_table(doc, doc[pi])
        except Exception as e: log(f"  OCR lỗi trang {pi + 1}: {e.__class__.__name__}: {e}"); r = None
        if r and len(r[0]) >= 20:
            res = {'page': pi + 1, 'grid': r[0], 'lines': r[1]}; break
    with open(cache, 'w', encoding='utf-8') as f: json.dump(res, f, ensure_ascii=False)
    return res


def repair_ocr(rows, tol=3.0):
    """Sửa ô OCR đọc sai bằng số học của biểu, chỉ khi có 2 ràng buộc độc lập cùng xác nhận; ghi lại trong canh_bao.
       dòng:  luy_ke = luy_ke_thang_truoc + thang_bao_cao        cột:  dòng tổng nhóm = xuất + nhập + nội địa (+ quá cảnh bốc dỡ)"""
    idx = {(r['nhom'], r['chieu'], r['don_vi']): r for r in rows}
    A, B, C, D = 'luy_ke_thang_truoc', 'thang_bao_cao', 'luy_ke', 'luy_ke_cung_ky'
    nfix = 0

    def fix(r, k, new):
        nonlocal nfix
        r['canh_bao'] = (r.get('canh_bao', '') + f"; sua theo so hoc {k}: doc duoc {fmt(r[k]) or 'trong'} -> {fmt(new)}").strip('; ')
        r[k] = round(new, 3); nfix += 1

    def row_ok(r): return None not in (r[A], r[B], r[C]) and abs(r[A] + r[B] - r[C]) <= tol
    for nhom in NHOM_ORDER:
        for unit in ('1000 tấn', '1000 TEU'):
            grp = [idx[(nhom, ch, unit)] for ch in CHIEU_ORDER if (nhom, ch, unit) in idx]
            if len(grp) < 4 or grp[0]['chieu'] != 'tong': continue
            tot, parts = grp[0], grp[1:]
            for _ in range(3):                                   # 1 dòng thành phần sai đúng 1 ô
                bad = [r for r in parts if not row_ok(r)]
                if len(bad) != 1: break
                r = bad[0]
                cand = [k for k in (A, B, C) if tot[k] is not None and all(p[k] is not None for p in parts if p is not r)
                        and (r[k] is None or abs(tot[k] - sum(p[k] or 0 for p in parts)) > tol)]
                done = False
                for k in cand:
                    if any(r[kk] is None for kk in (A, B, C) if kk != k): continue
                    new = r[A] + r[B] if k == C else (r[C] - r[B] if k == A else r[C] - r[A])
                    if new >= 0 and abs(new - (tot[k] - sum(p[k] for p in parts if p is not r))) <= tol:
                        fix(r, k, new); done = True; break
                if not done: break
            if all(row_ok(p) for p in parts):                    # các dòng thành phần đã tự khớp -> dòng tổng = tổng thành phần
                for k in (A, B, C):
                    sm = sum(p[k] for p in parts)
                    if tot[k] is None or abs(tot[k] - sm) > tol: fix(tot, k, sm)
            if all(p[D] is not None for p in parts) and None not in (tot[C], tot['yoy_pct']):
                sm = sum(p[D] for p in parts)                    # cùng kỳ: xác nhận bằng cột % so cùng kỳ
                if (tot[D] is None or abs(tot[D] - sm) > tol) and sm and abs(tot[C] / sm * 100 - tot['yoy_pct']) <= 0.6: fix(tot, D, sm)
    return nfix


# ---------------------------------------------------------------- crawl + tải
def get(sess, url, **kw):
    for i in range(3):
        try:
            r = sess.get(url, timeout=60, **kw)
            if r.status_code == 200: return r
            log(f"  HTTP {r.status_code} {url}")
        except requests.RequestException as e:
            log(f"  lỗi mạng ({e.__class__.__name__}) {url}")
        time.sleep(2 + 3 * i)
    return None


def safe_name(s):
    return re.sub(r'[\\/:*?"<>|]', '_', s)


def crawl_list(sess, known, full):
    posts, page = [], 0
    while page < 60:
        r = get(sess, LIST_URL, params={'page': page} if page else None)
        if r is None: break
        doc = LH.fromstring(r.content)
        found = []
        for a in doc.xpath("//div[contains(@class,'region-content')]//a[starts-with(@href,'/vi/noi-dung/')]"):
            slug, title = a.get('href').rsplit('/', 1)[1], nfc(a.text_content())
            if title and slug not in [p['slug'] for p in found]: found.append({'slug': slug, 'title': title})
        if not found or all(p['slug'] in [q['slug'] for q in posts] for p in found): break
        posts += [p for p in found if p['slug'] not in [q['slug'] for q in posts]]
        if not full and all(p['slug'] in known for p in found): break
        page += 1; time.sleep(0.3)
    return posts


def fetch_post(sess, post, force=False):
    url = f"{BASE}/vi/noi-dung/{post['slug']}"
    r = get(sess, url)
    if r is None: return None
    doc = LH.fromstring(r.content)
    art = (doc.xpath("//article") or doc.xpath("//div[contains(@class,'region-content')]") or [doc])[0]
    files, urls = [], []
    for a in art.xpath(".//a[@href]"):
        u = urljoin(BASE, a.get('href'))
        if not re.search(r'\.(docx?|xlsx?|pdf)(\?|$)', u, re.I) or u in urls: continue
        fn = f"{post['slug'][:80]}__{safe_name(unquote(u.rsplit('/', 1)[1].split('?')[0]))}"
        fp = os.path.join(RAW, fn)
        if force or not os.path.exists(fp) or os.path.getsize(fp) == 0:
            rr = get(sess, u)
            if rr is None: continue
            with open(fp, 'wb') as f: f.write(rr.content)
        files.append(fn); urls.append(u)
    kind = 'file' if files else 'empty'
    if not files:
        tabs = [t for t in art.xpath('.//table') if 'danh muc loai hang' in fold(t.text_content())]
        if tabs:
            fn = f"{post['slug'][:80]}__inline.html"
            html = '<html><head><meta charset="utf-8"></head><body>' + ''.join(
                LH.tostring(t, encoding='unicode') for t in tabs) + '</body></html>'
            with open(os.path.join(RAW, fn), 'w', encoding='utf-8') as f: f.write(html)
            files, urls, kind = [fn], [url], 'inline'
        elif art.xpath('.//table'): kind = 'bieu-khac'
    return {'slug': post['slug'], 'title': post['title'], 'url': url, 'kind': kind, 'files': '|'.join(files),
            'file_urls': '|'.join(urls), 'first_seen': datetime.now().strftime('%Y-%m-%d')}


def update_index(full=False, force=False, refresh=2):
    os.makedirs(RAW, exist_ok=True); os.makedirs(DATA, exist_ok=True)
    index = {r['slug']: r for r in read_csv(INDEX_CSV)}
    sess = requests.Session(); sess.headers.update(HEADERS)
    posts = crawl_list(sess, set(index), full or force or not index)
    log(f"Danh sách: {len(posts)} bài đã duyệt, {sum(1 for p in posts if p['slug'] not in index)} bài mới")
    n = 0
    for i, p in enumerate(posts):
        old = index.get(p['slug'])
        redo = force or old is None or (i < refresh and old.get('kind') == 'empty')
        if not redo: continue
        rec = fetch_post(sess, p, force)
        if rec is None: continue
        if old: rec['first_seen'] = old.get('first_seen') or rec['first_seen']
        index[p['slug']] = rec; n += 1
        log(f"  + {p['title']}  [{rec['kind']}] {rec['files']}")
        time.sleep(0.3)
    order = {p['slug']: i for i, p in enumerate(posts)}
    rows = sorted(index.values(), key=lambda r: order.get(r['slug'], 10 ** 6))
    write_csv(INDEX_CSV, rows, INDEX_FIELDS)
    return n


# ---------------------------------------------------------------- build
def parse_all(ocr='auto', force=False):
    index = read_csv(INDEX_CSV)
    override = {r['slug']: r['period'] for r in read_csv(OVERRIDE_CSV) if r.get('slug') and r.get('period')}
    allrows, nfile = [], 0
    for rank, post in enumerate(index):                       # index xếp bài mới nhất trước
        for fn, furl in zip(filter(None, post['files'].split('|')), post['file_urls'].split('|')):
            fp = os.path.join(RAW, fn)
            if not os.path.exists(fp): log(f"  thiếu file {fn}"); continue
            ext = os.path.splitext(fn)[1].lower()
            try:
                if ext == '.docx': items = [('docx', s, g, c) for s, g, c in grids_docx(fp)]
                elif ext == '.xlsx': items = [('xlsx', s, g, c) for s, g, c in grids_xlsx(fp)]
                elif ext == '.html': items = [('html', s, g, c) for s, g, c in grids_html(fp)]
                elif ext == '.pdf': items = list(grids_pdf(fp, ocr, force))
                else: log(f"  bỏ qua định dạng {ext}: {fn}"); continue
            except Exception as e:
                log(f"  LỖI đọc {fn}: {e.__class__.__name__}: {e}"); continue
            # file nhiều sheet (Biểu 28 gộp T1..T12): kỳ phải lấy từ trong sheet, không suy từ tiêu đề bài
            multi = ext == '.xlsx' and sum(1 for it in items if parse_grid(it[2], it[3], '')) > 1
            ok = 0
            for kind, sheet, grid, ctx in items:
                res = parse_grid(grid, ctx, '' if multi else post['title'], override.get(post['slug']),
                                 prefer_title=(kind == 'pdf-ocr'))
                if not res: continue
                if kind == 'pdf-ocr':
                    nf = repair_ocr(res['rows'])
                    if nf: log(f"  OCR: tự sửa {nf} ô theo số học - {fn[:70]}")
                    d = res['ngay_bao_cao']                   # ngày báo cáo đọc từ ảnh: chỉ giữ khi khớp tháng của kỳ
                    if d and res['thang'] and d[:7] != f"{res['nam']}-{res['thang']:02d}": res['ngay_bao_cao'] = ''
                if multi: kind = 'xlsx-bieu28'
                nw = validate(res['rows'])
                period = f"{res['nam']}-{res['thang']:02d}" if res['ky_loai'] == 'thang' and res['thang'] else str(res['nam'])
                for r in res['rows']:
                    r.update({'period': period, 'ky_loai': res['ky_loai'], 'nam': res['nam'], 'thang': res['thang'] or '',
                              'ngay_bao_cao': res['ngay_bao_cao'], 'nguon_loai': kind, 'nguon_file': fn, 'sheet': sheet,
                              'url': furl, 'ghi_chu': res['ghi_chu'], '_rank': rank})
                allrows += res['rows']; ok += 1
                if nw: log(f"  ! {period} [{kind}] {fn[:70]}: {nw} cảnh báo số học")
            if ok: nfile += 1
            else: log(f"  ! không parse được biểu: {fn}")
    log(f"Parse: {nfile} file -> {len(allrows)} dòng (mọi phiên bản)")
    return allrows


def item_key(r):
    return (NHOM_ORDER.index(r['nhom']), CHIEU_ORDER.index(r['chieu']), 0 if 'tấn' in r['don_vi'] else 1)


def build(allrows):
    ver = defaultdict(list)
    for r in allrows: ver[(r['period'], r['nguon_file'], r['sheet'])].append(r)
    best = {}
    for (period, fn, sheet), rows in ver.items():
        r = rows[0]
        score = (1 if r['nguon_loai'] == 'xlsx-bieu28' else 0,                 # file riêng của kỳ > sheet trong Biểu 28 gộp
                 sum(1 for x in rows if x.get('canh_bao')), -len(rows), r['_rank'])
        if period not in best or score < best[period][0]: best[period] = (score, (period, fn, sheet))
    chosen = {v[1] for v in best.values()}
    for r in allrows: r['chon'] = 1 if (r['period'], r['nguon_file'], r['sheet']) in chosen else 0
    master = [r for r in allrows if r['chon']]
    # tháng thực tế = chênh lệch lũy kế-đến-hết-tháng-trước giữa 2 kỳ liên tiếp (cột 3 của biểu chỉ là ước)
    idx = {(r['ky_loai'], r['nam'], r['thang'], r['nhom'], r['chieu'], r['don_vi']): r for r in master}
    for r in master:
        r['thang_thuc_te'] = None
        if r['ky_loai'] != 'thang' or r['luy_ke_thang_truoc'] is None: continue
        k = (r['nhom'], r['chieu'], r['don_vi'])
        nxt = idx.get(('thang', r['nam'], r['thang'] + 1) + k) if r['thang'] < 12 else None
        if nxt and nxt['luy_ke_thang_truoc'] is not None: end = nxt['luy_ke_thang_truoc']
        elif r['thang'] == 12 and idx.get(('nam', r['nam'], '') + k): end = idx[('nam', r['nam'], '') + k]['luy_ke']
        else: continue
        if end is None or end - r['luy_ke_thang_truoc'] < 0: continue
        r['thang_thuc_te'] = round(end - r['luy_ke_thang_truoc'], 3)
        est = r['thang_bao_cao']
        if r['chieu'] == 'tong' and est and abs(r['thang_thuc_te'] / est - 1) > 0.35:   # nguồn đổi phạm vi thống kê / sửa số giữa 2 kỳ
            r['canh_bao'] = (r.get('canh_bao', '') + '; thang_thuc_te lech uoc >35% (nguon doi pham vi/so lieu giua 2 ky?)').strip('; ')
    sk = lambda r: (r['period'] if len(r['period']) > 4 else r['period'] + '-99', r['_rank'], r['nguon_file'], r['sheet'], item_key(r))
    master.sort(key=sk); allrows.sort(key=sk)
    write_csv(MASTER_CSV, master, FIELDS); write_csv(ALL_CSV, allrows, ALL_FIELDS)
    periods = sorted({r['period'] for r in master if r['ky_loai'] == 'thang'})
    years = sorted({r['period'] for r in master if r['ky_loai'] == 'nam'})
    log(f"Master: {len(master)} dòng, {len(periods)} kỳ tháng ({periods[0] if periods else '-'} .. {periods[-1] if periods else '-'}), "
        f"{len(years)} kỳ năm ({', '.join(years)})")
    if periods:
        y0, m0 = map(int, periods[0].split('-')); y1, m1 = map(int, periods[-1].split('-'))
        want = [f"{y}-{m:02d}" for y in range(y0, y1 + 1) for m in range(1, 13) if (y, m) >= (y0, m0) and (y, m) <= (y1, m1)]
        miss = [p for p in want if p not in periods]
        if miss: log(f"Kỳ tháng KHÔNG có trên web ({len(miss)}): {', '.join(miss)}")
    nwarn = sum(1 for r in master if r.get('canh_bao'))
    if nwarn: log(f"Dòng có cảnh báo số học trong master: {nwarn} (xem cột canh_bao)")
    log(f"DONE -> {MASTER_CSV}")
    return master


def main():
    ap = argparse.ArgumentParser(description='Sản lượng hàng hóa thông qua cảng biển (Cục HH&ĐT VN) -> CSV tổng')
    ap.add_argument('--full', action='store_true', help='duyệt hết các trang danh sách')
    ap.add_argument('--force', action='store_true', help='tải lại file + OCR lại')
    ap.add_argument('--no-fetch', action='store_true', help='không vào mạng, chỉ parse raw/')
    ap.add_argument('--ocr', choices=['auto', 'off'], default='auto', help='OCR PDF scan (mặc định auto, có cache)')
    a = ap.parse_args()
    if not a.no_fetch: update_index(a.full, a.force)
    rows = parse_all(a.ocr, a.force)
    if not rows: log('ERROR: không có dữ liệu'); return 1
    build(rows)
    return 0


if __name__ == '__main__':
    sys.exit(main())
