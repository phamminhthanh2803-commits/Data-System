# -*- coding: utf-8 -*-
r"""
CHART PACK TTCK - gom san du lieu 12 thang cho 5 nhom bieu do vao 1 file Excel.

Sheet:
  00_Huong_dan          : muc luc, don vi, nguon, ngay du lieu
  01_KhoiNgoai_ChauA    : mua/ban rong khoi ngoai theo THANG (trieu USD) - 6 thi truong
                          Chau A + tong EM Chau A + luy ke
  01b_KhoiNgoai_VN_Ngay : khoi ngoai VN theo NGAY (ty VND, 3 san + tu doanh) + luy ke
  02_VonHoa_Nhom        : VNINDEX / VN30 / VNMidcap / VNSmallcap - chi so, rebase 100,
                          va von hoa uoc tinh (nghin ty VND) theo nhom quy mo
  03_DoRong_TT          : do rong thi truong theo ngay (tang/giam, A/D line,
                          % ma tren MA20/50/100/200, so ma dinh/day 52 tuan)
  04_Sector_12M         : chi so nganh gia quyen von hoa, rebase 100 (12 thang)
  04b_Sector_TongKet    : hieu suat 1T/3T/6T/12T/YTD, % ma tang, von hoa, dong gop
  05_Duoi_MA            : so co phieu NAM DUOI MA50 / MA200 / MA300 theo ngay (3 san + tong)

Nguon (da co san trong cac pipeline):
  D:\market-data\index-fetcher\tv-history.csv            gia daily TradingView (1.237 ma VN + chi so)
  D:\market-data\index-fetcher\flows-master.csv          khoi ngoai/tu doanh VN + khu vuc
  D:\market-data\index-fetcher\fx-master.csv             ty gia de quy USD
  D:\market-data\index-fetcher\raw\vn_screener_meta.csv  so CP luu hanh + nganh (TradingView)

Chay:  python build_chart_pack.py            (mac dinh 12 thang)
       python build_chart_pack.py --months 24
"""
import argparse
import os
import sys
import warnings

warnings.filterwarnings("ignore")
import numpy as np
import pandas as pd

sys.stdout.reconfigure(encoding="utf-8")

BASE = os.path.normpath(os.environ.get("MD_ROOT", "D:/market-data"))
IDXD = os.path.join(BASE, "index-fetcher")
RAW = os.path.join(IDXD, "raw")
SCR = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(SCR, "Chart_Pack_TTCK.xlsx")

TVH = os.path.join(IDXD, "tv-history.csv")
FLOWS = os.path.join(IDXD, "flows-master.csv")
FX = os.path.join(IDXD, "fx-master.csv")
META = os.path.join(RAW, "vn_screener_meta.csv")

BROKERS = ("SSI", "VND", "VCI", "HCM", "VIX", "SHS", "MBS", "FTS", "BSI", "CTS",
           "VDS", "AGR", "ORS", "TCX", "VPX")


def grp(n, s, i):
    """Phan nhom nganh - GIONG three_month.py / ytd_report.py de so sanh duoc."""
    i, s = str(i), str(s)
    if n in ("VIC", "VHM", "VRE", "VPL"):
        return "Vingroup"
    if "Bank" in i and "Investment" not in i:
        return "Ngan hang"
    if "Investment Banks" in i or "Brokers" in i or n in BROKERS:
        return "Chung khoan"
    if "Real Estate" in i or "Homebuilding" in i:
        return "BDS (ngoai Vin)"
    if "Steel" in i:
        return "Thep"
    if s == "Finance":
        return "Tai chinh khac"
    if s in ("Consumer Non-Durables", "Consumer Durables", "Retail Trade",
             "Consumer Services", "Distribution Services"):
        return "Tieu dung, ban le"
    if s in ("Utilities", "Energy Minerals"):
        return "Dien, dau khi"
    if s in ("Technology Services", "Electronic Technology", "Communications"):
        return "Cong nghe, vien thong"
    if s in ("Process Industries", "Non-Energy Minerals", "Producer Manufacturing",
             "Industrial Services"):
        return "Cong nghiep, vat lieu"
    if s == "Transportation":
        return "Van tai, cang, HK"
    return "Khac"


VI = {
    "Ngan hang": "Ngân hàng", "Chung khoan": "Chứng khoán", "BDS (ngoai Vin)": "BĐS (ngoài Vin)",
    "Thep": "Thép", "Tai chinh khac": "Tài chính khác", "Tieu dung, ban le": "Tiêu dùng, bán lẻ",
    "Dien, dau khi": "Điện, dầu khí", "Cong nghe, vien thong": "Công nghệ, viễn thông",
    "Cong nghiep, vat lieu": "Công nghiệp, vật liệu", "Van tai, cang, HK": "Vận tải, cảng, HK",
    "Khac": "Khác", "Vingroup": "Vingroup",
}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--months", type=int, default=12)
    ap.add_argument("--include-today", action="store_true",
                    help="giu ca phien hom nay du chua dong cua (mac dinh cat bo)")
    args = ap.parse_args()
    M = args.months

    print("Nap tv-history.csv ...")
    buf_year = pd.Timestamp.now().year - (M // 12) - 3   # du buffer cho MA300 + dinh/day 52 tuan
    tv = pd.read_csv(TVH, usecols=["date", "tv_symbol", "exchange", "symbol", "close"], dtype={"date": str})
    tv = tv[tv.date >= f"{buf_year}-01-01"]
    tv["date"] = pd.to_datetime(tv.date)
    vn = tv[tv.exchange.isin(["HOSE", "HNX", "UPCOM"])]

    meta = pd.read_csv(META).drop_duplicates("name").set_index("name")
    meta["nhom"] = [VI[grp(n, s, i)] for n, s, i in zip(meta.index, meta.sector, meta.industry)]

    stocks = vn[vn.symbol.isin(meta.index)]
    END = stocks.date.max()
    now = pd.Timestamp.now()
    if not args.include_today and END.normalize() == now.normalize() and now.hour < 15:
        # phien hom nay CHUA dong cua -> so lieu do rong/khoi ngoai con thieu, cat bo
        stocks = stocks[stocks.date < now.normalize()]
        tv = tv[tv.date < now.normalize()]
        END = stocks.date.max()
        print("   (bo phien %s vi chua dong cua)" % now.strftime("%d/%m"))
    START = END - pd.DateOffset(months=M)
    print("Ky du lieu: %s -> %s (%d ma)" % (START.date(), END.date(), stocks.symbol.nunique()))

    px_raw = stocks.pivot_table(index="date", columns="symbol", values="close", aggfunc="last").sort_index()
    px = px_raw.ffill(limit=20)
    sessions = px.index[(px.index >= START) & (px.index <= END)]
    exch = meta.exchange.reindex(px.columns)
    sh = meta.total_shares_outstanding_fundamental
    vni_px = tv[tv.tv_symbol == "HOSE:VNINDEX"].set_index("date").close.sort_index()

    sheets = {}

    # ==================================================== 1. KHOI NGOAI CHAU A
    print("1. Khoi ngoai chau A ...")
    fl = pd.read_csv(FLOWS, dtype={"date": str})
    fl["date"] = pd.to_datetime(fl.date)
    fx = pd.read_csv(FX, dtype={"date": str})
    fx["date"] = pd.to_datetime(fx.date)
    fxm = fx.assign(ym=fx.date.dt.to_period("M")).groupby(["currency", "ym"]).rate.mean().unstack(0)

    MKT = [("VNINDEX", "Việt Nam (HOSE)", "VND"), ("KOSPI", "Hàn Quốc (KOSPI)", "KRW"),
           ("KOSDAQ", "Hàn Quốc (KOSDAQ)", "KRW"), ("TAIEX", "Đài Loan (TWSE)", "TWD"),
           ("JCI", "Indonesia (IDX)", "IDR"), ("SET", "Thái Lan (SET)", "THB"),
           ("FBMKLCI", "Malaysia (Bursa)", "MYR")]
    ymi = pd.period_range(START.to_period("M"), END.to_period("M"), freq="M")
    usd = {}
    for code, ten, cur in MKT:
        d = fl[(fl.index_code == code) & (fl.flow_type == "foreign")].dropna(subset=["net_val"]).copy()
        d["ym"] = d.date.dt.to_period("M")
        # nguon cong bo theo THANG (SET, Bursa) co them vai dong ngay cua thang hien tai ->
        # thang nao co dong M thi lay M, thang nao chi co D thi cong don D (tranh dem 2 lan)
        m_ser = d[d.freq == "M"].groupby("ym").net_val.sum()
        d_ser = d[d.freq != "M"].groupby("ym").net_val.sum()
        s = m_ser.reindex(ymi)
        s = s.fillna(d_ser.reindex(ymi))
        usd[ten] = s / fxm[cur].reindex(ymi).ffill() / 1e6
    f1 = pd.DataFrame(usd, index=ymi)
    f1["EM Châu Á (tổng)"] = f1.sum(axis=1, min_count=1)
    f1 = pd.concat([f1, f1.cumsum().add_prefix("Luỹ kế ")], axis=1).round(1)
    f1.index = pd.Index(["%02d/%d" % (p.month, p.year) for p in ymi], name="Tháng")
    sheets["01_KhoiNgoai_ChauA"] = f1

    # ============================================= 1b. KHOI NGOAI VN THEO NGAY
    print("1b. Khoi ngoai VN theo ngay ...")
    vnf = fl[fl.index_code.isin(["VNINDEX", "HNXINDEX", "UPCOM"])
             & (fl.date >= START) & (fl.date <= END)]
    piv = vnf.pivot_table(index="date", columns=["flow_type", "index_code"], values="net_val", aggfunc="sum") / 1e9
    ren = {("foreign", "VNINDEX"): "KN ròng HOSE", ("foreign", "HNXINDEX"): "KN ròng HNX",
           ("foreign", "UPCOM"): "KN ròng UPCoM", ("prop", "VNINDEX"): "Tự doanh HOSE",
           ("prop", "HNXINDEX"): "Tự doanh HNX", ("prop", "UPCOM"): "Tự doanh UPCoM"}
    piv.columns = [ren.get(tuple(c), "|".join(map(str, c))) for c in piv.columns]
    piv = piv.reindex(columns=[v for v in ren.values() if v in piv.columns])
    kn_cols = [c for c in piv.columns if c.startswith("KN")]
    piv["KN ròng toàn TT"] = piv[kn_cols].sum(axis=1, min_count=1)
    piv["Luỹ kế KN toàn TT"] = piv["KN ròng toàn TT"].fillna(0).cumsum()
    piv["VN-Index"] = vni_px.reindex(piv.index).ffill()
    piv.index.name = "Ngày"
    sheets["01b_KhoiNgoai_VN_Ngay"] = piv.round(2)

    # ==================================================== 2. VON HOA CAC NHOM
    print("2. Von hoa VN30 / Mid / Small / VNI ...")
    CAPS = {"VNINDEX": "VN-Index", "VN30": "VN30", "VNMIDCAP": "VNMidcap", "VNSMALLCAP": "VNSmallcap"}
    lv = pd.DataFrame(index=sessions)
    for code, ten in CAPS.items():
        s = tv[tv.tv_symbol == "HOSE:" + code].set_index("date").close.sort_index()
        lv[ten] = s.reindex(sessions).ffill(limit=3)
    reb = (lv / lv.iloc[0] * 100)
    reb.columns = [c + " (rebase 100)" for c in reb.columns]

    hose_cols = [c for c in px.columns if exch.get(c) == "HOSE" and pd.notna(sh.get(c))]
    mc = px.loc[sessions, hose_cols] * sh.reindex(hose_cols)
    rank = mc.iloc[-1].rank(ascending=False)
    buckets = {"Vốn hoá top 30 (≈VN30, nghìn tỷ)": (rank <= 30).values,
               "Vốn hoá hạng 31-100 (≈Midcap, nghìn tỷ)": ((rank > 30) & (rank <= 100)).values,
               "Vốn hoá còn lại (≈Smallcap, nghìn tỷ)": (rank > 100).values}
    cap = pd.DataFrame(index=sessions)
    for ten, m in buckets.items():
        cap[ten] = mc.loc[:, m].sum(axis=1) / 1e12
    cap["Vốn hoá toàn HOSE (nghìn tỷ)"] = mc.sum(axis=1) / 1e12
    out2 = pd.concat([lv.round(2), reb.round(2), cap.round(1)], axis=1)
    out2.index.name = "Ngày"
    sheets["02_VonHoa_Nhom"] = out2

    # ==================================================== 3. DO RONG THI TRUONG
    print("3. Do rong thi truong ...")
    r1 = px_raw.pct_change()
    adv, dec = (r1 > 0).sum(axis=1), (r1 < 0).sum(axis=1)
    br = pd.DataFrame({"Số mã tăng": adv, "Số mã giảm": dec,
                       "Số mã đứng giá": ((r1 == 0) & px_raw.notna()).sum(axis=1),
                       "Số mã giao dịch": px_raw.notna().sum(axis=1)})
    br["% mã tăng"] = adv / (adv + dec).replace(0, np.nan) * 100
    br["Tăng - Giảm"] = adv - dec
    br["A/D line (luỹ kế)"] = (adv - dec).cumsum()
    for w in (20, 50, 100, 200):
        ma = px.rolling(w, min_periods=w).mean()
        ok = px.notna() & ma.notna()
        br["%% mã trên MA%d" % w] = ((px > ma) & ok).sum(axis=1) / ok.sum(axis=1).replace(0, np.nan) * 100
    hi52 = px.rolling(250, min_periods=120).max()
    lo52 = px.rolling(250, min_periods=120).min()
    br["Số mã đỉnh 52 tuần"] = ((px >= hi52) & px.notna()).sum(axis=1)
    br["Số mã đáy 52 tuần"] = ((px <= lo52) & px.notna()).sum(axis=1)
    br["Đỉnh - Đáy 52T"] = br["Số mã đỉnh 52 tuần"] - br["Số mã đáy 52 tuần"]
    br["VN-Index"] = vni_px.reindex(br.index).ffill()
    br = br.loc[sessions].round(2)
    br.index.name = "Ngày"
    sheets["03_DoRong_TT"] = br

    # ======================================================= 4. SECTOR 12 THANG
    print("4. Sector performance ...")
    nhom = meta.nhom.reindex(hose_cols)
    groups = sorted(nhom.dropna().unique())
    # ro co dinh: chi lay ma co gia ca dau ky lan cuoi ky -> chi so nganh khong bi nhay
    # khi co ma moi len san / ma bi huy niem yet
    base = mc.columns[mc.iloc[0].notna() & mc.iloc[-1].notna()]
    mcb = mc[base]
    sec = pd.DataFrame(index=sessions)
    for g in groups:
        cols = [c for c in base if nhom[c] == g]
        if not cols:
            continue
        v = mcb[cols].sum(axis=1)
        sec[g] = v / v.iloc[0] * 100
    tot = mcb.sum(axis=1)
    sec["Toàn HOSE"] = tot / tot.iloc[0] * 100
    sec.index.name = "Ngày"
    sheets["04_Sector_12M"] = sec.round(2)

    px_s = px.loc[sessions]

    def p_at(months=None, since=None):
        if since is not None:
            sub = sessions[sessions <= pd.Timestamp(since)]
            return px_s.loc[sub.max()] if len(sub) else px_s.iloc[0]
        sub = sessions[sessions <= END - pd.DateOffset(months=months)]
        return px_s.loc[sub.max()] if len(sub) else px_s.iloc[0]

    p1 = px_s.iloc[-1]
    mc_tot0 = mc.iloc[0].sum()

    def wret(cols, p0):
        """Loi nhuan gia quyen von hoa dau ky, bo qua ma thieu gia 1 trong 2 dau."""
        s0, s1, w = p0[cols], p1[cols], sh.reindex(cols)
        ok = s0.notna() & s1.notna() & (s0 > 0) & w.notna()
        if not ok.any():
            return np.nan
        wt = (s0[ok] * w[ok])
        return float(np.average((s1[ok] / s0[ok] - 1) * 100, weights=wt)) if wt.sum() > 0 else np.nan

    rows = []
    for g in groups + ["TOÀN HOSE"]:
        cols = hose_cols if g == "TOÀN HOSE" else [c for c in hose_cols if nhom[c] == g]
        row = {"Nhóm ngành": g, "Số mã": len(cols),
               "Vốn hoá hiện tại (nghìn tỷ)": mc[cols].iloc[-1].sum() / 1e12,
               "Tỷ trọng vốn hoá (%)": mc[cols].iloc[-1].sum() / mc.iloc[-1].sum() * 100}
        for lab, mm in [("1 tháng", 1), ("3 tháng", 3), ("6 tháng", 6), ("%d tháng" % M, M)]:
            row["Hiệu suất " + lab + " (%)"] = wret(cols, p_at(months=mm))
        row["Hiệu suất YTD %d (%%)" % END.year] = wret(cols, p_at(since="%d-12-31" % (END.year - 1)))
        p0 = p_at(months=M)
        r12 = (p1[cols] / p0[cols] - 1).dropna()
        row["%% mã tăng trong %d tháng" % M] = (r12 > 0).mean() * 100 if len(r12) else np.nan
        row["Đóng góp vốn hoá %dT (điểm %%)" % M] = (((p1[cols] - p0[cols]) * sh.reindex(cols)).sum()
                                                     / mc_tot0 * 100)
        rows.append(row)
    sheets["04b_Sector_TongKet"] = pd.DataFrame(rows).set_index("Nhóm ngành").round(2)

    # ================================================== 5. DUOI MA50/200/300
    print("5. So co phieu duoi MA50 / MA200 / MA300 ...")
    ma_tbl = pd.DataFrame(index=sessions)
    for w in (50, 200, 300):
        ma = px.rolling(w, min_periods=w).mean()
        below = ((px < ma) & px.notna() & ma.notna()).loc[sessions]
        valid = (px.notna() & ma.notna()).loc[sessions]
        for ex in ("HOSE", "HNX", "UPCOM"):
            cols = [c for c in px.columns if exch.get(c) == ex]
            ma_tbl["Dưới MA%d - %s" % (w, ex)] = below[cols].sum(axis=1)
        ma_tbl["Dưới MA%d - Toàn TT" % w] = below.sum(axis=1)
        ma_tbl["Số mã đủ dữ liệu MA%d" % w] = valid.sum(axis=1)
        ma_tbl["%% dưới MA%d" % w] = below.sum(axis=1) / valid.sum(axis=1).replace(0, np.nan) * 100
    ma_tbl["VN-Index"] = vni_px.reindex(sessions).ffill()
    ma_tbl.index.name = "Ngày"
    sheets["05_Duoi_MA"] = ma_tbl.round(2)

    # ================================================================ GHI FILE
    print("Ghi Excel ...")
    doc = pd.DataFrame([
        ("01_KhoiNgoai_ChauA", "Mua/bán ròng khối ngoại theo tháng: VN (HOSE), Hàn Quốc (KOSPI, KOSDAQ), "
                              "Đài Loan, Indonesia, Thái Lan, Malaysia + tổng EM châu Á + luỹ kế",
         "triệu USD (âm = bán ròng)", "flows-master.csv (VNDirect, Naver/KRX, TWSE, IDX, SET, Bursa) + fx-master.csv"),
        ("01b_KhoiNgoai_VN_Ngay", "Khối ngoại + tự doanh VN theo ngày, 3 sàn, luỹ kế", "tỷ VND",
         "flows-master.csv (VNDirect finfo v4)"),
        ("02_VonHoa_Nhom", "VN-Index / VN30 / VNMidcap / VNSmallcap: chỉ số, rebase 100, vốn hoá nhóm quy mô",
         "điểm · rebase 100 · nghìn tỷ VND", "tv-history.csv (TradingView) + số CP screener TradingView"),
        ("03_DoRong_TT", "Độ rộng: số mã tăng/giảm, A/D line, % mã trên MA20/50/100/200, đỉnh - đáy 52 tuần",
         "số mã · %", "tv-history.csv, toàn bộ HOSE + HNX + UPCoM"),
        ("04_Sector_12M", "Chỉ số ngành gia quyền vốn hoá, rebase 100 đầu kỳ", "rebase 100",
         "tv-history.csv + ngành/số CP screener TradingView (HOSE)"),
        ("04b_Sector_TongKet", "Hiệu suất 1T/3T/6T/12T/YTD, % mã tăng, vốn hoá, đóng góp", "% · nghìn tỷ", "như trên"),
        ("05_Duoi_MA", "Số cổ phiếu NẰM DƯỚI MA50 / MA200 / MA300 theo ngày, tách 3 sàn + toàn thị trường",
         "số mã · %", "tv-history.csv"),
    ], columns=["Sheet", "Nội dung", "Đơn vị", "Nguồn"])

    note = pd.DataFrame([
        ("Kỳ dữ liệu", "%s - %s (%d tháng)" % (START.strftime("%d/%m/%Y"), END.strftime("%d/%m/%Y"), M)),
        ("Ngày dựng file", pd.Timestamp.now().strftime("%d/%m/%Y %H:%M")),
        ("Số mã VN dùng để tính", "%d mã (HOSE %d, HNX %d, UPCoM %d)" % (
            len(px.columns), int((exch == "HOSE").sum()), int((exch == "HNX").sum()),
            int((exch == "UPCOM").sum()))),
        ("Giá", "giá đóng cửa đã điều chỉnh (TradingView); MA tính trên giá điều chỉnh"),
        ("Vốn hoá nhóm quy mô", "ƯỚC TÍNH = số CP lưu hành hiện tại × giá; nhóm chia theo thứ hạng vốn hoá HOSE "
                                "cuối kỳ (top 30 / 31-100 / còn lại), KHÔNG phải rổ chính thức của HOSE. "
                                "Cột chỉ số VN30/VNMidcap/VNSmallcap mới là rổ chính thức."),
        ("Khối ngoại Thái Lan / Malaysia", "nguồn chỉ công bố theo THÁNG (SET investor-type, Bursa PDF)"),
        ("Quy đổi USD", "tỷ giá bình quân tháng (fx-master.csv)"),
        ("Cách dựng lại", "python D:\\market-data\\chart-pack\\build_chart_pack.py [--months N]"),
    ], columns=["Mục", "Chi tiết"])

    with pd.ExcelWriter(OUT, engine="openpyxl", datetime_format="dd/mm/yyyy") as xw:
        doc.to_excel(xw, sheet_name="00_Huong_dan", index=False, startrow=0)
        note.to_excel(xw, sheet_name="00_Huong_dan", index=False, startrow=len(doc) + 3)
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name)

    from openpyxl import load_workbook
    from openpyxl.styles import Font, Alignment
    wb = load_workbook(OUT)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2" if ws.title == "00_Huong_dan" else "B2"
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[1].height = 32
        for col in ws.columns:
            w = max((len(str(c.value)) for c in col[:80] if c.value is not None), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max(11, w + 1), 36)
    wb.save(OUT)

    print("\nXONG -> %s" % OUT)
    for name, df in sheets.items():
        print("  %-22s %5d dong x %3d cot" % (name, df.shape[0], df.shape[1]))


if __name__ == "__main__":
    main()
