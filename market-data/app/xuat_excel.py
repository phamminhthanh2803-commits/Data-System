# -*- coding: utf-8 -*-
"""Xuat Excel "dung nhu dang thay": ghi lai MOI bieu do / bang / KPI app ve trong 1 lan chay
(dung khung thoi gian, phan nganh, bo loc dang chon) roi dung 1 file .xlsx co data + chart
Excel that (LineChart / BarChart cua openpyxl, sua duoc trong Excel).

Luong chay:
  - app goi bat_dau(PAGE) dau script, cac ham ve goi ghi(...), bang goi ghi_bang(...),
    kpi goi ghi_kpi(...), tab boc bang tabs(...) de biet dang o tab nao.
  - cuoi script app goi ket_thuc(PAGE): dung file trang dang xem khi co yeu cau, hoac chay
    tiep hang doi "toan app" (moi luot rerun 1 trang, gom ban ghi vao session_state).
"""
from __future__ import annotations

import io
import re
from datetime import datetime

import numpy as np
import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.axis import DateAxis
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

K_RUN, K_CTX, K_ALL, K_Q = "_xuat_run", "_xuat_ctx", "_xuat_all", "_xuat_queue"
K_TRY, K_SNAP, K_FILE, K_REQ, K_LOI = "_xuat_try", "_xuat_snap", "_xuat_file", "_xuat_req", "_xuat_loi"
MAX_DONG = 6000          # bang dai hon -> chi giu MAX_DONG dong cuoi (ghi chu trong file)
TAT = {"Tổng quan": "TQ", "Việt Nam": "VN", "Khu vực & thế giới": "KV", "Vĩ mô & tiền tệ": "VM",
       "Trái phiếu & nhà đầu tư": "TP", "Biểu đồ nhúng": "BD", "Kho dữ liệu": "KHO"}


# ================================================================== GHI LAI
def _ss():
    return st.session_state


def bat_dau(page: str):
    _ss()[K_RUN] = []
    _ss()[K_CTX] = {"page": page, "tab": [], "tieu_de": ""}


def _ctx():
    return _ss().get(K_CTX) or {"page": "", "tab": [], "tieu_de": ""}


def dat_tieu_de(t):
    if t:
        _ctx()["tieu_de"] = re.sub(r"[*_`#]", "", str(t)).strip().split("\n")[0][:120]


def _them(rec):
    if K_RUN not in _ss():
        return
    c = _ctx()
    rec.update(page=c["page"], tab=" / ".join(c["tab"]) or "Chung")
    _ss()[K_RUN].append(rec)


def ghi(loai: str, tieu_de: str, df, don_vi: str = "", **kw):
    """loai: line | bar | stack | hbar | bar_line | nguong | custom | treemap."""
    try:
        if isinstance(df, pd.Series):
            df = df.to_frame(df.name or "Giá trị")
        if df is None or df.empty:
            return
        _them({"loai": loai, "tieu_de": tieu_de or _ctx()["tieu_de"] or "Biểu đồ",
               "df": df.copy(), "don_vi": don_vi or "", **kw})
    except Exception:  # noqa: BLE001 - ghi loi khong duoc lam hong app
        pass


def ghi_bang(data, tieu_de=None):
    try:
        df = getattr(data, "data", data)            # Styler -> DataFrame
        if isinstance(df, pd.Series):
            df = df.to_frame()
        if not isinstance(df, pd.DataFrame) or df.empty:
            return
        css, ep = {}, None
        if hasattr(data, "_compute"):               # Styler: lay mau chu / dam / nen tung o
            data._compute()
            css = {k: list(v) for k, v in data.ctx.items() if v}
            ep = _styler_precision(data)
        _them({"loai": "bang", "tieu_de": tieu_de or _ctx()["tieu_de"] or "Bảng", "df": df.copy(),
               "don_vi": "", "css": css, "do_chinh_xac": ep})
    except Exception:  # noqa: BLE001
        pass


def _styler_precision(sty):
    """Styler.format(precision=n) -> n (doc qua ham hien thi cua 1 o so bat ky)."""
    try:
        for (i, j), f in sty._display_funcs.items():
            v = sty.data.iat[i, j]
            if isinstance(v, (float, np.floating)) and np.isfinite(v):
                t = f(1.123456789)
                return len(t.split(".")[1]) if "." in t else 0
    except Exception:  # noqa: BLE001
        pass
    return None


def ghi_kpi(label, value, delta=None):
    _them({"loai": "kpi", "tieu_de": str(label), "gia_tri": str(value),
           "delta": "" if delta is None else str(delta)})


def ghi_altair(ch, ten=""):
    """Bieu do Altair tu dung (khong qua ham ve chung): tach data tung lop -> bang wide."""
    try:
        lops = list(getattr(ch, "layer", []) or []) or [ch]
        goc = getattr(ch, "data", None)
        for i, lp in enumerate(lops):
            d = getattr(lp, "data", None)
            d = d if isinstance(d, pd.DataFrame) else goc
            if not isinstance(d, pd.DataFrame) or d.empty:
                continue
            enc = lp.encoding

            def truong(k):
                e = getattr(enc, k, None)
                if e is None or type(e).__name__ == "UndefinedType":
                    return None
                f = getattr(e, "field", None)
                if isinstance(f, str):
                    return f
                sh = getattr(e, "shorthand", None)
                return sh.split(":")[0] if isinstance(sh, str) and sh else None
            x, y, c = truong("x"), truong("y"), truong("color")
            mark = lp.mark if isinstance(lp.mark, str) else getattr(lp.mark, "type", "line")
            if x in d and y in d:
                if c in d and c not in (x, y):
                    w = d.pivot_table(index=x, columns=c, values=y, aggfunc="sum")
                else:
                    w = d.groupby(x)[[y]].sum()
                loai = "line" if mark in ("line", "area", "point", "circle") else "bar"
                ghi(loai, f"{_ctx()['tieu_de'] or ten}" + (f" (lớp {i + 1})" if len(lops) > 1 else ""), w)
            else:
                ghi_bang(d, f"{_ctx()['tieu_de'] or ten} - dữ liệu")
    except Exception:  # noqa: BLE001
        pass


class _Tab:
    """Boc st.tabs de biet bieu do dang nam o tab nao (ho tro tab long nhau)."""

    def __init__(self, tab, ten):
        self._tab, self._ten = tab, ten

    def __enter__(self):
        _ctx()["tab"].append(self._ten)
        return self._tab.__enter__()

    def __exit__(self, *a):
        if _ctx()["tab"]:
            _ctx()["tab"].pop()
        return self._tab.__exit__(*a)

    def __getattr__(self, k):
        return getattr(self._tab, k)


def tabs(names, **kw):
    return [_Tab(t, n) for t, n in zip(st.tabs(names, **kw), names)]


# ============================================================ LUONG XUAT
def dang_xuat_all() -> bool:
    return bool(_ss().get(K_Q))


def bat_dau_xuat_all(trang_hien_tai: str, cac_trang: list, giu_khoa=()):
    """Nut 'Toan bo app': luot hien tai (trang dang xem) da ve xong -> ban ghi cua no lay o
    ket_thuc(); cac trang con lai xep hang doi."""
    _ss()[K_ALL] = {}
    _ss()[K_TRY] = {}
    _ss()[K_LOI] = []
    _ss()[K_Q] = [p for p in cac_trang if p != trang_hien_tai]
    _ss()["_xuat_dau"] = trang_hien_tai
    # snapshot gia tri widget co key cua trang dang xem (Streamlit xoa state widget khong ve)
    _ss()[K_SNAP] = {k: v for k, v in _ss().items()
                     if isinstance(k, str) and not k.startswith(("_", "dl_", "FormSubmitter"))
                     and k not in giu_khoa and not callable(v)}
    _ss().pop(K_FILE, None)


def trang_luot_nay(page: str) -> str:
    """Dang xuat toan app -> tra trang ke tiep trong hang doi (trang loi 2 lan -> bo qua)."""
    if _ss().get("_xuat_dau"):          # luot bam nut: ve trang dang xem truoc
        return page
    q = _ss().get(K_Q) or []
    while q:
        p = q[0]
        n = _ss()[K_TRY].get(p, 0) + 1
        _ss()[K_TRY][p] = n
        if n <= 1:
            return p
        _ss()[K_LOI].append(p)          # luot truoc chet giua chung (loi / st.stop) -> bo qua
        q.pop(0)
    return page


def ket_thuc(page: str, cai_dat: dict, cho_nut=None):
    """Goi o CUOI script. cho_nut: st.empty() o thanh ben de hien nut tai / tien do."""
    rec = _ss().get(K_RUN, [])
    if _ss().pop(K_REQ, False):                               # xuat trang dang xem
        _ss()[K_FILE] = (dung_excel({page: rec}, cai_dat), f"MarketData_{_ten_file(page)}")
    if K_Q in _ss():
        _ss()[K_ALL][page] = rec
        q = _ss()[K_Q]
        if q and q[0] == page:
            q.pop(0)
        if _ss().get("_xuat_dau") == page:
            _ss()["_xuat_dau"] = None
        if q:
            st.rerun()
        # xong hang doi -> dung file, khoi phuc widget trang dang xem, ve lai trang do
        thu_tu = [p for p in cai_dat.get("thu_tu_trang", []) if p in _ss()[K_ALL]]
        tat_ca = {p: _ss()[K_ALL][p] for p in thu_tu}
        cai_dat = dict(cai_dat, trang_loi=", ".join(_ss().get(K_LOI, [])) or "không")
        _ss()[K_FILE] = (dung_excel(tat_ca, cai_dat), "MarketData_ToanApp")
        for k, v in (_ss().pop(K_SNAP, {}) or {}).items():
            if k not in _ss():
                try:
                    _ss()[k] = v
                except Exception:  # noqa: BLE001
                    pass
        for k in (K_Q, K_ALL, K_TRY):
            _ss().pop(k, None)
        st.rerun()
    if cho_nut is not None and K_FILE in _ss():
        data, ten = _ss()[K_FILE]
        with cho_nut.container():
            st.download_button("⬇️ Tải file Excel", data, file_name=f"{ten}_{datetime.now():%Y%m%d_%H%M}.xlsx",
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                               type="primary", width="stretch", key="_xuat_tai")
            st.caption(f"{ten} · {len(data) / 1e6:.1f} MB")


def _ten_file(page):
    return re.sub(r"[^A-Za-z0-9]+", "", _bo_dau(_ten_trang(page))) or "Trang"


def _ten_trang(page):
    return re.sub(r"^[^\wÀ-ỹ]+", "", page).strip()


def _bo_dau(s):
    import unicodedata
    s = unicodedata.normalize("NFD", s).replace("đ", "d").replace("Đ", "D")
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


# ================================================================ DUNG FILE
# Dinh dang = dung nhu app: bang mau TH dang chon (Bao cao / Toi), mau tung series theo thu tu
# Vega gan (domain sap xep tang dan), cot duong/am doi mau, luoi ngang, khong luoi doc, nen,
# chu truc 11px mau ink2, dinh dang so theo _fmt cua app, truc ngay co dinh khung dang xem.
from openpyxl.chart.layout import Layout, ManualLayout  # noqa: E402
from openpyxl.chart.marker import DataPoint  # noqa: E402
from openpyxl.chart.shapes import GraphicalProperties  # noqa: E402
from openpyxl.chart.text import RichText  # noqa: E402
from openpyxl.chart.title import Title  # noqa: E402
from openpyxl.chart.text import Text  # noqa: E402
from openpyxl.drawing.line import LineProperties  # noqa: E402
from openpyxl.drawing.text import (CharacterProperties, Font as DFont, Paragraph,  # noqa: E402
                                   ParagraphProperties, RegularTextRun, RichTextProperties)

K_TH = "_xuat_th"
TH_MAC_DINH = {"bg": "#ffffff", "ink": "#262626", "ink2": "#595959", "grid": "#D9D9D9",
               "bar": "#BFBFBF", "pos": "#00B050", "neg": "#C00000", "nhan": "#ED7D31",
               "range": ["#262626", "#ED7D31", "#A6A6A6", "#F4B183", "#C00000", "#00B050",
                         "#4472C4", "#843C0C", "#7F7F7F", "#BF9000", "#2E75B6", "#D9D9D9"]}
FONT = "Source Sans Pro"          # font cua app (Streamlit); may khong co -> Excel tu thay
PX = 0.0265                       # 1 px man hinh ~ 0,0265 cm
EMU_PX = 9525                     # 1 px = 9525 EMU (do day net)


def dat_giao_dien(th: dict):
    """App goi moi luot chay: luu bang mau dang dung de file xuat giong het man hinh."""
    _ss()[K_TH] = {k: th[k] for k in ("bg", "ink", "ink2", "grid", "bar", "pos", "neg", "nhan", "range")}


def _th():
    try:
        return _ss().get(K_TH) or TH_MAC_DINH
    except Exception:  # noqa: BLE001 - ngoai streamlit (test)
        return TH_MAC_DINH


def _h(c):
    return str(c).lstrip("#").upper()[:6]


def mau_series(names, rng=None):
    """Mau Vega gan cho tung series: domain nominal sap xep tang dan, lap vong theo range."""
    rng = rng or _th()["range"]
    dom = sorted({str(n) for n in names})
    return {str(n): rng[dom.index(str(n)) % len(rng)] for n in names}


def fmt_excel(d3: str) -> str:
    """',.2f' -> '#,##0.00' (dau phan cach theo Windows cua may mo file)."""
    m = re.search(r"\.(\d)f", d3 or "")
    n = int(m.group(1)) if m else 0
    return "#,##0" + ("." + "0" * n if n else "")


def fmt_ngay(strf: str) -> str:
    return {"%d/%m": "dd/mm", "%m/%Y": "mm/yyyy", "%Y": "yyyy"}.get(strf or "", "mm/yyyy")


def _ser(ts) -> float:
    return (pd.Timestamp(ts) - pd.Timestamp("1899-12-30")).days


def _chu(sz_pt=8.25, mau=None, dam=False):
    return CharacterProperties(sz=int(sz_pt * 100), b=dam, solidFill=_h(mau or _th()["ink2"]),
                               latin=DFont(typeface=FONT))


def _txpr(sz_pt=8.25, mau=None, dam=False):
    return RichText(bodyPr=RichTextProperties(), p=[Paragraph(pPr=ParagraphProperties(
        defRPr=_chu(sz_pt, mau, dam)), endParaRPr=_chu(sz_pt, mau, dam))])


def _net(mau, w_px=1.0, dash=None):
    ln = LineProperties(solidFill=_h(mau), w=int(w_px * EMU_PX))
    if dash:
        ln.prstDash = dash
    return ln


def _tieu_de(text):
    """Tieu de dam, can trai, tren cung - giong dong **tieu de** markdown cua app."""
    cp = _chu(12, _th()["ink"], True)
    t = Title(tx=Text(rich=RichText(bodyPr=RichTextProperties(), p=[
        Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=text)])])),
        overlay=False)
    t.layout = Layout(manualLayout=ManualLayout(x=0.005, y=0.0, xMode="edge", yMode="edge"))
    return t


def _nice(lo, hi, zero, n=6):
    """Mien truc Y kieu Vega nice (tickCount 6)."""
    if zero:
        lo, hi = min(lo, 0.0), max(hi, 0.0)
    if hi <= lo:
        hi = lo + (abs(lo) or 1.0)
    raw = (hi - lo) / n
    mu = 10 ** np.floor(np.log10(raw))
    step = min((s * mu for s in (1, 2, 2.5, 5, 10) if s * mu >= raw), default=10 * mu)
    return float(np.floor(lo / step) * step), float(np.ceil(hi / step) * step), float(step)


def _dinh_dang_chart(ch, rec, vals, la_ngay, cot_ngang=False):
    th, k = _th(), rec.get("kieu", {})
    # nen + khung giong app (khong vien)
    ch.graphical_properties = GraphicalProperties(solidFill=_h(th["bg"]), ln=LineProperties(noFill=True))
    ch.plot_area.graphicalProperties = GraphicalProperties(solidFill=_h(th["bg"]), ln=LineProperties(noFill=True))
    cat_ax, val_ax = ch.x_axis, ch.y_axis
    for ax in (cat_ax, val_ax):
        ax.delete = False
        ax.txPr = _txpr(8.25, th["ink2"])
        ax.spPr = GraphicalProperties(ln=_net(th["grid"], 1))
        ax.majorTickMark = "out"
    cat_ax.majorGridlines = None                                  # grid=False truc X
    val_ax.majorGridlines.spPr = GraphicalProperties(ln=_net(th["grid"], 0.8))
    val_ax.spPr = GraphicalProperties(ln=LineProperties(noFill=True))   # truc Y chi co nhan
    val_ax.majorTickMark = "none"
    cat_ax.tickLblPos = "low"
    if rec.get("don_vi"):
        val_ax.title = rec["don_vi"]
        val_ax.title.overlay = False
        val_ax.title.tx.rich.p[0].pPr = ParagraphProperties(defRPr=_chu(8.25, th["ink2"]))
        for p in val_ax.title.tx.rich.p:
            for r in (p.r or []):
                r.rPr = _chu(8.25, th["ink2"])
    # dinh dang so + mien truc giong app
    y_fmt = k.get("y_fmt")
    if vals is not None and len(vals):
        lo, hi = float(np.nanmin(vals)), float(np.nanmax(vals))
        if not y_fmt:
            m = max(abs(lo), abs(hi))
            y_fmt = ",.2f" if m < 10 else (",.1f" if m < 1000 else ",.0f")
        a, b, step = _nice(lo, hi, k.get("zero", True))
        val_ax.scaling.min, val_ax.scaling.max, val_ax.majorUnit = a, b, step
    val_ax.number_format = fmt_excel(y_fmt or ",.0f")
    if la_ngay and not cot_ngang:
        xf = k.get("x_fmt")
        cat_ax.number_format = fmt_ngay(xf)
        if isinstance(cat_ax, DateAxis):
            x0, x1 = k.get("x0"), k.get("x1")
            if x0 is not None and x1 is not None:              # truc X co dinh = khung dang xem
                cat_ax.scaling.min, cat_ax.scaling.max = _ser(x0), _ser(x1)
            span = (cat_ax.scaling.max or 0) - (cat_ax.scaling.min or 0)
            if span > 0:                                        # ~8 moc nhu tickCount=8 cua app
                if xf == "%Y":
                    cat_ax.majorTimeUnit, cat_ax.majorUnit = "years", max(1, round(span / 365.25 / 8))
                elif xf == "%d/%m":
                    cat_ax.majorTimeUnit, cat_ax.majorUnit = "days", max(1, round(span / 8))
                else:
                    cat_ax.majorTimeUnit, cat_ax.majorUnit = "months", max(1, round(span / 30.44 / 8))
                cat_ax.baseTimeUnit = "days"
        else:                                                   # truc chu (cot): moi n nhan hien 1
            n = int(k.get("_n_cat") or 1)
            cat_ax.tickLblSkip = max(1, int(np.ceil(n / 8)))
            cat_ax.tickMarkSkip = cat_ax.tickLblSkip
    if ch.legend is not None:
        ch.legend.position = "b"
        ch.legend.overlay = False
        ch.legend.txPr = _txpr(8.25, th["ink2"])
    ch.title = _tieu_de(rec["tieu_de"][:110])


def _mau_cua(rec, names):
    """Mau theo thu tu cot: uu tien mau app truyen sang, khong co thi tinh nhu Vega."""
    k = rec.get("kieu", {})
    m = k.get("mau") or {}
    tu_tinh = mau_series(names)
    return [m.get(str(n), tu_tinh[str(n)]) for n in names]


def _ve_chart(ws, rec, df, r_hdr, r_end, ncol, la_ngay, anchor):
    th, k = _th(), rec.get("kieu", {})
    loai = rec["loai"]
    if ncol < 2 or r_end <= r_hdr:
        return 0
    names = list(df.columns[1:])
    cats = Reference(ws, min_col=1, min_row=r_hdr + 1, max_row=r_end)
    data = Reference(ws, min_col=2, max_col=ncol, min_row=r_hdr, max_row=r_end)
    so = df.iloc[:, 1:].apply(pd.to_numeric, errors="coerce")
    vals = so.to_numpy(dtype=float).ravel()
    vals = vals[np.isfinite(vals)]
    cao_cm = (k.get("cao_px") or 340) * PX + 1.2                   # + cho tieu de
    rong_cm = k.get("rong_cm", 24.0)
    cot_ngang = False

    if loai in ("line", "nguong", "custom_line"):
        ch = LineChart()
        if la_ngay:
            ch.x_axis = DateAxis(crossAx=100)
            ch.y_axis.crossAx = 500
        ch.add_data(data, titles_from_data=True)
        ch.set_categories(cats)
        mau = _mau_cua(rec, names)
        dut = set(k.get("net_dut", []))
        for s, n, m in zip(ch.series, names, mau):
            w = k.get("lw", 1.9) if n not in dut else 1.6
            s.graphicalProperties.line = _net(m, w, "dash" if n in dut else None)
            s.smooth = False
            s.marker.symbol = "none"
        ch.display_blanks = "span"
    elif loai in ("bar", "stack", "hbar", "custom_bar"):
        ch = BarChart()
        ch.x_axis.auto = False                                   # truc chu: cot deu nhau nhu app
        cot_ngang = loai == "hbar"
        ch.type = "bar" if cot_ngang else "col"
        if loai == "stack":
            ch.grouping, ch.overlap = "stacked", 100
            vals = np.concatenate([so.clip(lower=0).sum(axis=1).to_numpy(), so.clip(upper=0).sum(axis=1).to_numpy()])
        ch.gapWidth = 30 if not cot_ngang else 25
        ch.add_data(data, titles_from_data=True)
        ch.set_categories(cats)
        if len(names) == 1:                                          # 1 series: mau don / theo dau
            s = ch.series[0]
            goc = k.get("mau_cot") or (th["nhan"] if cot_ngang else th["bar"])
            s.graphicalProperties = GraphicalProperties(solidFill=_h(th["pos"] if k.get("sign") else goc),
                                                        ln=LineProperties(noFill=True))
            s.invertIfNegative = False
            if k.get("sign"):
                for i, v in enumerate(so.iloc[:, 0].to_numpy()):
                    if np.isfinite(v) and v < 0:
                        dp = DataPoint(idx=i, invertIfNegative=False)
                        dp.graphicalProperties = GraphicalProperties(solidFill=_h(th["neg"]),
                                                                     ln=LineProperties(noFill=True))
                        s.dPt.append(dp)
            ch.legend = None
        else:
            for s, m in zip(ch.series, _mau_cua(rec, names)):
                s.graphicalProperties = GraphicalProperties(solidFill=_h(m), ln=LineProperties(noFill=True))
                s.invertIfNegative = False
        if cot_ngang:                                                 # hang 1 o tren, truc gia tri o duoi
            ch.x_axis.scaling.orientation = "maxMin"
            ch.y_axis.crosses = "max"
            cao_cm = max(k.get("cao_px") or 24 * (r_end - r_hdr), 180) * PX + 1.2
            rong_cm = k.get("rong_cm", 16.0)
    elif loai == "bar_line":
        ch = BarChart()
        ch.x_axis.auto = False
        ch.add_data(Reference(ws, min_col=2, max_col=2, min_row=r_hdr, max_row=r_end), titles_from_data=True)
        ch.set_categories(cats)
        ch.gapWidth = 30
        s = ch.series[0]
        s.graphicalProperties = GraphicalProperties(solidFill=_h(k.get("mau_cot") or th["bar"]),
                                                    ln=LineProperties(noFill=True))
        s.invertIfNegative = False
        if k.get("sign"):
            for i, v in enumerate(so.iloc[:, 0].to_numpy()):
                dp = DataPoint(idx=i, invertIfNegative=False)
                dp.graphicalProperties = GraphicalProperties(
                    solidFill=_h(th["pos"] if np.isfinite(v) and v >= 0 else th["neg"]), ln=LineProperties(noFill=True))
                s.dPt.append(dp)
        if ncol > 2:
            ln = LineChart()
            ln.add_data(Reference(ws, min_col=3, max_col=ncol, min_row=r_hdr, max_row=r_end), titles_from_data=True)
            m_ln = mau_series(names[1:], [th["ink"], th["nhan"]])
            for sl, n in zip(ln.series, names[1:]):
                sl.graphicalProperties.line = _net(m_ln[str(n)], 2)
                sl.smooth = False
                sl.marker.symbol = "none"
            ch += ln
    else:
        return 0
    if len(names) > 6 and ch.legend is not None:        # chu thich 4 cot o duoi nhu app
        cao_cm += 0.5 * np.ceil(len(names) / 4)
    rec = dict(rec, kieu=dict(k, _n_cat=r_end - r_hdr))
    _dinh_dang_chart(ch, rec, vals, la_ngay, cot_ngang)
    ch.height, ch.width = cao_cm, rong_cm
    ws.add_chart(ch, anchor)
    return cao_cm


# ---------------------------------------------------------------- BANG + KPI
def _ten_sheet(ten, dung):
    ten = re.sub(r"[\[\]\:\*\?\/\\]", "-", ten)[:31].strip() or "Sheet"
    goc, i = ten, 2
    while ten.lower() in dung:
        hau = f"~{i}"
        ten = goc[:31 - len(hau)] + hau
        i += 1
    dung.add(ten.lower())
    return ten


def _chuan_df(df: pd.DataFrame):
    """-> (df phang, so cot index da dua vao bang)."""
    df = df.copy()
    if isinstance(df.columns, pd.MultiIndex):
        df.columns = [" | ".join(str(x) for x in c if str(x) != "") for c in df.columns]
    df.columns = [str(c) for c in df.columns]
    n_idx = 0
    if isinstance(df.index, pd.MultiIndex) or not isinstance(df.index, pd.RangeIndex):
        n_idx = df.index.nlevels
        df = df.reset_index()
        df.columns = ["Ngày" if (c in ("index", "date") and pd.api.types.is_datetime64_any_dtype(df[c]))
                      else ("Mục" if c == "index" else str(c)) for c in df.columns]
    return df, n_idx


def _gia_tri(v):
    if v is None:
        return None
    if isinstance(v, (pd.Timestamp, datetime)):
        return None if pd.isna(v) else pd.Timestamp(v).to_pydatetime().replace(tzinfo=None)
    if isinstance(v, (np.floating, float)):
        return None if not np.isfinite(v) else float(v)
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (np.bool_,)):
        return bool(v)
    if isinstance(v, pd.Period):
        return str(v)
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    return v if isinstance(v, (int, float, str, bool)) else str(v)


def _fmt_cot(s: pd.Series, ep=None) -> str:
    """Dinh dang so cua 1 cot: Styler precision neu co, khong thi theo do lon (nhu _fmt app)."""
    if pd.api.types.is_integer_dtype(s):          # so nguyen: precision cua Styler khong ap
        return "#,##0"
    if ep is not None:
        return "#,##0" + ("." + "0" * ep if ep else "")
    v = pd.to_numeric(s, errors="coerce").replace([np.inf, -np.inf], np.nan).dropna()
    if v.empty:
        return "General"
    if (v == v.round()).all():
        return "#,##0"
    m = v.abs().max()
    return "#,##0.00" if m < 10 else ("#,##0.0" if m < 1000 else "#,##0")


def _css_mau(v):
    v = str(v).strip()
    if v.startswith("#"):
        return _h(v) if len(v) >= 7 else "".join(ch * 2 for ch in v.lstrip("#"))[:6].upper()
    m = re.match(r"rgba?\((\d+),\s*(\d+),\s*(\d+)", v)
    return "".join(f"{int(x):02X}" for x in m.groups()) if m else None


def _ghi_bang(ws, r0, df, rec=None):
    """Ghi bang giong st.dataframe cua app (header xam nhat, chu ink, so theo do lon, mau Styler).
    Tra (dong header, dong cuoi, so cot, cot dau la ngay?)."""
    th = _th()
    toi = _h(th["bg"]) != "FFFFFF"
    nen_hdr = "262730" if toi else "F0F2F6"
    chu, chu2 = _h(th["ink"]), _h(th["ink2"])
    fnt = lambda **k: Font(name=FONT, size=10, **k)  # noqa: E731
    la_ngay = [pd.api.types.is_datetime64_any_dtype(df[c]) for c in df.columns]
    so = [pd.api.types.is_numeric_dtype(df[c]) and not pd.api.types.is_bool_dtype(df[c]) for c in df.columns]
    for j, c in enumerate(df.columns, 1):          # cot so can phai, chu/ngay can trai (nhu st.dataframe)
        o = ws.cell(r0, j, c)
        o.font, o.fill = fnt(bold=True, color=chu2), PatternFill("solid", fgColor=nen_hdr)
        o.alignment = Alignment(wrap_text=True, vertical="center",
                                horizontal="right" if so[j - 1] else "left")
    ep = (rec or {}).get("do_chinh_xac")
    fmts = [fmt_ngay_bang(df[c]) if la_ngay[j] else (_fmt_cot(df[c], ep) if so[j] else None)
            for j, c in enumerate(df.columns)]
    nen_bang = PatternFill("solid", fgColor=_h(th["bg"])) if toi else None
    for i, row in enumerate(df.itertuples(index=False), r0 + 1):
        for j, v in enumerate(row, 1):
            o = ws.cell(i, j, _gia_tri(v))
            o.font = fnt(color=chu)
            if nen_bang:
                o.fill = nen_bang
            if fmts[j - 1]:
                o.number_format = fmts[j - 1]
    # mau / dam / nen tu Styler cua app (to xanh-do theo dau, background_gradient...)
    n_idx, css = (rec or {}).get("n_idx", 0), (rec or {}).get("css") or {}
    for (ri, ci), props in css.items():
        rr, cc = r0 + 1 + ri, 1 + n_idx + ci
        if rr > r0 + len(df) or cc > len(df.columns):
            continue
        o = ws.cell(rr, cc)
        mau_chu, dam, nen = None, False, None
        for p, val in props:
            if p == "color":
                mau_chu = _css_mau(val)
            elif p == "font-weight":
                dam = str(val) in ("bold", "600", "700", "800")
            elif p in ("background-color", "background"):
                nen = _css_mau(val)
        o.font = fnt(color=mau_chu or chu, bold=dam)
        if nen:
            o.fill = PatternFill("solid", fgColor=nen)
    for j, c in enumerate(df.columns, 1):
        w = 12 if la_ngay[j - 1] else min(max(10, len(str(c)) * 0.9 + 2), 30)
        L = get_column_letter(j)
        ws.column_dimensions[L].width = max(ws.column_dimensions[L].width or 0, w)
    return r0, r0 + len(df), len(df.columns), bool(la_ngay and la_ngay[0])


def fmt_ngay_bang(s: pd.Series) -> str:
    """Ngay trong bang: chuoi thang/quy -> mm/yyyy, con lai dd/mm/yyyy."""
    d = pd.to_datetime(s, errors="coerce").dropna()
    if len(d) > 2 and (d.dt.day > 25).all() or (len(d) > 2 and (d.dt.day == 1).all()):
        return "mm/yyyy"
    return "dd/mm/yyyy"


def _ghi_kpi(ws, r, kpis):
    """KPI giong st.metric: nhan xam nho, gia tri dam lon, delta xanh/do co mui ten."""
    th = _th()
    ws.cell(r, 1, "Chỉ số nhanh").font = Font(name=FONT, bold=True, size=12, color=_h(th["ink"]))
    r += 1
    moi_hang = 5
    for i, x in enumerate(kpis):
        hang, cot = divmod(i, moi_hang)
        rr, cc = r + hang * 4, 1 + cot * 2
        ws.cell(rr, cc, x["tieu_de"]).font = Font(name=FONT, size=9, color=_h(th["ink2"]))
        ws.cell(rr + 1, cc, x["gia_tri"]).font = Font(name=FONT, size=16, color=_h(th["ink"]))
        d = x["delta"].strip()
        if d:
            am = d.startswith(("-", "−"))
            duong = d.startswith("+") or (d[:1].isdigit())
            mau = th["neg"] if am else (th["pos"] if duong else th["ink2"])
            ws.cell(rr + 2, cc, ("↓ " if am else "↑ " if duong else "") + d).font = Font(
                name=FONT, size=9, color=_h(mau))
        for c2 in (cc, cc + 1):
            L = get_column_letter(c2)
            ws.column_dimensions[L].width = max(ws.column_dimensions[L].width or 0, 13)
    return r + (len(kpis) - 1) // moi_hang * 4 + 4


def dung_excel(trang_rec: dict, cai_dat: dict) -> bytes:
    """trang_rec: {ten trang: [ban ghi]} -> bytes xlsx (Muc luc + moi (trang, tab) 1 sheet)."""
    th = _th()
    INK, INK2 = _h(th["ink"]), _h(th["ink2"])
    wb = Workbook()
    ml = wb.active
    ml.title = "Mục lục"
    dung = {"mục lục"}
    ml["A1"] = "Market Data — xuất từ app"
    ml["A1"].font = Font(name=FONT, bold=True, size=18, color=INK)
    dong = 3
    for k, v in [("Thời điểm xuất", f"{datetime.now():%d/%m/%Y %H:%M}")] + list(cai_dat.get("hien_thi", {}).items()):
        ml.cell(dong, 1, k).font = Font(name=FONT, bold=True, color=INK2)
        ml.cell(dong, 2, str(v)).font = Font(name=FONT, color=INK)
        dong += 1
    if cai_dat.get("trang_loi") not in (None, "không"):
        ml.cell(dong, 1, "Trang bị bỏ qua (lỗi)").font = Font(name=FONT, bold=True, color=_h(th["neg"]))
        ml.cell(dong, 2, cai_dat["trang_loi"])
        dong += 1
    dong += 1
    for j, h in enumerate(["Trang", "Tab", "Nội dung", "Loại", "Sheet", "Đi tới"], 1):
        o = ml.cell(dong, j, h)
        o.font, o.fill = Font(name=FONT, bold=True, color=INK2), PatternFill("solid", fgColor="F0F2F6")
    dong += 1
    TEN_LOAI = {"line": "Biểu đồ đường", "bar": "Biểu đồ cột", "stack": "Cột chồng", "hbar": "Cột ngang",
                "bar_line": "Cột + đường", "nguong": "Đường + ngưỡng", "custom_line": "Biểu đồ đường",
                "custom_bar": "Biểu đồ cột", "treemap": "Bản đồ nhiệt (chỉ data)", "bang": "Bảng"}

    for page, recs in trang_rec.items():
        ten_tr = _ten_trang(page)
        theo_tab = {}
        for r in recs:
            theo_tab.setdefault(r["tab"], []).append(r)
        for tab, rs in theo_tab.items():
            sn = _ten_sheet(f"{TAT.get(ten_tr, ten_tr[:4])}-{tab.split(' / ')[-1] if tab != 'Chung' else ten_tr}", dung)
            ws = wb.create_sheet(sn)
            ws.sheet_view.showGridLines = False                     # nen tron nhu app
            if _h(th["bg"]) != "FFFFFF":
                ws.sheet_properties.tabColor = _h(th["nhan"])
            ws["A1"] = f"{ten_tr} › {tab}"
            ws["A1"].font = Font(name=FONT, bold=True, size=18, color=INK)
            ws["A2"] = cai_dat.get("dong_mo_ta", "")
            ws["A2"].font = Font(name=FONT, italic=True, size=9, color=INK2)
            r = 4
            kpis = [x for x in rs if x["loai"] == "kpi"]
            if kpis:
                r = _ghi_kpi(ws, r, kpis) + 1
            for rec in [x for x in rs if x["loai"] != "kpi"]:
                df, n_idx = _chuan_df(rec["df"])
                rec = dict(rec, n_idx=n_idx)
                cat = len(df) > MAX_DONG
                if cat:
                    df = df.tail(MAX_DONG)
                    rec["css"] = {(i - (len(rec["df"]) - MAX_DONG), j): v for (i, j), v in (rec.get("css") or {}).items()
                                  if i >= len(rec["df"]) - MAX_DONG}
                o = ws.cell(r, 1, rec["tieu_de"])
                o.font = Font(name=FONT, bold=True, size=12, color=INK)
                ml_row = r
                ghi_chu = TEN_LOAI.get(rec["loai"], rec["loai"])
                if rec.get("don_vi"):
                    ghi_chu += f" · đơn vị: {rec['don_vi']}"
                if cat:
                    ghi_chu += f" · chỉ giữ {MAX_DONG:,} dòng cuối"
                ws.cell(r + 1, 1, ghi_chu).font = Font(name=FONT, italic=True, size=9, color=INK2)
                r_hdr, r_end, ncol, la_ngay = _ghi_bang(ws, r + 2, df, rec)
                cao_cm = 0
                if rec["loai"] not in ("bang", "treemap"):
                    try:
                        cao_cm = _ve_chart(ws, rec, df, r_hdr, r_end, ncol, la_ngay,
                                           f"{get_column_letter(ncol + 2)}{r}") or 0
                    except Exception as e:  # noqa: BLE001
                        ws.cell(r + 1, 4, f"(không dựng được chart: {e})")
                r = max(r_end, r + int(cao_cm / 0.53) + 1) + 3   # 1 dong mac dinh ~0,53 cm
                vals = [ten_tr, tab, rec["tieu_de"], TEN_LOAI.get(rec["loai"], rec["loai"]), sn]
                for j, v in enumerate(vals, 1):
                    ml.cell(dong, j, v).font = Font(name=FONT, color=INK)
                o = ml.cell(dong, 6, "→ mở")
                o.hyperlink = f"#'{sn}'!A{ml_row}"
                o.font = Font(name=FONT, color="0563C1", underline="single")
                dong += 1
            ws.freeze_panes = "A4"
            ws.sheet_view.zoomScale = 90
    for L, w in zip("ABCDEF", [24, 30, 60, 20, 26, 8]):
        ml.column_dimensions[L].width = w
    ml.sheet_view.showGridLines = False
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
